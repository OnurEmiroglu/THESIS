"""Generate thesis_30.docx from thesis_29.docx.

Adds a standalone Literature Review / Related Work chapter after the
Introduction and renumbers later chapters. This script does not run
experiments, regenerate figures, or modify protected evidence.
"""

from __future__ import annotations

import io
import re
import shutil
import sys
from pathlib import Path

from docx import Document
from docx.text.paragraph import Paragraph

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")

SRC = Path("manuscript/thesis_29.docx")
DST = Path("manuscript/thesis_30.docx")
LIT_DRAFT = Path("docs/internal/literature_review_chapter_draft_2026-05-27.md")


def _paragraph_after(paragraph: Paragraph, text: str, style: str | None = None) -> Paragraph:
    new_para = paragraph._parent.add_paragraph(text, style=style)
    paragraph._p.addnext(new_para._p)
    return new_para


def _paragraph_before(doc: Document, ref_para: Paragraph, text: str, style: str | None = None) -> Paragraph:
    new_para = doc.add_paragraph(text, style=style)
    ref_para._p.addprevious(new_para._p)
    return new_para


def _heading_before(doc: Document, ref_para: Paragraph, text: str, level: int) -> Paragraph:
    new_para = doc.add_heading(text, level=level)
    ref_para._p.addprevious(new_para._p)
    return new_para


def _find_paragraph(doc: Document, exact_text: str) -> Paragraph:
    for para in doc.paragraphs:
        if para.text.strip() == exact_text:
            return para
    raise RuntimeError(f"Paragraph not found: {exact_text}")


def _replace_text(para: Paragraph, text: str) -> None:
    if para.runs:
        para.runs[0].text = text
        for run in para.runs[1:]:
            run.text = ""
    else:
        para.add_run(text)


def _fix_mojibake(text: str) -> str:
    replacements = {
        "LÄ°TERATÃœR TARAMASI": "LİTERATÜR TARAMASI",
        "GuÃ©ant": "Guéant",
        "GaÅ¡perov": "Gašperov",
        "KostanjÄar": "Kostanjčar",
        "Ä‡": "č",
        "Å¡": "š",
        "Ã©": "é",
        "â€“": "-",
        "â€”": "-",
        "â€˜": "'",
        "â€™": "'",
        "â€œ": '"',
        "â€": '"',
    }
    for bad, good in replacements.items():
        text = text.replace(bad, good)
    return text


def _load_literature_review() -> list[tuple[int, str]]:
    """Return (heading_level, text) where level 0 means normal paragraph."""
    raw = LIT_DRAFT.read_text(encoding="utf-8-sig")
    raw = _fix_mojibake(raw)
    raw = raw.replace("# 3. LİTERATÜR TARAMASI", "# 3. LİTERATÜR TARAMASI / RELATED WORK")
    blocks: list[tuple[int, str]] = []
    paragraphs = [p.strip() for p in raw.split("\n\n") if p.strip()]
    for para in paragraphs:
        if para.startswith("## "):
            blocks.append((2, para[3:].strip()))
        elif para.startswith("# "):
            blocks.append((1, para[2:].strip()))
        else:
            text = " ".join(line.strip() for line in para.splitlines())
            blocks.append((0, text))
    return _shorten_and_polish_lit(blocks)


def _shorten_and_polish_lit(blocks: list[tuple[int, str]]) -> list[tuple[int, str]]:
    polished: list[tuple[int, str]] = []
    in_33 = False
    removed_one_33_sentence = False
    for level, text in blocks:
        if level == 2:
            in_33 = text.startswith("3.3 ")
        if in_33 and level == 0 and not removed_one_33_sentence:
            text = text.replace(
                "A market-making policy that ignores changes in market conditions may therefore be poorly calibrated to risk.",
                "",
            ).strip()
            removed_one_33_sentence = True
        if text.startswith("The thesis uses this logic to avoid relying only on conventional significance tests."):
            text = text.replace(
                "TOST supports this interpretation by testing whether the performance difference lies inside a predefined interval of practical equivalence.",
                (
                    "In plain terms, TOST asks whether the entire plausible range of the observed difference is small enough "
                    "to fall within a pre-specified practical-equivalence band."
                ),
            )
        polished.append((level, text))
    return polished


def _renumber_existing_headings(doc: Document) -> None:
    heading_map = {
        "3. TEORİ VE METODOLOJİ": "4. TEORİ VE METODOLOJİ",
        "4. SONUÇLAR VE TARTIŞMA": "5. SONUÇLAR VE TARTIŞMA",
        "5. SİNYAL BİLGİLENDİRİCİLİK SÜPÜRMESİ": "6. SİNYAL BİLGİLENDİRİCİLİK SÜPÜRMESİ",
        "6. SONUÇ": "7. SONUÇ",
    }
    for para in doc.paragraphs:
        text = para.text.strip()
        style = para.style.name if para.style else ""
        if "Heading" not in style:
            continue
        if text in heading_map:
            _replace_text(para, heading_map[text])
            continue
        match = re.match(r"^([3-6])\.(\d+(?:\.\d+)?)\s+(.*)$", text)
        if match:
            chapter = int(match.group(1))
            suffix = match.group(2)
            title = match.group(3)
            _replace_text(para, f"{chapter + 1}.{suffix} {title}")


def _revise_introduction(doc: Document) -> None:
    old_as = (
        "Klasik yaklaşımlar arasında Avellaneda ve Stoikov (2008) tarafından geliştirilen stokastik kontrol çerçevesi öne çıkmaktadır. "
        "Bu model, envanter riskini rezervasyon fiyatı kavramıyla içselleştirerek optimal bid-ask spread'ini analitik olarak türetmektedir. "
        "Ancak bu yaklaşım, sabit bir volatilite varsayımına dayanmakta ve piyasa rejimlerindeki değişimlere uyum sağlayamamaktadır."
    )
    new_as = (
        "Klasik piyasa yapıcılığı literatürü, özellikle Avellaneda-Stoikov çerçevesi, envanter riski ve volatilitenin kotasyon kararlarındaki merkezi rolünü göstermektedir. "
        "Bu tez bu geleneği analitik bir referans strateji olarak korurken, asıl olarak öğrenilmiş PPO politikalarında volatilite bilgisinin nasıl temsil edildiğini sınamaktadır."
    )
    old_rq = (
        "Bu çalışmanın temel araştırma sorusu şudur: Volatilite rejim bilgisine erişimi olan bir PPO ajanı, bu bilgiden yoksun olan eşdeğerine kıyasla daha iyi performans sergiler mi? "
        "Bu soruyu yanıtlamak için aşağıdaki katkılar sunulmaktadır:"
    )
    new_rq = (
        "Bu çalışmanın temel araştırma sorusu şudur: PPO politikasının gözleminde sürekli bir gerçekleşmiş volatilite vekili olan sigma_hat zaten mevcutken, açık kategorik volatilite-rejim etiketi sağlam ve ek performans değeri üretir mi? "
        "Bu soruyu yanıtlamak için aşağıdaki katkılar sunulmaktadır:"
    )
    for para in doc.paragraphs:
        if para.text.strip() == old_as:
            _replace_text(para, new_as)
        elif para.text.strip() == old_rq:
            _replace_text(para, new_rq)
    item2 = _find_paragraph(
        doc,
        "(2) Rejim farkındalığı ablasyonu gerçekleştirilmiş; 5 PPO varyantı 20 bağımsız tohum ile OOS protokolüyle karşılaştırılmıştır.",
    )
    _replace_text(
        item2,
        "(2) Gözlem uzayı ablasyonu gerçekleştirilmiş; sigma_hat, tahmini rejim etiketi, gerçek rejim etiketi ve bunların birleşimlerini kullanan 5 PPO varyantı 20 bağımsız tohum ile OOS protokolüyle karşılaştırılmıştır.",
    )
    item4 = _find_paragraph(
        doc,
        "(4) Detector robustness analizi gerçekleştirilmiş; null sonucun dedektör seçiminden bağımsızlığı doğrulanmıştır.",
    )
    _replace_text(
        item4,
        "(4) Detector robustness, rejime koşullu ödül, model misspecification ve sinyal bilgilendiriciliği analizleriyle bulgunun duyarlılığı değerlendirilmiştir.",
    )
    _paragraph_after(
        item4,
        (
            "Elde edilen kanıtlar, test edilen sentetik HFMM ortamında açık kategorik rejim etiketlerinin sigma_hat'in ötesinde sağlam bir ek performans değeri sağlamadığı ve bulguların sinyal yedekliliği yorumu ile tutarlı olduğu yönündedir. "
            "Bu sonuç volatilite rejimlerinin önemsiz olduğu anlamına gelmez; yalnızca bu kontrollü gözlem tasarımında kategorik etiketin ek değerini sınırlar."
        ),
    )


def _insert_literature_review(doc: Document) -> None:
    ref = _find_paragraph(doc, "4. TEORİ VE METODOLOJİ")
    for level, text in reversed(_load_literature_review()):
        if level == 1:
            ref = _heading_before(doc, ref, text, level=1)
        elif level == 2:
            ref = _heading_before(doc, ref, text, level=2)
        else:
            ref = _paragraph_before(doc, ref, text)


def _append_bibliography(doc: Document) -> None:
    refs = [
        "[7] Guéant, O., Lehalle, C.-A., & Fernandez-Tapia, J. (2013). Dealing with the inventory risk: a solution to the market making problem. Mathematics and Financial Economics, 7, 477-507.",
        "[8] Fodra, P., & Labadie, M. (2012). High-frequency market-making with inventory constraints and directional bets. Preprint, arXiv:1206.4810.",
        "[9] Fodra, P., & Labadie, M. (2013). High-frequency market-making for multi-dimensional Markov processes. Preprint, arXiv:1303.7177.",
        "[10] Spooner, T., & Savani, R. (2020). Robust market making via adversarial reinforcement learning. Proceedings of the Twenty-Ninth International Joint Conference on Artificial Intelligence, 4590-4596.",
        "[11] Gašperov, B., & Kostanjčar, Z. (2021). Market making with signals through deep reinforcement learning. IEEE Access, 9, 61611-61622.",
        "[12] Gašperov, B., Begušić, S., Posedel Šimović, P., & Kostanjčar, Z. (2021). Reinforcement learning approaches to optimal market making. Mathematics, 9(21), 2689.",
        "[13] Gašperov, B., & Kostanjčar, Z. (2022). Deep reinforcement learning for market making under a Hawkes process-based limit order book model. IEEE Control Systems Letters, 6, 2485-2490.",
        "[14] Zimmer, R., & Costa, O. L. V. (2025). Reinforcement learning-based market making as a stochastic control on non-stationary limit order book dynamics. Preprint, arXiv:2509.12456.",
        "[15] Lakens, D. (2017). Equivalence tests: A practical primer for t tests, correlations, and meta-analyses. Social Psychological and Personality Science, 8(4), 355-362.",
        "[16] Lakens, D., Scheel, A. M., & Isager, P. M. (2018). Equivalence testing for psychological research: A tutorial. Advances in Methods and Practices in Psychological Science, 1(2), 259-269.",
    ]
    appendix = _find_paragraph(doc, "Appendix B: Source Code File Index")
    ref_para = appendix
    for ref in reversed(refs):
        ref_para = _paragraph_before(doc, ref_para, ref)


def _update_internal_references(doc: Document) -> None:
    replacements = {
        "§3.3 ve §4.6": "§4.3 ve §5.6",
        "Section 4.1'deki": "Section 5.1'deki",
        "Section 4.7'de": "Section 5.7'de",
        "§3.1'de tanımlanan": "§4.1'de tanımlanan",
    }
    for para in doc.paragraphs:
        text = para.text
        new_text = text
        for old, new in replacements.items():
            new_text = new_text.replace(old, new)
        if new_text != text:
            _replace_text(para, new_text)


def _sanity_check(doc: Document) -> None:
    headings = [p.text.strip() for p in doc.paragraphs if p.style and "Heading" in p.style.name]
    required = [
        "2. GİRİŞ",
        "3. LİTERATÜR TARAMASI / RELATED WORK",
        "4. TEORİ VE METODOLOJİ",
        "5. SONUÇLAR VE TARTIŞMA",
        "6. SİNYAL BİLGİLENDİRİCİLİK SÜPÜRMESİ",
        "7. SONUÇ",
    ]
    missing = [h for h in required if h not in headings]
    if missing:
        raise RuntimeError(f"Missing required headings: {missing}")
    full_text = "\n".join(p.text for p in doc.paragraphs)
    unsafe = [
        "regimes are useless",
        "labels contain zero information",
        "categorical interference is conclusively proven",
        "findings generalize to all markets",
        "live-trading deployment",
        "LÄ°TERATÃœR",
        "GuÃ©ant",
        "GaÅ¡perov",
        "KostanjÄ",
    ]
    hits = [s for s in unsafe if s in full_text]
    if hits:
        raise RuntimeError(f"Unsafe or mojibake text found: {hits}")
    stale_refs = ["§3.3 ve §4.6", "Section 4.1'deki", "Section 4.7'de", "§3.1'de tanımlanan"]
    stale_hits = [s for s in stale_refs if s in full_text]
    if stale_hits:
        raise RuntimeError(f"Stale internal references found: {stale_hits}")


def main() -> None:
    if not SRC.exists():
        raise FileNotFoundError(SRC)
    shutil.copy2(SRC, DST)
    doc = Document(DST)
    _revise_introduction(doc)
    _renumber_existing_headings(doc)
    _insert_literature_review(doc)
    _append_bibliography(doc)
    _update_internal_references(doc)
    _sanity_check(doc)
    doc.save(DST)
    print(f"Wrote {DST}")


if __name__ == "__main__":
    main()
