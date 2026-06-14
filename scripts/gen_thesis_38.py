"""Generate thesis_38.docx/pdf from thesis_37 with final typography polish.

This documentation-only generator does not run experiments, regenerate
protected CSVs, or regenerate performance figures. It copies thesis_37 and
applies only equation typography, plus-minus notation, administrative
placeholder wording, and manuscript-version consistency updates.
"""

from __future__ import annotations

import hashlib
import shutil
import subprocess
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt


ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "manuscript" / "thesis_37.docx"
DST = ROOT / "manuscript" / "thesis_38.docx"
PDF = ROOT / "manuscript" / "thesis_38.pdf"

PROTECTED_HASHES = {
    ROOT / "results" / "metrics_detector_compare.csv": "28E7AD40BB47214F8576132846E9E1D4CD643F623CF1187743091FC367A206ED",
    ROOT / "docs" / "internal" / "wp6_sweep_full" / "summary_condition_variant.csv": "6DD627E81637A49A60163F58AC1D3EF23B8D694E39AC55BA64FBF808E978C6EA",
    ROOT / "docs" / "internal" / "wp6_sweep_full" / "summary_paired_combined_vs_sigma.csv": "4BABCAAACE1DD5228C674E2CED9D977236F8D3ACB503C098FAEB06FF6C10B796",
    ROOT / "docs" / "internal" / "wp6_sweep_full" / "summary_paired_combined_vs_regime.csv": "2087FEFBE5DC39AF23372EA2D8999AC0F1071D0FEE90BC1B3668F2158130E8F9",
}

SECTION9_EQUATIONS = {
    "(16)": "(16)   Sharpe-like = mean(ΔWₜ) / std(ΔWₜ) × √(1/Δt)",
    "(17)": "(17)   inv_p99 = percentile₉₉(|qₜ|)",
    "(18)": "(18)   fill_rate = filled quote opportunities / total quote opportunities",
    "(19)": "(19)   t = mean(dᵢ) / (sd(dᵢ) / √n)",
}

VERSION_REPLACEMENTS = {
    "The final thesis_37 manuscript embeds existing figures": (
        "The final thesis_38 manuscript embeds existing figures"
    ),
    "The final thesis_37 manuscript reuses frozen performance results.": (
        "The final thesis_38 manuscript reuses frozen performance results."
    ),
    (
        "final manuscript manuscript/thesis_37.pdf and .docx; "
        "source manuscript manuscript/thesis_36.pdf and .docx"
    ): (
        "final manuscript manuscript/thesis_38.pdf and .docx; "
        "source manuscript manuscript/thesis_37.pdf and .docx"
    ),
    "thesis_37 finalization did not rerun experiments": (
        "thesis_38 finalization did not rerun experiments"
    ),
    "this final thesis_37 generation": "this final thesis_38 generation",
    "Status in thesis_37": "Status in thesis_38",
    (
        "thesis_37 is the final administrative and visual cleanup built on thesis_36; "
        "frozen evidence remains unchanged."
    ): (
        "thesis_38 is the final formatting-consistency pass built on thesis_37; "
        "frozen evidence remains unchanged."
    ),
    "Documented validation; not rerun for thesis_37": (
        "Documented validation; not rerun for thesis_38"
    ),
}


def iter_all_paragraphs(doc: Document):
    for paragraph in doc.paragraphs:
        yield paragraph
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs
    for section in doc.sections:
        yield from section.header.paragraphs
        yield from section.footer.paragraphs


def has_drawing(paragraph) -> bool:
    return any(run._r.xpath(".//w:drawing") for run in paragraph.runs)


def set_paragraph_text(paragraph, text: str) -> None:
    if has_drawing(paragraph) or paragraph.text == text:
        return
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            if not run._r.xpath(".//w:drawing"):
                run.text = ""
    else:
        paragraph.add_run(text)


def restyle_section9_equations(doc: Document) -> None:
    found = set()
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        number = next(
            (key for key in SECTION9_EQUATIONS if text.startswith(key + " ")),
            None,
        )
        if number is None:
            continue
        set_paragraph_text(paragraph, SECTION9_EQUATIONS[number])
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_before = Pt(6)
        paragraph.paragraph_format.space_after = Pt(4)
        paragraph.paragraph_format.keep_with_next = True
        for run in paragraph.runs:
            run.font.name = "Cambria Math"
            run.font.size = Pt(11)
            run.bold = True
            r_fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
            r_fonts.set(qn("w:ascii"), "Cambria Math")
            r_fonts.set(qn("w:hAnsi"), "Cambria Math")
        found.add(number)
    missing = sorted(set(SECTION9_EQUATIONS) - found)
    if missing:
        raise RuntimeError("Missing Section 9 equations: " + ", ".join(missing))


def standardize_plus_minus(doc: Document) -> int:
    replacements = 0
    for paragraph in iter_all_paragraphs(doc):
        if "+/-" not in paragraph.text or has_drawing(paragraph):
            continue
        run_replacements = 0
        for run in paragraph.runs:
            count = run.text.count("+/-")
            if count:
                run.text = run.text.replace("+/-", "±")
                run_replacements += count
        if run_replacements == 0:
            count = paragraph.text.count("+/-")
            set_paragraph_text(paragraph, paragraph.text.replace("+/-", "±"))
            run_replacements = count
        replacements += run_replacements
    return replacements


def clean_title_placeholders(doc: Document) -> list[str]:
    values = {
        "Supervisor": "[To be provided]",
        "Matriculation No.": "[To be provided]",
        "Course / Module": "[To be provided]",
        "Advisor / Co-supervisor": "[To be provided, or Not applicable]",
    }
    cover = doc.tables[0]
    found = []
    for row in cover.rows:
        label = row.cells[0].text.strip()
        if label in values:
            row.cells[1].text = values[label]
            found.append(label)
    missing = sorted(set(values) - set(found))
    if missing:
        raise RuntimeError("Missing title-page fields: " + ", ".join(missing))
    return found


def update_version_references(doc: Document) -> None:
    for paragraph in iter_all_paragraphs(doc):
        if has_drawing(paragraph):
            continue
        new_text = paragraph.text
        for old, new in VERSION_REPLACEMENTS.items():
            new_text = new_text.replace(old, new)
        if new_text != paragraph.text:
            set_paragraph_text(paragraph, new_text)


def verify_protected_hashes() -> None:
    matches = 0
    for path, expected in PROTECTED_HASHES.items():
        actual = hashlib.sha256(path.read_bytes()).hexdigest().upper()
        if actual != expected:
            raise RuntimeError(
                f"Protected artifact hash mismatch: {path.relative_to(ROOT)}"
            )
        matches += 1
    print(f"Protected CSV hashes: {matches}/{len(PROTECTED_HASHES)} MATCH")


def export_pdf_with_word() -> None:
    tmp_pdf = PDF.with_name(PDF.stem + "_export_tmp.pdf")
    ps = (
        "$ErrorActionPreference = 'Stop'; "
        f"$tmp = '{str(tmp_pdf)}'; "
        "if (Test-Path -LiteralPath $tmp) { Remove-Item -LiteralPath $tmp -Force }; "
        "$word = New-Object -ComObject Word.Application; "
        "$word.Visible = $false; "
        "try { "
        f"$doc = $word.Documents.Open('{str(DST)}', $false, $false); "
        "$doc.Fields.Update() | Out-Null; "
        "foreach ($toc in $doc.TablesOfContents) { $toc.Update() | Out-Null }; "
        "$doc.Save(); "
        "$doc.ExportAsFixedFormat($tmp, 17); "
        "$doc.Close($false); "
        "} finally { $word.Quit() }; "
        f"Move-Item -LiteralPath $tmp -Destination '{str(PDF)}' -Force"
    )
    subprocess.run(
        ["powershell", "-Command", ps],
        cwd=str(ROOT),
        check=True,
        capture_output=True,
        text=True,
    )


def main() -> int:
    if not SRC.exists():
        raise FileNotFoundError(SRC)
    verify_protected_hashes()
    shutil.copy2(SRC, DST)

    doc = Document(DST)
    before_shapes = len(doc.inline_shapes)
    restyle_section9_equations(doc)
    plus_minus_count = standardize_plus_minus(doc)
    placeholders = clean_title_placeholders(doc)
    update_version_references(doc)

    after_shapes = len(doc.inline_shapes)
    if after_shapes != before_shapes:
        raise RuntimeError(
            f"Figure preservation failed: inline_shapes {before_shapes} -> {after_shapes}"
        )

    doc.save(DST)
    print(f"Wrote {DST.relative_to(ROOT)}")
    print(f"Preserved inline_shapes: {after_shapes}")
    print(f"Standardized +/- occurrences: {plus_minus_count}")
    print("Fields requiring user input: " + ", ".join(placeholders))

    if sys.platform.startswith("win"):
        export_pdf_with_word()
        print(f"Wrote {PDF.relative_to(ROOT)}")
    else:
        print("PDF export skipped: Microsoft Word automation requires Windows.")

    verify_protected_hashes()
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
