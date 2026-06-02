"""Generate a Turkish reading/audit copy of thesis_34.

This script creates manuscript/thesis_34_TR.docx and attempts to export
manuscript/thesis_34_TR.pdf. It is a translation artifact only: it does not run
experiments, regenerate figures, or modify the official English thesis_34 files.
"""

from __future__ import annotations

import shutil
import subprocess
import sys
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "manuscript" / "thesis_34.docx"
DST = ROOT / "manuscript" / "thesis_34_TR.docx"
PDF = ROOT / "manuscript" / "thesis_34_TR.pdf"


TRANSLATIONS = {
    "High-Frequency Market Making via Reinforcement Learning under Different Volatility Regimes": "Farklı Volatilite Rejimleri Altında Pekiştirmeli Öğrenme ile Yüksek Frekanslı Piyasa Yapıcılığı",
    "Technical Report Template Adaptation / MSc Thesis Draft": "Teknik Rapor Şablonu Uyarlaması / MSc Tez Taslağı",
    "Draft integrity note: thesis_34 is a documentation and defense-clarity polish pass. Frozen experimental evidence is reused and not regenerated.": "Taslak bütünlüğü notu: thesis_34 bir dokümantasyon ve savunma açıklığı cilalama geçişidir. Dondurulmuş deneysel kanıt yeniden kullanılmıştır ve yeniden üretilmemiştir.",
    "In the tested controlled synthetic HFMM environment, explicit categorical volatility-regime labels do not provide robust incremental value once sigma_hat is already observed by the PPO policy. The results are consistent with signal redundancy.": "Test edilen kontrollü sentetik HFMM ortamında, PPO politikası sigma_hat sinyalini zaten gözlemlerken açık kategorik volatilite-rejimi etiketleri güvenilir bir ek performans katkısı sağlamamaktadır. Sonuçlar sinyal fazlalığı (signal redundancy) yorumu ile tutarlıdır.",
    "Abstract": "Özet",
    "This report studies whether explicit volatility-regime information improves a Proximal Policy Optimization (PPO)-based high-frequency market-making (HFMM) agent in a controlled synthetic limit-order-book (LOB) environment. The simulator uses an arithmetic Brownian-motion mid-price, a sticky Markov volatility-regime process, rolling realized-volatility detection, and Poisson-arrival fills. The PPO agent chooses bid and ask quotes through a discrete half-spread and skew parameterization, and it is compared with a fixed-spread strategy and an Avellaneda-Stoikov analytical baseline.": "Bu rapor, açık volatilite-rejimi bilgisinin kontrollü sentetik bir limit emir defteri (limit order book, LOB) ortamında Proximal Policy Optimization (PPO) tabanlı bir yüksek frekanslı piyasa yapıcılığı (HFMM) ajanını iyileştirip iyileştirmediğini inceler. Simülatör aritmetik Brownian hareketli orta fiyatı, kalıcı geçişlere sahip bir Markov volatilite-rejimi sürecini, kayan gerçekleşmiş volatilite tespitini ve Poisson-varışlı dolumları kullanır. PPO ajanı, ayrık yarı-spread ve skew parametreleştirmesiyle alış ve satış kotasyonlarını seçer; sabit-spread stratejisi ve Avellaneda-Stoikov analitik referans modeli ile karşılaştırılır.",
    "In the tested controlled synthetic HFMM environment, explicit categorical volatility-regime labels do not provide robust incremental value once sigma_hat is already observed by the PPO policy. The results are consistent with signal redundancy. The evidence comes from the canonical frozen experiments: out-of-sample evaluation, detector robustness, oracle-label ablations, regime-conditional reward shaping, mild model misspecification, and a signal-informativeness sweep. In the five-variant ablation, sigma_only achieved the highest mean Sharpe-like value (0.753), while oracle_full did not provide a statistically significant or practically meaningful improvement over sigma_hat alone (paired t-test p = 0.115; TOST +/-0.10 p = 0.00067; 90% CI [-0.001, +0.063]). The strongest interpretation is therefore conditional and synthetic-market bounded: the explicit regime label mostly repackages volatility information already present in the continuous signal. The report makes no production-deployment claim and does not claim to prove the internal PPO mechanism.": "Test edilen kontrollü sentetik HFMM ortamında, PPO politikası sigma_hat sinyalini zaten gözlemlerken açık kategorik volatilite-rejimi etiketleri güvenilir bir ek performans katkısı sağlamamaktadır. Sonuçlar sinyal fazlalığı (signal redundancy) yorumu ile tutarlıdır. Kanıt, kanonik dondurulmuş deneylerden gelir: örneklem dışı (OOS) değerlendirme, dedektör seçimine karşı sağlamlık, oracle-etiket ablasyonları, rejime koşullu ödül şekillendirme, hafif model yanlış-belirlenimi ve sinyalin bilgi değerini test eden tarama. Beş-varyant ablasyonda sigma_only en yüksek ortalama Sharpe-like değerine ulaşmıştır (0.753); oracle_full ise yalnızca sigma_hat'e göre istatistiksel olarak anlamlı ya da pratik olarak anlamlı bir iyileşme sağlamamıştır (eşleştirilmiş t-testi p = 0.115; TOST +/-0.10 p = 0.00067; 90% CI [-0.001, +0.063]). Bu nedenle en güçlü yorum koşullu ve sentetik piyasa ile sınırlıdır: açık rejim etiketi çoğunlukla sürekli sinyalde zaten bulunan volatilite bilgisini yeniden paketlemektedir. Rapor gerçek işlem/production ortamında kullanım iddiası taşımaz ve PPO'nun iç mekanizmasını kanıtladığını iddia etmez.",
    "Symbols, Abbreviations, and Glossary": "Semboller, Kısaltmalar ve Sözlük",
    "1 Introduction": "1 Giriş",
    "1.1 Problem Statement and Motivation": "1.1 Problem Tanımı ve Motivasyon",
    "Market makers continuously choose bid and ask quotes while balancing spread capture against inventory risk. When volatility changes, the same quote width can become either too passive or too exposed. A natural response is to provide a learning agent with regime labels such as low, medium, and high volatility. The central question is whether that categorical label adds decision-relevant information when the policy already observes a continuous volatility estimate.": "Piyasa yapıcılar (market makers), spread kazancını envanter riskiyle dengelerken sürekli olarak alış ve satış kotasyonları seçer. Volatilite değiştiğinde aynı kotasyon genişliği ya fazla pasif ya da fazla riskli hâle gelebilir. Doğal bir tepki, öğrenen ajana düşük, orta ve yüksek volatilite gibi rejim etiketleri sağlamaktır. Temel soru, politika zaten sürekli bir volatilite tahmini gözlemlerken bu kategorik etiketin karar açısından ilgili ek bilgi sağlayıp sağlamadığıdır.",
    "1.2 Research Objective": "1.2 Araştırma Amacı",
    "The objective is to test the incremental value of explicit categorical volatility-regime labels in PPO market making under a controlled synthetic environment. The comparison is intentionally information-design focused: the regime label is evaluated against a continuous sigma_hat signal, not against an observation space with no volatility information.": "Amaç, kontrollü sentetik bir ortamda PPO piyasa yapıcılığında açık kategorik volatilite-rejimi etiketlerinin ilave bilgi değeri taşıyıp taşımadığını test etmektir. Karşılaştırma bilinçli olarak gözlem/sinyal tasarımı odaklıdır: rejim etiketi, volatilite bilgisi hiç içermeyen bir gözlem uzayına karşı değil, sürekli sigma_hat sinyaline karşı değerlendirilir.",
    "1.3 Background and Related Work": "1.3 Arka Plan ve İlgili Çalışmalar",
    "The literature relevant to this thesis falls into four connected strands: classical inventory-aware market making, reinforcement-learning market making, volatility regimes and non-stationarity, and statistical interpretation of small or practically negligible differences. The purpose of the review is not to claim a new analytical quoting rule or a new reinforcement-learning algorithm. It is to locate a narrower information-design question inside an established market-making literature.": "Bu tezle ilgili literatür dört bağlantılı alanda toplanır: klasik envanter-duyarlı piyasa yapıcılığı, pekiştirmeli öğrenme ile piyasa yapıcılığı, volatilite rejimleri ve piyasa koşullarının zamanla değişmesi (non-stationarity), ayrıca küçük ya da pratik olarak ihmal edilebilir farkların istatistiksel yorumu. Bu incelemenin amacı yeni bir analitik kotasyon kuralı ya da yeni bir pekiştirmeli öğrenme algoritması iddia etmek değildir. Amaç, daha dar bir gözlem/sinyal tasarımı sorusunu yerleşik piyasa yapıcılığı literatürü içinde konumlandırmaktır.",
    "Classical market-making models provide the inventory-risk and quote-placement foundation. A market maker earns spread income by posting bid and ask quotes, but stochastic execution creates inventory exposure whose value changes with the mid-price. Avellaneda and Stoikov (2008) formalize this tradeoff through a reservation price and spread term that depend on inventory, risk aversion, volatility, remaining horizon, and order-arrival intensity. Gueant, Lehalle and Fernandez-Tapia (2013) extend this inventory-control view and provide tractable approximations for optimal quotes under inventory constraints.": "Klasik piyasa yapıcılığı modelleri envanter riski ve kotasyon yerleşimi temelini sağlar. Bir piyasa yapıcı alış ve satış kotasyonları göndererek spread geliri elde eder; ancak stokastik gerçekleşme, değeri orta fiyatla değişen envanter maruziyeti yaratır. Avellaneda ve Stoikov (2008), bu ödünleşimi envanter, riskten kaçınma, volatilite, kalan ufuk ve emir-varış yoğunluğuna bağlı bir rezervasyon fiyatı ve spread terimi üzerinden biçimselleştirir. Gueant, Lehalle ve Fernandez-Tapia (2013), bu envanter-kontrol bakışını genişletir ve envanter kısıtları altında optimal kotasyonlar için hesaplanabilir yaklaşımlar sunar.",
    "This analytical literature motivates the thesis design in two direct ways: it explains why inventory, volatility, quote distance, and fill intensity are central state variables, and it supplies the Avellaneda-Stoikov baseline used as a non-learning comparator. Related preprints by Fodra and Labadie (2012, 2013) are useful additional context for inventory-constrained and state-dependent market-making models, but they are treated cautiously as preprints. The thesis does not attempt to derive a new closed-form quoting rule; it uses the analytical tradition to frame the simulation and benchmark design.": "Bu analitik literatür tez tasarımını iki doğrudan yolla motive eder: envanter, volatilite, kotasyon mesafesi ve dolum yoğunluğunun neden merkezi durum değişkenleri olduğunu açıklar ve öğrenmeyen karşılaştırıcı olarak kullanılan Avellaneda-Stoikov referans modelini sağlar. Fodra ve Labadie'nin (2012, 2013) ilgili ön baskıları envanter kısıtlı ve durum-bağımlı piyasa yapıcılığı modelleri için yararlı ek bağlam sunar; ancak ön baskı oldukları için ihtiyatla ele alınır. Tez yeni bir kapalı-form kotasyon kuralı türetmeye çalışmaz; simülasyon ve benchmark tasarımını çerçevelemek için analitik geleneği kullanır.",
    "A second strand replaces closed-form rules with learned policies. Spooner et al. (2018) demonstrate that reinforcement learning can learn market-making behavior in a simulated limit-order-book setting and show why reward design and inventory control matter. Spooner and Savani (2020) extend the discussion toward robustness under adversarially varied market conditions. These studies support the use of simulator-based evaluation for learned quoting policies, while also showing that market-making performance can be sensitive to the assumptions used to generate training and evaluation episodes.": "İkinci alan, kapalı-form kuralların yerine öğrenilmiş politikaları koyar. Spooner vd. (2018), pekiştirmeli öğrenmenin simüle edilmiş bir limit emir defteri ortamında piyasa yapıcılığı davranışını öğrenebildiğini gösterir ve ödül tasarımı ile envanter kontrolünün neden önemli olduğunu ortaya koyar. Spooner ve Savani (2020), tartışmayı sistematik olarak değiştirilen piyasa koşulları altında sağlamlığa doğru genişletir. Bu çalışmalar, öğrenilmiş kotasyon politikaları için simülatör tabanlı değerlendirmenin kullanımını destekler; aynı zamanda piyasa yapıcılığı performansının eğitim ve değerlendirme bölümlerini üretmek için kullanılan varsayımlara duyarlı olabileceğini gösterir.",
    "Deep reinforcement-learning studies with market signals are closest to the present information-design question. Gasperov and Kostanjcar (2021) study market making with signals through deep reinforcement learning; Gasperov et al. (2021) review reinforcement-learning approaches to optimal market making; and Gasperov and Kostanjcar (2022) train agents in a richer Hawkes-process limit-order-book simulator. This literature motivates learned quoting, state representation, reward shaping, robustness checks, and simulator-based evidence. The present thesis is narrower: PPO is used as a controlled testbed for whether one additional volatility-regime channel has incremental value, not as a claim of algorithmic optimality.": "Piyasa sinyalleriyle derin pekiştirmeli öğrenme çalışmaları, mevcut gözlem/sinyal tasarımı sorusuna en yakın çalışmalardır. Gasperov ve Kostanjcar (2021) derin pekiştirmeli öğrenme yoluyla sinyalli piyasa yapıcılığını inceler; Gasperov vd. (2021) optimal piyasa yapıcılığına yönelik pekiştirmeli öğrenme yaklaşımlarını gözden geçirir; Gasperov ve Kostanjcar (2022) ise daha zengin bir Hawkes-süreci limit emir defteri simülatöründe ajanlar eğitir. Bu literatür öğrenilmiş kotasyonu, durum temsilini, ödül şekillendirmeyi, sağlamlık kontrollerini ve simülatör tabanlı kanıtı motive eder. Mevcut tez daha dardır: PPO, algoritmik optimalite iddiası olarak değil, ek bir volatilite-rejimi kanalının ilave bilgi değeri olup olmadığını test eden kontrollü bir test yatağı olarak kullanılır.",
    "Volatility regimes and non-stationarity are economically meaningful. A high-volatility state can increase inventory risk, change fill-risk tradeoffs, and make a quote width that was appropriate in a calm period too aggressive in a turbulent period. The thesis therefore models a latent three-state Markov volatility process and a causal rolling realized-volatility detector. This design should not be read as a claim that volatility regimes are irrelevant. The tested question is more specific: whether categorical regime labels add incremental value once the policy already observes a continuous realized-volatility estimate.": "Volatilite rejimleri ve piyasa koşullarının zamanla değişmesi (non-stationarity) ekonomik olarak anlamlıdır. Yüksek volatilite durumu envanter riskini artırabilir, dolum-risk ödünleşimlerini değiştirebilir ve sakin bir dönemde uygun olan kotasyon genişliğini çalkantılı bir dönemde fazla agresif hâle getirebilir. Bu nedenle tez, gizil üç-durumlu bir Markov volatilite süreci ve nedensel kayan gerçekleşmiş-volatilite dedektörü modeller. Bu tasarım, volatilite rejimlerinin önemsiz olduğu iddiası olarak okunmamalıdır. Test edilen soru daha özgüldür: politika zaten sürekli bir gerçekleşmiş volatilite tahmini gözlemlerken kategorik rejim etiketlerinin ek fayda sağlayıp sağlamadığıdır.",
    "The key state-representation issue is redundancy. The continuous signal sigma_hat gives the PPO policy a graded estimate of recent realized volatility. A low/medium/high label derived from related volatility information may help if it summarizes latent structure that sigma_hat does not expose, but it may add little if the continuous signal already contains the useful variation for quote placement. This gives the exact gap addressed by the thesis: When a PPO market-making policy already observes a continuous realized-volatility estimate, does adding an explicit categorical volatility-regime label improve out-of-sample performance in a controlled synthetic HFMM environment?": "Temel durum-temsili sorunu fazlalıktır. Sürekli sigma_hat sinyali PPO politikasına yakın dönem gerçekleşmiş volatilitenin dereceli bir tahminini verir. İlgili volatilite bilgisinden türetilen düşük/orta/yüksek etiketi, sigma_hat'in açığa çıkarmadığı gizil yapıyı özetliyorsa yardımcı olabilir; ancak sürekli sinyal kotasyon yerleşimi için yararlı değişimi zaten içeriyorsa az katkı sağlayabilir. Bu, tezin ele aldığı kesin boşluğu verir: Bir PPO piyasa yapıcılığı politikası zaten sürekli bir gerçekleşmiş-volatilite tahmini gözlemlerken açık kategorik volatilite-rejimi etiketi eklemek kontrollü sentetik HFMM ortamında örneklem dışı performansı iyileştirir mi?",
    "Finally, the empirical interpretation requires care because the thesis evaluates an incremental signal, not the existence of volatility effects in general. Lakens (2017) and Lakens, Scheel and Isager (2018) motivate equivalence testing as a way to distinguish a merely non-significant difference from evidence that a difference is small relative to a pre-specified practical bound. The TOST framing is therefore used as background for the later statistics section; it does not by itself prove that labels contain no information.": "Son olarak, ampirik yorum dikkat gerektirir; çünkü tez genel olarak volatilite etkilerinin varlığını değil, artı bir sinyali değerlendirir. Lakens (2017) ve Lakens, Scheel ve Isager (2018), yalnızca anlamlı olmayan bir farkı, önceden belirlenmiş pratik bir sınıra göre küçük bir farka ilişkin kanıttan ayırmanın yolu olarak eşdeğerlik testini motive eder. Bu nedenle TOST çerçevesi sonraki istatistik bölümünün arka planı olarak kullanılır; tek başına etiketlerin bilgi içermediğini kanıtlamaz.",
    "1.4 Contributions": "1.4 Katkılar",
    "A controlled HFMM simulator with Poisson-arrival fills, fees, latency, and Markov volatility regimes.": "Poisson-varışlı dolumlar, ücretler, gecikme ve Markov volatilite rejimleri içeren kontrollü bir HFMM simülatörü.",
    "A Gymnasium environment in which PPO controls quote half-spread and skew.": "PPO'nun kotasyon yarı-spread ve skew değerlerini kontrol ettiği bir Gymnasium ortamı.",
    "An OOS comparison of fixed-spread, Avellaneda-Stoikov, regime-aware PPO, and regime-blind PPO strategies.": "Sabit-spread, Avellaneda-Stoikov, rejim-farkında PPO ve rejim-kör PPO stratejilerinin OOS karşılaştırması.",
    "A five-variant ablation separating continuous volatility information from estimated and oracle regime labels.": "Sürekli volatilite bilgisini tahmini ve oracle rejim etiketlerinden ayıran beş-varyant ablasyon.",
    "Detector, reward-shaping, misspecification, and signal-informativeness checks that bound the interpretation.": "Yorumu sınırlandıran dedektör seçimi, ödül şekillendirme, yanlış-belirlenim ve sinyal bilgi-değeri kontrolleri.",
    "1.5 Scope and Limitations": "1.5 Kapsam ve Sınırlılıklar",
    "The thesis is a controlled synthetic-market study. It does not claim production deployment viability, real-order-book external validity, or a proven PPO representation mechanism. The result is conditional on the simulator, signal design, PPO hyperparameters, and degradation calibration tested here.": "Tez kontrollü bir sentetik piyasa çalışmasıdır. Gerçek işlem/production ortamında kullanılabilirlik, gerçek limit emir defteri dış geçerliliği ya da kanıtlanmış bir PPO temsil mekanizması iddia etmez. Sonuç burada test edilen simülatöre, sinyal tasarımına, PPO hiperparametrelerine ve sinyal bozulması kalibrasyonuna koşulludur.",
    "1.6 Report Structure": "1.6 Rapor Yapısı",
    "Sections 2-6 define the mathematical model, use case, architecture, layered model, and algorithms. Sections 7-9 describe the experimental setup, results, and metrics. Sections 10-12 discuss the interpretation, reproducibility, and conclusion. Appendices provide code maps, sanity checks, and extended evidence.": "Bölüm 2-6 matematiksel modeli, kullanım senaryosunu, mimariyi, katmanlı modeli ve algoritmaları tanımlar. Bölüm 7-9 deneysel kurulumu, sonuçları ve metrikleri açıklar. Bölüm 10-12 yorumu, yeniden üretilebilirliği ve sonucu tartışır. Ekler kod haritaları, sağduyu kontrolleri ve genişletilmiş kanıt sunar.",
    "2 Mathematical Formulation": "2 Matematiksel Formülasyon",
    "2.1 Synthetic Mid-Price Process": "2.1 Sentetik Orta Fiyat Süreci",
    "Here S_t is the mid-price, sigma_t is the step volatility in ticks or price units according to the simulator configuration, dt is the time step, and epsilon_t is an independent standard-normal shock.": "Burada S_t orta fiyattır; sigma_t simülatör yapılandırmasına göre tick ya da fiyat birimleri cinsinden adım volatilitesidir; dt zaman adımıdır; epsilon_t bağımsız standart-normal şoktur.",
    "2.2 Markov Volatility-Regime Process": "2.2 Markov Volatilite-Rejimi Süreci",
    "Here z_t is the latent volatility regime at step t, L/M/H denote low, medium, and high volatility, and P_ij is the sticky transition probability from regime i to regime j.": "Burada z_t, t adımındaki gizil volatilite rejimidir; L/M/H düşük, orta ve yüksek volatiliteyi gösterir; P_ij rejim i'den rejim j'ye yapışkan geçiş olasılığıdır.",
    "Rows correspond to the current regime L/M/H, and columns correspond to the next regime L/M/H. The large diagonal probabilities make the regimes persistent, or sticky, so volatility states tend to last for many steps rather than switching erratically.": "Satırlar mevcut L/M/H rejimine, sütunlar bir sonraki L/M/H rejimine karşılık gelir. Büyük diyagonal olasılıklar rejimleri kalıcı, yani yapışkan, kılar; böylece volatilite durumları düzensiz biçimde değişmek yerine birçok adım sürme eğilimindedir.",
    "This persistence matters because rolling realized-volatility detection is only meaningful when regime-dependent volatility differences last long enough to be observed through a finite window.": "Bu kalıcılık önemlidir; çünkü kayan gerçekleşmiş-volatilite tespiti ancak rejime bağlı volatilite farkları sonlu bir pencere içinde gözlemlenebilecek kadar uzun sürdüğünde anlamlıdır.",
    "The regime controls the volatility multiplier applied to the base sigma parameter. The canonical full experiments use sigma multipliers [0.6, 1.0, 1.8] for L, M, and H.": "Rejim, temel sigma parametresine uygulanan volatilite çarpanını kontrol eder. Kanonik tam deneyler L, M ve H için sigma çarpanları [0.6, 1.0, 1.8] kullanır.",
    "2.3 Realized Volatility Signal and Regime Detection": "2.3 Gerçekleşmiş Volatilite Sinyali ve Rejim Tespiti",
    "Here sigma_hat_t is the rolling realized-volatility signal, w is the rolling window length, and Delta S_i is the mid-price increment over step i.": "Burada sigma_hat_t kayan gerçekleşmiş-volatilite sinyalidir; w kayan pencere uzunluğudur; Delta S_i ise i adımındaki orta fiyat artışıdır.",
    "Estimated regime labels are assigned by thresholding sigma_hat after a warmup period. The main reported training and evaluation pipelines use the causal rv_baseline detector. The rv_dwell detector is retained only as an auxiliary/offline robustness comparison, while the HMM detector is an additional robustness variant.": "Tahmini rejim etiketleri, ısınma döneminden sonra sigma_hat eşiklenerek atanır. Raporlanan ana eğitim ve değerlendirme boru hatları nedensel rv_baseline dedektörünü kullanır. rv_dwell dedektörü yalnızca yardımcı/çevrimdışı sağlamlık karşılaştırması olarak tutulur; HMM dedektörü ise ek bir sağlamlık varyantıdır.",
    "2.3.1 Illustrative Synthetic Environment Path": "2.3.1 Açıklayıcı Sentetik Ortam Patikası",
    "This subsection visualizes one deterministic synthetic path generated under the canonical HFMM environment configuration. The figure is included to make the simulated market environment transparent: the mid-price evolves under a regime-switching arithmetic Brownian-motion process, the realized-volatility signal is computed from rolling mid-price changes, and the latent Markov regime switches between low, medium, and high volatility states. The dashed vertical line marks the chronological 70/30 train-test split used in the out-of-sample protocol. The figure is illustrative of the methodology only and is not a new performance result.": "Bu alt bölüm, kanonik HFMM ortam yapılandırması altında üretilen deterministik bir sentetik patikayı görselleştirir. Şekil, simüle edilen piyasa ortamını şeffaf kılmak için eklenmiştir: orta fiyat rejim-değiştiren aritmetik Brownian hareket süreci altında evrilir, gerçekleşmiş-volatilite sinyali kayan orta fiyat değişimlerinden hesaplanır ve gizil Markov rejimi düşük, orta ve yüksek volatilite durumları arasında geçiş yapar. Kesikli dikey çizgi, örneklem dışı protokolde kullanılan kronolojik 70/30 eğitim-test ayrımını gösterir. Şekil yalnızca metodolojiyi açıklayıcıdır ve yeni bir performans sonucu değildir.",
    "Figure 2.1. Illustrative synthetic HFMM environment path. The upper panel shows the generated mid-price path. The middle panel compares the rolling realized-volatility signal observed by the policy with the true regime-dependent volatility level. The lower panel shows the latent Markov volatility-regime sequence. The dashed vertical line marks the chronological 70/30 train-test split used for out-of-sample evaluation. This figure is generated by scripts/figures/gen_synthetic_environment_figure.py and is a methodology illustration, not a regenerated performance artifact.": "Figure 2.1. Açıklayıcı sentetik HFMM ortam patikası. Üst panel üretilen orta fiyat patikasını gösterir. Orta panel, politikanın gözlemlediği kayan gerçekleşmiş-volatilite sinyalini gerçek rejime bağlı volatilite düzeyiyle karşılaştırır. Alt panel gizil Markov volatilite-rejimi dizisini gösterir. Kesikli dikey çizgi, örneklem dışı değerlendirme için kullanılan kronolojik 70/30 eğitim-test ayrımını işaretler. Bu şekil scripts/figures/gen_synthetic_environment_figure.py tarafından üretilmiştir ve yeniden üretilmiş bir performans artefaktı değil, metodoloji açıklamasıdır.",
    "2.4 Limit-Order Fill Model": "2.4 Limit Emir Dolum Modeli",
    "Here lambda(delta) is the fill intensity at quote distance delta, A is the baseline arrival scale, and k controls the exponential decay as quotes move farther from the mid-price.": "Burada lambda(delta), delta kotasyon mesafesindeki dolum yoğunluğudur; A temel varış ölçeğidir; k ise kotasyonlar orta fiyattan uzaklaştıkça üstel azalmayı kontrol eder.",
    "Intuitively, quotes closer to the mid-price fill more often because they are more attractive to incoming market orders. Quotes farther away earn more spread if they fill, but the exponential intensity makes those fills less likely.": "Sezgisel olarak, orta fiyata daha yakın kotasyonlar gelen piyasa emirleri için daha cazip olduklarından daha sık dolar. Daha uzaktaki kotasyonlar dolarsa daha fazla spread kazandırır; ancak üstel yoğunluk bu dolumları daha az olası kılar.",
    "Here P(fill | delta) is the per-step fill probability and dt is the simulator step length.": "Burada P(fill | delta) adım başına dolum olasılığıdır ve dt simülatör adım uzunluğudur.",
    "2.5 Quote Parameterization": "2.5 Kotasyon Parametreleştirmesi",
    "Here h_idx and m_idx are the two discrete PPO action components, h is the half-spread in ticks, and m is the quote skew in ticks.": "Burada h_idx ve m_idx iki ayrık PPO eylem bileşenidir; h tick cinsinden yarı-spread, m ise tick cinsinden kotasyon skew değeridir.",
    "The half-spread h controls quote width and aggressiveness: small h is more aggressive and large h is more passive. The skew m shifts the bid and ask distances asymmetrically, allowing the policy to lean against inventory or other state pressure while still quoting both sides.": "Yarı-spread h kotasyon genişliğini ve agresifliğini kontrol eder: küçük h daha agresif, büyük h daha pasiftir. Skew m alış ve satış mesafelerini asimetrik olarak kaydırır; böylece politika her iki tarafta da kotasyon verirken envantere ya da başka durum baskılarına karşı eğilebilir.",
    "Here delta_bid and delta_ask are the bid and ask quote distances in ticks. The max operator enforces a minimum quote distance of one tick.": "Burada delta_bid ve delta_ask tick cinsinden alış ve satış kotasyon mesafeleridir. max operatörü en az bir tick kotasyon mesafesini zorunlu kılar.",
    "2.6 Inventory, Cash, Equity, and Reward": "2.6 Envanter, Nakit, Özsermaye ve Ödül",
    "Here q_t is inventory, F_t^{bid} is the bid-side fill indicator or count, and F_t^{ask} is the ask-side fill indicator or count.": "Burada q_t envanterdir; F_t^{bid} alış tarafı dolum göstergesi ya da sayısıdır; F_t^{ask} satış tarafı dolum göstergesi ya da sayısıdır.",
    "Here X_t is cash, P_t^{bid} and P_t^{ask} are the executed bid and ask prices, and fees_t denotes transaction costs.": "Burada X_t nakittir; P_t^{bid} ve P_t^{ask} gerçekleşen alış ve satış fiyatlarıdır; fees_t işlem maliyetlerini gösterir.",
    "Here W_t is mark-to-market wealth or equity, X_t is cash, q_t is inventory, and S_t is the mid-price.": "Burada W_t piyasa fiyatına göre değerlenen servet ya da özsermayedir (mark-to-market wealth); X_t nakit, q_t envanter ve S_t orta fiyattır.",
    "The equity measure is marked to the mid-price for consistency across simulated strategies. This is an evaluation convention inside the controlled simulator and should not be interpreted as immediately liquidatable wealth in a real limit-order book. In practice, liquidating inventory would generally require crossing the spread, facing queue priority and available depth, and potentially incurring additional adverse-selection or market-impact costs. These effects are part of the real-market boundary discussed later, not part of the current controlled experiment.": "Özsermaye ölçüsü, simüle edilen stratejiler arasında tutarlılık için orta fiyata göre değerlenir. Bu, kontrollü simülatör içindeki bir değerlendirme konvansiyonudur ve gerçek bir limit emir defterinde hemen nakde çevrilebilir servet olarak yorumlanmamalıdır. Pratikte envanteri kapatmak genellikle spread'i geçmeyi, kuyruk önceliği ve mevcut derinlikle karşılaşmayı ve ek ters-seçim ya da piyasa-etkisi maliyetleri doğurmayı gerektirir. Bu etkiler daha sonra tartışılan gerçek-piyasa sınırının parçasıdır; mevcut kontrollü deneyin parçası değildir.",
    "Here R_t is the reward and eta is the inventory-penalty coefficient. Fees are included in the cash update and are not counted a second time in the reward.": "Burada R_t ödüldür ve eta envanter-cezası katsayısıdır. Ücretler nakit güncellemesine dahildir ve ödülde ikinci kez sayılmaz.",
    "The eta q^2 term discourages the agent from earning apparent PnL by accumulating large directional inventory. It connects the learning objective to market-making risk control rather than pure speculation on the mid-price path.": "eta q^2 terimi, ajanın büyük yönlü envanter biriktirerek görünürde PnL kazanmasını caydırır. Öğrenme amacını orta fiyat patikası üzerinde saf spekülasyondan ziyade piyasa yapıcılığı risk kontrolüne bağlar.",
    "2.7 Avellaneda-Stoikov Baseline": "2.7 Avellaneda-Stoikov Baz Çizgisi",
    "Here r_t is the reservation price, gamma is inventory-risk aversion, sigma is volatility, and tau is remaining horizon.": "Burada r_t rezervasyon fiyatıdır; gamma envanter-riskinden kaçınma katsayısıdır; sigma volatilitedir; tau kalan ufuktur.",
    "Here delta_AS is the AS half-spread and k is the same fill-intensity decay parameter used in the execution model. The implemented deltas are clipped to configured minimum and maximum bounds.": "Burada delta_AS, AS yarı-spread değeridir ve k yürütme modelinde kullanılan aynı dolum-yoğunluğu azalma parametresidir. Uygulanan delta değerleri yapılandırılmış alt ve üst sınırlara kırpılır.",
    "2.8 PPO Objective": "2.8 PPO Amacı",
    "Here rho_t(theta) is the PPO probability ratio, pi_theta is the current policy, pi_theta_old is the behavior policy used to collect the rollout, a_t is the action, and s_t is the state.": "Burada rho_t(theta) PPO olasılık oranıdır; pi_theta mevcut politika, pi_theta_old rollout toplamak için kullanılan davranış politikası, a_t eylem ve s_t durumdur.",
    "Here theta denotes policy parameters, rho_t(theta) is the PPO probability ratio, A_hat_t is the advantage estimate, and eps is the clipping parameter. The report uses this only as a concise training-objective reference.": "Burada theta politika parametrelerini, rho_t(theta) PPO olasılık oranını, A_hat_t avantaj tahminini ve eps kırpma parametresini gösterir. Rapor bunu yalnızca kısa bir eğitim-amacı referansı olarak kullanır.",
    "PPO is used because the thesis needs a stable, widely used policy-gradient algorithm for stochastic sequential control, not because it claims PPO is the best possible market-making RL algorithm. The action is discrete but two-dimensional MultiDiscrete: the agent chooses half-spread h and skew m, which together form a 25-action quote grid. DQN would require value-based handling of this grid and is less natural for the policy-gradient control framing used here; SAC is more naturally aligned with continuous actions and would expand the scope; and A2C is simpler but often less stable or sample-efficient than PPO. PPO therefore provides a reliable baseline algorithm for controlled ablation of the observation channels.": "PPO, tezin stokastik sıralı kontrol için istikrarlı ve yaygın kullanılan bir politika-gradyanı algoritmasına ihtiyaç duyması nedeniyle kullanılır; PPO'nun mümkün olan en iyi piyasa yapıcılığı RL algoritması olduğu iddia edildiği için değil. Eylem ayrık fakat iki boyutlu MultiDiscrete yapıdadır: ajan yarı-spread h ve skew m seçer; bunlar birlikte 25-eylemli bir kotasyon ızgarası oluşturur. DQN bu ızgaranın değer-tabanlı ele alınmasını gerektirir ve burada kullanılan politika-gradyanı kontrol çerçevesi için daha az doğaldır; SAC sürekli eylemlerle daha doğal uyumludur ve kapsamı genişletir; A2C daha basittir fakat çoğu zaman PPO'dan daha az kararlı ya da örnek-verimli olabilir. Bu nedenle PPO, gözlem kanallarının kontrollü ablasyonu için güvenilir bir baz algoritma sağlar.",
    "3 Use-Case Scenario": "3 Kullanım Senaryosu",
    "3.1 Synthetic HFMM Scenario": "3.1 Sentetik HFMM Senaryosu",
    "The use case is a synthetic market maker posting one bid and one ask quote at each step. The environment supplies mid-price dynamics, volatility estimates, and regime labels; the agent receives stochastic fills and is evaluated on wealth, risk-adjusted performance, inventory tails, and fill behavior.": "Kullanım senaryosu, her adımda bir alış ve bir satış kotasyonu gönderen sentetik bir piyasa yapıcıdır. Ortam orta fiyat dinamiklerini, volatilite tahminlerini ve rejim etiketlerini sağlar; ajan stokastik dolumlar alır ve servet, risk-ayarlı performans, envanter kuyrukları ve dolum davranışı üzerinden değerlendirilir.",
    "3.2 Trading Agents and Strategy Variants": "3.2 İşlem Ajanları ve Strateji Varyantları",
    "3.3 Information-Design Question": "3.3 Bilgi-Tasarımı Sorusu",
    "4 System Architecture": "4 Sistem Mimarisi",
    "4.1 Run Lifecycle and Reproducibility Layer": "4.1 Çalıştırma Yaşam Döngüsü ve Yeniden Üretilebilirlik Katmanı",
    "`run.py` dispatches jobs from JSON configuration files, creates a timestamped run directory, snapshots the config, records git metadata, writes logs and metrics, and finalizes run status. Long WP6 jobs support resume validation against the saved config snapshot.": "`run.py`, işleri JSON yapılandırma dosyalarından yönlendirir, zaman damgalı bir çalışma dizini oluşturur, yapılandırmanın anlık görüntüsünü alır, git meta verisini kaydeder, log ve metrikleri yazar ve çalışma durumunu sonlandırır. Uzun WP6 işleri, kaydedilmiş yapılandırma anlık görüntüsüne karşı devam doğrulamasını destekler.",
    "4.2 Synthetic Market and Detector Layer": "4.2 Sentetik Piyasa ve Dedektör Katmanı",
    "`src/wp1/sim.py` implements the fill and inventory simulator. `src/wp2/synth_regime.py` generates synthetic regime paths, mid-prices, rolling realized volatility, and detector outputs.": "`src/wp1/sim.py` dolum ve envanter simülatörünü uygular. `src/wp2/synth_regime.py` sentetik rejim patikalarını, orta fiyatları, kayan gerçekleşmiş volatiliteyi ve dedektör çıktılarını üretir.",
    "4.3 Gymnasium Environment Layer": "4.3 Gymnasium Ortam Katmanı",
    "`src/wp3/env.py` wraps the simulator as a Gymnasium environment. The observation vector is [q_norm, sigma_hat, tau, regime_L, regime_M, regime_H]. Warmup, invalid, or disabled regime labels produce a zero one-hot vector rather than an artificial medium label.": "`src/wp3/env.py` simülatörü bir Gymnasium ortamı olarak sarar. Gözlem vektörü [q_norm, sigma_hat, tau, regime_L, regime_M, regime_H] şeklindedir. Isınma, geçersiz ya da devre dışı rejim etiketleri yapay bir orta etiket yerine sıfır one-hot vektörü üretir.",
    "4.4 Strategy and Training Layer": "4.4 Strateji ve Eğitim Katmanı",
    "The training layer fits PPO policies, the main evaluation layer compares strategies and ablations, and the signal-informativeness layer tests whether regime labels gain value as sigma_hat is degraded. The fixed-spread and AS baselines provide non-learning comparators.": "Eğitim katmanı PPO politikalarını eğitir; ana değerlendirme katmanı stratejileri ve ablasyonları karşılaştırır; sinyal bilgi-değeri katmanı ise sigma_hat bozuldukça rejim etiketlerinin değer kazanıp kazanmadığını test eder. Sabit-spread ve AS referans stratejileri öğrenmeyen karşılaştırıcılar sağlar.",
    "4.5 Evaluation and Evidence Layer": "4.5 Değerlendirme ve Kanıt Katmanı",
    "Evaluation writes CSV metrics, figures, and summaries. The thesis_34 adaptation embeds existing figures and quotes frozen numerical evidence; it does not regenerate protected evidence artifacts.": "Değerlendirme CSV metrikleri, figürler ve özetler yazar. thesis_34 uyarlaması mevcut figürleri belgeye ekler ve dondurulmuş sayısal kanıtı alıntılar; korunan kanıt artefaktlarını yeniden üretmez.",
    "4.6 Methodology Pipeline Illustration": "4.6 Metodoloji Boru Hattı Gösterimi",
    "Figure 4.1 summarizes the evidence pipeline at a methodological level. A synthetic path generates the latent volatility regime and mid-price sequence; rolling realized volatility and detector labels form the observed signals; the Gymnasium environment converts these signals, inventory, time, quotes, and fills into sequential observations and rewards; PPO variants and analytical baselines are then evaluated on held-out out-of-sample metrics. The figure is included only to clarify the workflow and is not a regenerated performance artifact.": "Figure 4.1 kanıt boru hattını metodolojik düzeyde özetler. Sentetik bir patika gizil volatilite rejimini ve orta fiyat dizisini üretir; kayan gerçekleşmiş volatilite ve dedektör etiketleri gözlemlenen sinyalleri oluşturur; Gymnasium ortamı bu sinyalleri, envanteri, zamanı, kotasyonları ve dolumları sıralı gözlemlere ve ödüllere dönüştürür; ardından PPO varyantları ve analitik baz çizgileri elde tutulan örneklem dışı metrikler üzerinde değerlendirilir. Şekil yalnızca iş akışını açıklığa kavuşturmak için eklenmiştir ve yeniden üretilmiş bir performans artefaktı değildir.",
    "Figure 4.1. Methodology-only HFMM evaluation pipeline. Synthetic market paths produce sigma_hat and regime_hat signals, which feed the Gymnasium environment, PPO variants, baselines, and out-of-sample metrics. Generated by scripts/figures/gen_pipeline_architecture_figure.py; this illustration is not performance evidence.": "Figure 4.1. Yalnızca metodolojiye yönelik HFMM değerlendirme boru hattı. Sentetik piyasa patikaları sigma_hat ve regime_hat sinyallerini üretir; bu sinyaller Gymnasium ortamını, PPO varyantlarını, baz çizgilerini ve örneklem dışı metrikleri besler. scripts/figures/gen_pipeline_architecture_figure.py tarafından üretilmiştir; bu gösterim performans kanıtı değildir.",
    "5 Layered Experimental Model": "5 Katmanlı Deneysel Model",
    "Section 4 describes the software/evidence architecture, while this section restates the same workflow as a conceptual layered experimental model.": "Bölüm 4 yazılım/kanıt mimarisini açıklar; bu bölüm ise aynı iş akışını kavramsal bir katmanlı deneysel model olarak yeniden ifade eder.",
    "5.1 Layer 1: Market State Generation": "5.1 Katman 1: Piyasa Durumu Üretimi",
    "Generate Markov regimes and mid-price paths under controlled volatility multipliers.": "Kontrollü volatilite çarpanları altında Markov rejimlerini ve orta fiyat patikalarını üret.",
    "5.2 Layer 2: Signal Construction": "5.2 Katman 2: Sinyal İnşası",
    "Construct sigma_hat from rolling realized volatility and derive estimated regime labels.": "Kayan gerçekleşmiş volatiliteden sigma_hat oluştur ve tahmini rejim etiketlerini türet.",
    "5.3 Layer 3: Observation Encoding": "5.3 Katman 3: Gözlem Kodlama",
    "Encode inventory, volatility, time-to-horizon, and optional regime one-hot channels.": "Envanteri, volatiliteyi, ufka kalan zamanı ve isteğe bağlı rejim one-hot kanallarını kodla.",
    "5.4 Layer 4: Action and Execution": "5.4 Katman 4: Eylem ve Yürütme",
    "Decode PPO actions into bid/ask quote distances and sample Poisson-arrival fills.": "PPO eylemlerini alış/satış kotasyon mesafelerine çöz ve Poisson-varışlı dolumları örnekle.",
    "5.5 Layer 5: Learning and Evaluation": "5.5 Katman 5: Öğrenme ve Değerlendirme",
    "Train PPO on the chronological train segment and evaluate deterministic policies OOS.": "PPO'yu kronolojik eğitim kesitinde eğit ve deterministik politikaları OOS değerlendir.",
    "6 Algorithm Specification and Pseudocode": "6 Algoritma Belirtimi ve Sözde Kod",
    "6.1 Plain-Language Overview": "6.1 Sade Dil Özeti",
    "The experiment creates a synthetic market path, estimates volatility signals, trains policies on the first 70% of the path, and evaluates all strategies on the held-out 30%. Ablations modify which volatility channels the PPO policy can observe.": "Deney sentetik bir piyasa patikası oluşturur, volatilite sinyallerini tahmin eder, politikaları patikanın ilk 70%'lik bölümünde eğitir ve tüm stratejileri ayrılan 30%'luk test bölümünde değerlendirir. Ablasyonlar PPO politikasının hangi volatilite kanallarını gözlemleyebileceğini değiştirir.",
    "6.2 Synthetic Market and Regime Generation": "6.2 Sentetik Piyasa ve Rejim Üretimi",
    "Input: seed, transition matrix P, base volatility, volatility multipliers, n_steps": "Girdi: seed, geçiş matrisi P, temel volatilite, volatilite çarpanları, n_steps",
    "Initialize z_0 and S_0": "z_0 ve S_0 değerlerini başlat",
    "For t = 0 ... n_steps - 1:": "t = 0 ... n_steps - 1 için:",
    "    sample z_{t+1} from P[z_t]": "    z_{t+1} değerini P[z_t] içinden örnekle",
    "    set sigma_t = base_sigma * multiplier[z_t]": "    sigma_t = base_sigma * multiplier[z_t] olarak ayarla",
    "    sample mid-price increment and update S_{t+1}": "    orta fiyat artışını örnekle ve S_{t+1} değerini güncelle",
    "Compute rolling sigma_hat after warmup window": "Isınma penceresinden sonra kayan sigma_hat değerini hesapla",
    "Calibrate thresholds on warmup data and assign regime_hat": "Eşikleri ısınma verisi üzerinde kalibre et ve regime_hat ata",
    "Output: exogenous table with mid, sigma_hat, regime_true, regime_hat": "Çıktı: mid, sigma_hat, regime_true, regime_hat içeren dışsal tablo",
    "6.3 Market-Making Environment Step": "6.3 Piyasa Yapıcılığı Ortam Adımı",
    "Input: action (h_idx, m_idx), current simulator state, exogenous row": "Girdi: eylem (h_idx, m_idx), mevcut simülatör durumu, dışsal satır",
    "Decode h = h_idx + 1 and m = m_idx - 2": "h = h_idx + 1 ve m = m_idx - 2 olarak çöz",
    "Set delta_bid = max(1, h + m), delta_ask = max(1, h - m)": "delta_bid = max(1, h + m), delta_ask = max(1, h - m) olarak ayarla",
    "Compute fill intensities and per-step fill probabilities": "Dolum yoğunluklarını ve adım başına dolum olasılıklarını hesapla",
    "Sample bid and ask fills": "Alış ve satış dolumlarını örnekle",
    "Update inventory and cash, including fees": "Ücretler dahil envanter ve nakdi güncelle",
    "Advance mid-price and compute W_{t+1}": "Orta fiyatı ilerlet ve W_{t+1} değerini hesapla",
    "Return observation, reward = delta_equity - eta * inventory^2, done flag, info": "Gözlemi, reward = delta_equity - eta * inventory^2 değerini, done bayrağını ve info bilgisini döndür",
    "6.4 PPO Training and OOS Evaluation Protocol": "6.4 PPO Eğitimi ve OOS Değerlendirme Protokolü",
    "For each seed and variant:": "Her seed ve varyant için:",
    "    generate one exogenous synthetic path": "    bir dışsal sentetik patika üret",
    "    split path chronologically into train and test segments": "    patikayı kronolojik olarak eğitim ve test kesitlerine ayır",
    "    configure observation channels for the variant": "    varyant için gözlem kanallarını yapılandır",
    "    train PPO on the train segment for the configured timesteps": "    PPO'yu yapılandırılmış zaman adımı sayısı boyunca eğitim kesitinde eğit",
    "    evaluate the deterministic policy on the test segment": "    deterministik politikayı test kesitinde değerlendir",
    "    log Sharpe-like ratio, final equity, inventory p99, fill rate, and per-regime metrics": "    Sharpe-like oranını, final equity değerini, inventory p99 değerini, fill rate değerini ve rejim başına metrikleri kaydet",
    "Aggregate seed-paired statistics across variants": "Varyantlar arasında seed-eşleştirilmiş istatistikleri topla",
    "6.5 Experiment 5: Signal-Informativeness Sweep": "6.5 Deney 5: Sinyal Bilgi-Değeri Taraması",
    "For each degradation condition in {full, noisy, lagged, coarsened, none}:": "{full, noisy, lagged, coarsened, none} içindeki her sinyal bozulması koşulu için:",
    "    transform sigma_hat according to the condition": "    sigma_hat değerini koşula göre dönüştür",
    "    for each valid variant and seed:": "    her geçerli varyant ve seed için:",
    "        train PPO and evaluate OOS": "        PPO'yu eğit ve OOS değerlendir",
    "Aggregate condition-variant means and paired comparisons": "Koşul-varyant ortalamalarını ve eşleştirilmiş karşılaştırmaları topla",
    "Test whether combined gains value as sigma_hat is degraded": "sigma_hat bozuldukça combined varyantının değer kazanıp kazanmadığını test et",
    "6.6 Reference Implementation Map": "6.6 Referans Uygulama Haritası",
    "6.7 Complexity Analysis": "6.7 Karmaşıklık Analizi",
    "Simulation and deterministic evaluation are linear in the number of environment steps. PPO training cost is approximately linear in total timesteps, policy-network forward/backward passes, and the number of seed-variant cells. The WP6 full sweep is expensive because it multiplies conditions, variants, and seeds; this is why its outputs are treated as frozen evidence.": "Simülasyon ve deterministik değerlendirme ortam adımı sayısına göre doğrusaldır. PPO eğitim maliyeti yaklaşık olarak toplam zaman adımı sayısına, politika ağının ileri/geri geçişlerine ve seed-varyant hücrelerinin sayısına göre doğrusaldır. WP6 tam süpürmesi pahalıdır; çünkü koşulları, varyantları ve seed değerlerini çarpar. Bu nedenle çıktıları dondurulmuş kanıt olarak ele alınır.",
    "The implementation prioritizes reproducibility and auditability over maximum simulation throughput. Future versions could improve runtime through vectorized environments, parallel rollouts, or lower-level simulator components for the execution loop.": "Uygulama, maksimum simülasyon iş hacminden ziyade yeniden üretilebilirliği ve denetlenebilirliği önceliklendirir. Gelecek sürümler, vektörleştirilmiş ortamlar, paralel rollout'lar ya da yürütme döngüsü için daha düşük seviyeli simülatör bileşenleriyle çalışma süresini iyileştirebilir.",
    "6.8 Executable Example Commands": "6.8 Çalıştırılabilir Örnek Komutlar",
    "Commands": "Komutlar",
    "7 Experimental Setup": "7 Deneysel Kurulum",
    "7.1 Canonical Configuration Protocol": "7.1 Kanonik Yapılandırma Protokolü",
    "Canonical experiments are driven by JSON configs under `config/`. Shared parameters include mid0 = 100.0, tick_size = 0.01, dt = 0.2, baseline sigma_mid_ticks = 0.8, A = 5.0, k = 1.5, fee_bps = 0.2, and latency_steps = 1.": "Kanonik deneyler `config/` altındaki JSON yapılandırmalarıyla yürütülür. Ortak parametreler mid0 = 100.0, tick_size = 0.01, dt = 0.2, baseline sigma_mid_ticks = 0.8, A = 5.0, k = 1.5, fee_bps = 0.2 ve latency_steps = 1 değerlerini içerir.",
    "Market context. The simulator is asset-class agnostic and uses normalized price units. It should be read as a stylized liquid electronic limit-order-book market rather than as a calibrated model of a specific venue such as equities, futures, crypto, or BIST. The fee and latency parameters are controlled friction assumptions used to make strategy comparisons internally consistent. They are not intended to identify a particular exchange microstructure.": "Piyasa bağlamı. Simülatör varlık-sınıfından bağımsızdır ve normalize edilmiş fiyat birimleri kullanır. Hisse senetleri, vadeli işlemler, kripto ya da BIST gibi belirli bir piyasa yerinin kalibre edilmiş modeli olarak değil, stilize likit elektronik limit emir defteri piyasası olarak okunmalıdır. Ücret ve gecikme parametreleri, strateji karşılaştırmalarını iç tutarlı kılmak için kullanılan kontrollü sürtünme varsayımlarıdır. Belirli bir borsa mikro yapısını tanımlamayı amaçlamazlar.",
    "7.2 Data Generation and Preprocessing": "7.2 Veri Üretimi ve Ön İşleme",
    "Each run generates synthetic mid-price and regime paths, computes rolling realized volatility, and assigns detector labels. The main reported pipelines use causal rv_baseline labels.": "Her çalışma sentetik orta fiyat ve rejim patikaları üretir, kayan gerçekleşmiş volatiliteyi hesaplar ve dedektör etiketleri atar. Raporlanan ana boru hatları nedensel rv_baseline etiketlerini kullanır.",
    "7.3 Train/Test Split": "7.3 Eğitim/Test Ayrımı",
    "The main out-of-sample evaluation and signal-informativeness sweep use a chronological 70/30 split on the exogenous series. PPO trains on the first segment and is evaluated deterministically on the held-out OOS segment.": "Ana örneklem dışı değerlendirme ve sinyalin bilgi değerini test eden tarama, dışsal seri üzerinde kronolojik 70/30 ayrım kullanır. PPO ilk kesitte eğitilir ve elde tutulan OOS kesitinde deterministik olarak değerlendirilir.",
    "7.4 Strategy Variants": "7.4 Strateji Varyantları",
    "Experiment 1, the Main Out-of-Sample Evaluation, evaluates naive, AS, ppo_aware, and ppo_blind. Experiment 2, the Five-Variant Signal Ablation, and Experiment 5, the Signal-Informativeness Sweep, use sigma_only, regime_only, combined, oracle_pure, and oracle_full.": "Deney 1, Ana Örneklem Dışı Değerlendirme, naive, AS, ppo_aware ve ppo_blind stratejilerini değerlendirir. Deney 2, Beş-Varyant Sinyal Ablasyonu, ve Deney 5, Sinyal Bilgi-Değeri Taraması, sigma_only, regime_only, combined, oracle_pure ve oracle_full varyantlarını kullanır.",
    "7.5 Detector Variants": "7.5 Dedektör Varyantları",
    "7.6 PPO Hyperparameters": "7.6 PPO Hiperparametreleri",
    "The PPO hyperparameters are held fixed across PPO variants to keep the signal-channel comparison fair. They should be read as canonical experiment settings rather than as a claim of globally optimal PPO tuning. The inventory penalty eta = 0.001 is the canonical reward scale used in the main experiments after the inventory-penalty ablation.": "Sinyal-kanalı karşılaştırmasını adil tutmak için PPO hiperparametreleri PPO varyantları arasında sabit tutulur. Bunlar küresel olarak optimal PPO ayarı iddiası değil, kanonik deney ayarları olarak okunmalıdır. Envanter cezası eta = 0.001, envanter-cezası ablasyonundan sonra ana deneylerde kullanılan kanonik ödül ölçeğidir.",
    "7.7 Protected Evidence and No-Rerun Policy": "7.7 Korunan Kanıt ve Yeniden Çalıştırmama Politikası",
    "The thesis_34 adaptation reuses frozen performance results. It does not rerun PPO training, WP5/WP6 experiments, detector robustness, ablations, misspecification checks, protected CSV generation, or frozen performance-figure generation. The only newly generated visuals are methodology-only illustrations produced by `scripts/figures/gen_synthetic_environment_figure.py` and `scripts/figures/gen_pipeline_architecture_figure.py`.": "thesis_34 uyarlaması dondurulmuş performans sonuçlarını yeniden kullanır. PPO eğitimini, WP5/WP6 deneylerini, dedektör seçimine karşı sağlamlık analizlerini, ablasyonları, yanlış-belirlenim kontrollerini, korunan CSV üretimini ya da dondurulmuş performans şekillerini yeniden çalıştırmaz. Yeni üretilen tek görseller, `scripts/figures/gen_synthetic_environment_figure.py` ve `scripts/figures/gen_pipeline_architecture_figure.py` tarafından üretilen yalnızca metodolojiye yönelik gösterimlerdir.",
    "8 Results and Visualisation": "8 Sonuçlar ve Görselleştirme",
    "8.1 Experiment 1: Main Out-of-Sample Evaluation": "8.1 Deney 1: Ana Örneklem Dışı Değerlendirme",
    "Experiment 1, the 20-seed Main Out-of-Sample Evaluation, shows that PPO variants produce much higher risk-adjusted performance than the fixed-spread and AS baselines. AS can produce higher raw equity, but with substantially larger inventory exposure.": "Deney 1, 20-seed Ana Örneklem Dışı Değerlendirme, PPO varyantlarının sabit-spread ve AS baz çizgilerinden çok daha yüksek risk-ayarlı performans ürettiğini gösterir. AS daha yüksek ham özsermaye üretebilir; ancak bunu önemli ölçüde daha büyük envanter maruziyetiyle yapar.",
    "AS has higher raw equity but carries substantially larger inventory tail risk; therefore the main comparison is risk-adjusted performance and inventory control, not raw equity alone.": "AS daha yüksek ham özsermayeye sahiptir; ancak önemli ölçüde daha büyük envanter kuyruk riski taşır. Bu nedenle ana karşılaştırma yalnızca ham özsermaye değil, risk-ayarlı performans ve envanter kontrolüdür.",
    "The AS implementation is used as a canonical analytical inventory-aware comparator rather than as an exhaustively optimized trading system. The PPO-versus-AS comparison establishes a reference point for learned quoting and inventory control, but the central thesis claim is the within-PPO signal-design comparison between sigma_hat, estimated labels, and oracle labels.": "AS uygulaması, kapsamlı biçimde optimize edilmiş bir işlem sistemi olarak değil, kanonik analitik envanter-duyarlı karşılaştırıcı olarak kullanılır. PPO-AS karşılaştırması öğrenilmiş kotasyon ve envanter kontrolü için bir referans noktası kurar; ancak tezin merkezi iddiası sigma_hat, tahmini etiketler ve oracle etiketler arasındaki PPO-içi sinyal-tasarımı karşılaştırmasıdır.",
    "Figure 8.1. Experiment 1 main OOS results: PPO variants dominate naive and AS on risk-adjusted performance, while AS carries much larger inventory tail risk.": "Figure 8.1. Deney 1 ana OOS sonuçları: PPO varyantları risk-ayarlı performansta naive ve AS stratejilerine baskındır; AS ise çok daha büyük envanter kuyruk riski taşır.",
    "Figure 8.2. Seed-paired PPO-aware versus PPO-blind comparison: the estimated regime channel does not create a robust Sharpe-like advantage.": "Figure 8.2. Seed-eşleştirilmiş PPO-aware ve PPO-blind karşılaştırması: tahmini rejim kanalı sağlam bir Sharpe-like avantajı yaratmaz.",
    "The paired-seed view shows that the aware policy does not dominate the blind policy seed by seed. The equity panel even favors the blind policy in the canonical paired test.": "Seed-eşleştirilmiş görünüm, aware politikanın blind politikaya seed bazında baskın olmadığını gösterir. Özsermaye paneli kanonik eşleştirilmiş testte blind politikayı bile destekler.",
    "8.2 Five-Variant Ablation": "8.2 Beş-Varyant Ablasyon",
    "The strongest ablation result is that sigma_only has the highest mean Sharpe-like value, while oracle_full does not significantly beat it. TOST supports practical equivalence under the +/-0.10 Sharpe-like bound.": "En güçlü ablasyon sonucu, sigma_only varyantının en yüksek ortalama Sharpe-like değerine sahip olması ve oracle_full varyantının onu anlamlı biçimde geçmemesidir. TOST, +/-0.10 Sharpe-like sınırı altında pratik eşdeğerliği destekler.",
    "The label-quality objection is also bounded by the oracle_pure result. oracle_pure receives the true categorical regime label but not sigma_hat, yet it remains below sigma_only in mean Sharpe-like performance in the five-variant ablation. This suggests that the limitation is not only detector noise: in this calibration, the continuous volatility signal is more useful for quote control than the discretized regime category.": "Etiket-kalitesi itirazı da oracle_pure sonucu ile sınırlanır. oracle_pure gerçek kategorik rejim etiketini alır fakat sigma_hat almaz; buna rağmen beş-varyant ablasyonda ortalama Sharpe-like performansta sigma_only altında kalır. Bu, sınırlamanın yalnızca dedektör gürültüsü olmadığını düşündürür: bu kalibrasyonda sürekli volatilite sinyali kotasyon kontrolü için ayrıklaştırılmış rejim kategorisinden daha yararlıdır.",
    "Figure 8.3. Five-variant ablation summary: sigma_only is the strongest mean Sharpe-like variant, and adding categorical labels does not improve it.": "Figure 8.3. Beş-varyant ablasyon özeti: sigma_only en güçlü ortalama Sharpe-like varyantıdır ve kategorik etiket eklemek onu iyileştirmez.",
    "Figure 8.4. Oracle-label paired-seed comparison: even true regime labels do not reliably improve on sigma_hat alone.": "Figure 8.4. Oracle-etiket seed-eşleştirilmiş karşılaştırması: gerçek rejim etiketleri bile yalnızca sigma_hat'e göre güvenilir bir iyileşme sağlamaz.",
    "The compact statistical summary below collects the canonical paired tests used to interpret the ablation and robustness results.": "Aşağıdaki kompakt istatistik özeti, ablasyon ve sağlamlık sonuçlarını yorumlamak için kullanılan kanonik eşleştirilmiş testleri toplar.",
    "The t-test results show no significant Sharpe improvement from explicit regime labels, while the TOST result provides positive evidence of practical equivalence under the stated bound.": "t-testi sonuçları açık rejim etiketlerinden anlamlı bir Sharpe iyileşmesi olmadığını gösterirken, TOST sonucu belirtilen sınır altında pratik eşdeğerliğe ilişkin pozitif kanıt sağlar.",
    "8.3 Detector Robustness": "8.3 Dedektör Sağlamlığı",
    "The ANOVA statistic is retained as a descriptive robustness summary across detector-specific PPO-aware Sharpe values. Because the same seeds are shared across detector conditions, the primary inferential evidence remains the seed-paired detector-specific tests; the ANOVA should not be read as the sole formal test of a repeated-measures design.": "ANOVA istatistiği, dedektöre özgü PPO-aware Sharpe değerleri boyunca betimleyici bir sağlamlık özeti olarak tutulur. Aynı seed değerleri dedektör koşulları arasında paylaşıldığı için birincil çıkarımsal kanıt seed-eşleştirilmiş dedektöre özgü testler olarak kalır; ANOVA tekrarlı-ölçümler tasarımının tek formel testi olarak okunmamalıdır.",
    "The detector robustness table comes from Experiment 3, the Detector Robustness experiment (w5_detector_full), while the main OOS table in Section 8.1 comes from Experiment 1 (w5_main). Small differences in PPO-aware/blind means therefore reflect distinct frozen experiment batches, not an inconsistency.": "Dedektör sağlamlığı tablosu Deney 3, Dedektör Sağlamlığı deneyi (w5_detector_full), içinden gelir; Bölüm 8.1'deki ana OOS tablosu ise Deney 1 (w5_main) içinden gelir. Bu nedenle PPO-aware/blind ortalamalarındaki küçük farklar bir tutarsızlığı değil, ayrı dondurulmuş deney partilerini yansıtır.",
    "Figure 8.5. Detector robustness: rv_baseline, rv_dwell, and HMM all fail to create a reliable regime-aware PPO advantage.": "Figure 8.5. Dedektör sağlamlığı: rv_baseline, rv_dwell ve HMM'nin hiçbiri güvenilir bir rejim-farkında PPO avantajı yaratmaz.",
    "8.4 Reward-Shaping and Misspecification Checks": "8.4 Ödül Şekillendirme ve Yanlış-Belirlenim Kontrolleri",
    "The regime-conditional eta run tests whether explicit labels become useful when the reward penalizes high-volatility inventory more strongly. sigma_only still beats combined on Sharpe-like performance (p = 0.0016). Under mild regime-dependent execution misspecification, sigma_only and oracle_full remain statistically indistinguishable and practically equivalent under the reported TOST bound.": "Rejime koşullu eta çalışması, ödül yüksek-volatilite envanterini daha güçlü cezalandırdığında açık etiketlerin yararlı hâle gelip gelmediğini test eder. sigma_only, Sharpe-like performansta combined varyantını hâlâ geçer (p = 0.0016). Hafif rejime bağlı yürütme yanlış-belirlenimi altında sigma_only ve oracle_full, raporlanan TOST sınırı altında istatistiksel olarak ayırt edilemez ve pratik olarak eşdeğer kalır.",
    "In the misspecification check, the mean difference between sigma_only and oracle_full is very small, but the +/-0.05 TOST result is close to the threshold; it should therefore be read as supportive but less decisive than the main five-variant ablation equivalence result.": "Yanlış-belirlenim kontrolünde sigma_only ile oracle_full arasındaki ortalama fark çok küçüktür; ancak +/-0.05 TOST sonucu eşiğe yakındır. Bu nedenle ana beş-varyant ablasyon eşdeğerlik sonucundan daha az belirleyici, fakat destekleyici olarak okunmalıdır.",
    "Figure 8.6. Regime-conditional eta summary: changing the reward channel by regime still favors sigma_only over combined.": "Figure 8.6. Rejime koşullu eta özeti: ödül kanalını rejime göre değiştirmek hâlâ sigma_only varyantını combined karşısında destekler.",
    "Figure 8.7. Mild model-misspecification summary: sigma_only and oracle_full remain close under regime-dependent execution parameters.": "Figure 8.7. Hafif model yanlış-belirlenimi özeti: sigma_only ve oracle_full rejime bağlı yürütme parametreleri altında yakın kalır.",
    "8.5 Experiment 5: Signal-Informativeness Sweep": "8.5 Deney 5: Sinyal Bilgi-Değeri Taraması",
    "Entries are mean +/- 95% confidence-interval half-width across 20 seeds, not standard deviations.": "Girdiler 20 seed boyunca ortalama +/- 95% güven aralığı yarı genişliğidir; standart sapma değildir.",
    "Experiment 5, the Signal-Informativeness Sweep, did not support the original informativeness-threshold hypothesis within the tested calibration band. sigma_only remained high under informative conditions, while combined was directionally below sigma_only.": "Deney 5, sinyalin bilgi değerini test eden tarama, test edilen kalibrasyon bandı içinde özgün bilgi-değeri eşiği hipotezini desteklemedi. sigma_only bilgilendirici koşullar altında yüksek kaldı; combined ise yönsel olarak sigma_only altında kaldı.",
    "The Signal-Informativeness Sweep figures should be read as a refinement of the main finding rather than a new mechanism proof: they show that the tested sigma_hat degradation path did not reveal a regime-label advantage.": "Sinyal bilgi-değeri taramasına ait şekiller, yeni bir mekanizma kanıtı değil, ana bulgunun daha ince bir okuması olarak değerlendirilmelidir: test edilen sigma_hat bozulması patikasının bir rejim-etiketi avantajı ortaya çıkarmadığını gösterirler.",
    "Figure 8.8. Experiment 5 monotonic-gap plot: the expected narrowing of the sigma_only versus combined gap does not appear in the tested calibration band.": "Figure 8.8. Deney 5 monotonik-fark grafiği: sigma_only ile combined arasındaki farkın beklenen daralması test edilen kalibrasyon bandında görünmez.",
    "One possible interpretation is that the degraded continuous signal still preserves enough ordinal volatility information for quote-width adaptation, whereas the categorical label is coarser and does not add useful variation in this calibration. In noisy or lagged settings, adding the label may also increase the observation dimension without improving the control-relevant signal. This should be read as a performance pattern in the tested calibration band, not as proof of the PPO policy's internal representation mechanism.": "Olası bir yorum şudur: bozulmuş sürekli sinyal, kotasyon genişliği uyarlaması için yeterli sıralı volatilite bilgisini hâlâ korurken kategorik etiket daha kaba kalmakta ve bu kalibrasyonda yararlı bir değişkenlik eklememektedir. Gürültülü ya da gecikmeli sinyal senaryolarında etiket eklemek, kontrolle ilgili sinyali iyileştirmeden gözlem boyutunu da artırabilir. Bu, PPO politikasının iç temsil mekanizmasının kanıtı olarak değil, test edilen kalibrasyon bandında gözlenen performans deseni olarak okunmalıdır.",
    "The strongest degradation pattern appears in the lagged condition, where combined is materially below sigma_only (mean difference about -0.101; Cohen's dz about -0.91). This suggests that a stale categorical channel can be harmful when it adds delayed/coarse state information rather than control-relevant volatility information.": "En belirgin performans düşüşü gecikmeli sinyal senaryosunda görünür; burada combined, sigma_only değerinin belirgin biçimde altındadır (ortalama fark yaklaşık -0.101; Cohen's dz yaklaşık -0.91). Bu, güncel olmayan kategorik rejim sinyalinin kontrolle ilgili volatilite bilgisi yerine gecikmiş ve kaba durum bilgisi eklediğinde zararlı olabileceğini düşündürür.",
    "The paired-seed views below show whether aggregate differences are broad across seeds or driven by a few runs.": "Aşağıdaki seed-eşleştirilmiş görünümler, toplu farkların seed'ler boyunca yaygın mı yoksa birkaç çalışma tarafından mı sürüklendiğini gösterir.",
    "Figure 8.9. Experiment 5 paired-seed combined versus sigma_only: combined is directionally below sigma_only in informative conditions.": "Figure 8.9. Deney 5 seed-eşleştirilmiş combined ve sigma_only karşılaştırması: combined bilgilendirici koşullarda yönsel olarak sigma_only altındadır.",
    "Figure 8.10. Experiment 5 paired-seed combined versus regime_only: this diagnostic bounds how much sigma_hat is used inside the combined variant.": "Figure 8.10. Deney 5 seed-eşleştirilmiş combined ve regime_only karşılaştırması: bu tanı, combined varyantı içinde sigma_hat kullanımının ne kadar olduğunu sınırlar.",
    "9 Evaluation Metrics": "9 Değerlendirme Metrikleri",
    "9.1 Final Equity": "9.1 Final Equity",
    "Final equity is the terminal marked-to-market wealth. It is economically meaningful but not sufficient alone because high raw equity can be earned by taking large inventory risk.": "Final equity, terminal piyasa fiyatına göre değerlenen servettir. Ekonomik olarak anlamlıdır; ancak tek başına yeterli değildir, çünkü yüksek ham özsermaye büyük envanter riski alınarak kazanılabilir.",
    "9.2 Sharpe-Like Ratio": "9.2 Sharpe-Like Oranı",
    "Let Delta W_t = W_t - W_{t-1}. The reported Sharpe-like metric is computed as mean(Delta W_t) / std(Delta W_t), using the sample standard deviation with ddof = 1, multiplied by sqrt(1 / dt) when the standard deviation is positive; otherwise it is reported as 0. It is therefore a simulator-scale risk-adjusted PnL metric rather than an annualized market Sharpe ratio.": "Delta W_t = W_t - W_{t-1} olsun. Raporlanan Sharpe-like metriği, örnek standart sapması ddof = 1 kullanılarak mean(Delta W_t) / std(Delta W_t) biçiminde hesaplanır ve standart sapma pozitif olduğunda sqrt(1 / dt) ile çarpılır; aksi hâlde 0 olarak raporlanır. Bu nedenle yıllıklaştırılmış piyasa Sharpe oranı değil, simülatör ölçekli risk-ayarlı PnL metriğidir.",
    "9.3 Inventory Tail Risk": "9.3 Envanter Kuyruk Riski",
    "Here q_t is inventory. inv_p99 is a practical tail-risk diagnostic for whether a strategy earns returns by carrying large inventory.": "Burada q_t envanterdir. inv_p99, bir stratejinin büyük envanter taşıyarak getiri kazanıp kazanmadığını anlamak için pratik bir kuyruk-riski tanısıdır.",
    "9.4 Fill Rate": "9.4 Dolum Oranı",
    "The numerator and denominator follow the logged project convention. Fill rate is diagnostic rather than the primary thesis metric.": "Pay ve payda loglanan proje konvansiyonunu izler. Fill rate birincil tez metriği değil, tanısal bir metriktir.",
    "9.5 Paired t-Tests": "9.5 Eşleştirilmiş t-Testleri",
    "Here d_i is the paired seed-level difference between two strategies, sd(d_i) is its sample standard deviation, and n is the number of paired seeds.": "Burada d_i iki strateji arasındaki seed-düzeyinde eşleştirilmiş farktır; sd(d_i) bunun örnek standart sapmasıdır; n eşleştirilmiş seed sayısıdır.",
    "Because the manuscript reports several paired comparisons, isolated marginal p-values should be read descriptively rather than as a stand-alone familywise discovery claim. The main conclusion is based on the repeated pattern across the main OOS evaluation, detector robustness, oracle ablations, reward-shaping, misspecification, and signal-informativeness diagnostics.": "Manuskript birden fazla eşleştirilmiş karşılaştırma raporladığı için, tekil marjinal p-değerleri bağımsız bir aile-düzeyi keşif iddiası olarak değil, betimleyici olarak okunmalıdır. Ana sonuç; ana OOS değerlendirmesi, dedektör seçimine karşı sağlamlık, oracle ablasyonları, ödül şekillendirme, yanlış-belirlenim ve sinyal bilgi-değeri tanıları boyunca tekrarlanan desene dayanır.",
    "9.6 TOST Equivalence Tests": "9.6 TOST Eşdeğerlik Testleri",
    "TOST evaluates whether the plausible range of a paired difference lies inside a pre-specified practical-equivalence band. At alpha = 0.05, the corresponding equivalence reading uses the 90% confidence interval. The Five-Variant Signal Ablation uses TOST to support practical equivalence for sigma_only versus oracle_full; the Signal-Informativeness Sweep also uses related paired tests to examine non-equivalence and mean indistinguishability patterns.": "TOST, eşleştirilmiş bir farkın makul aralığının önceden belirlenmiş pratik-eşdeğerlik bandı içinde kalıp kalmadığını değerlendirir. alpha = 0.05 düzeyinde ilgili eşdeğerlik okuması 90% güven aralığını kullanır. Beş-Varyant Sinyal Ablasyonu, sigma_only ve oracle_full için pratik eşdeğerliği desteklemek amacıyla TOST kullanır; Sinyal Bilgi-Değeri Taraması da eşdeğer olmama ve ortalama ayırt edilemezliği desenlerini incelemek için ilgili eşleştirilmiş testleri kullanır.",
    "The equivalence bounds should be read as practical smallest-effect-size thresholds for the Sharpe-like metric, not as universal constants. A wider +/-0.10 band is used for the main five-variant ablation, where the question is whether oracle regime information creates a practically meaningful improvement over sigma_hat alone. Tighter +/-0.05 bounds are used in narrower robustness or misspecification checks. The thesis therefore does not treat TOST as a stand-alone proof; it reports TOST together with paired tests, confidence intervals, and directional seed-level evidence.": "Eşdeğerlik sınırları, evrensel sabitler olarak değil, Sharpe-like metriği için pratik en küçük etki büyüklüğü eşikleri olarak okunmalıdır. Daha geniş +/-0.10 bandı, sorunun oracle rejim bilgisinin yalnızca sigma_hat'e göre pratik olarak anlamlı bir iyileşme yaratıp yaratmadığı olduğu ana beş-varyant ablasyonda kullanılır. Daha dar +/-0.05 sınırları daha dar sağlamlık ya da yanlış-belirlenim kontrollerinde kullanılır. Bu nedenle tez TOST'u tek başına kanıt olarak ele almaz; TOST'u eşleştirilmiş testler, güven aralıkları ve yönsel seed-düzeyi kanıtla birlikte raporlar.",
    "10 Discussion": "10 Tartışma",
    "10.1 Main Interpretation: Signal Redundancy": "10.1 Ana Yorum: Sinyal Fazlalığı",
    "The interpretation is that the continuous realized-volatility proxy already supplies the economically useful volatility information needed by the policy in this environment. The categorical label may be coarser or redundant when presented alongside sigma_hat.": "Yorum şudur: sürekli gerçekleşmiş-volatilite vekili, bu ortamda politikanın ihtiyaç duyduğu ekonomik olarak yararlı volatilite bilgisini zaten sağlar. Kategorik etiket sigma_hat yanında sunulduğunda daha kaba ya da fazla olabilir.",
    "The evidence should therefore be read slightly more precisely than pure redundancy. If a categorical label only repeated the useful information in sigma_hat, combined and sigma_only would be expected to be approximately indistinguishable. In several experiments, however, combined is directionally below sigma_only. This pattern is consistent with signal redundancy plus mild categorical-channel degradation: the label may be coarser than the continuous volatility estimate, may add observation dimensionality, or may interact with policy learning without adding control-relevant information. This is an interpretation of the observed performance pattern, not a proof of the PPO policy's internal mechanism.": "Bu nedenle kanıt, saf fazlalıktan biraz daha hassas okunmalıdır. Kategorik bir etiket sigma_hat içindeki yararlı bilgiyi yalnızca tekrarlasaydı, combined ve sigma_only varyantlarının yaklaşık olarak ayırt edilemez olması beklenirdi. Ancak birkaç deneyde combined yönsel olarak sigma_only altındadır. Bu desen, sinyal fazlalığına ek olarak kategorik rejim etiketinin bazı senaryolarda performansı hafifçe bozması (mild categorical-channel degradation) ile tutarlıdır: etiket sürekli volatilite tahmininden daha kaba olabilir, gözlem boyutunu artırabilir ya da kontrolle ilgili bilgi eklemeden politika öğrenimiyle olumsuz etkileşebilir. Bu, gözlenen performans deseninin yorumudur; PPO politikasının iç mekanizmasının kanıtı değildir.",
    "10.2 Why the Result Is Not a Weak Null": "10.2 Sonuç Neden Zayıf Bir Null Değil",
    "The conclusion does not rest only on p > 0.05. It is supported by detector robustness, oracle labels, TOST equivalence, reward-channel checks, mild misspecification, and signal-degradation tests.": "Sonuç yalnızca p > 0.05 üzerine dayanmaz. Dedektör seçimine karşı sağlamlık, oracle etiketler, TOST eşdeğerliği, ödül-kanalı kontrolleri, hafif yanlış-belirlenim ve sinyal bozulması testleri tarafından desteklenir.",
    "10.3 Synthetic-Market Boundary": "10.3 Sentetik-Piyasa Sınırı",
    "The result is bounded to the tested synthetic HFMM simulator. The simulator deliberately isolates quote distance, Poisson-arrival fills, fees, latency, inventory, and volatility regimes so that the incremental information value of the regime channel can be tested cleanly.": "Sonuç, test edilen sentetik HFMM simülatörüyle sınırlıdır. Simülatör, rejim kanalının artı bilgi değerinin temiz biçimde test edilebilmesi için kotasyon mesafesini, Poisson-varışlı dolumları, ücretleri, gecikmeyi, envanteri ve volatilite rejimlerini bilinçli olarak izole eder.",
    "Real limit-order books include additional mechanisms that are outside the main simulator: queue position, execution priority, adverse selection, order-book imbalance, richer latency effects, clustered or self-exciting order arrivals, non-stationary liquidity, and exchange-specific transaction-cost and microstructure rules. These mechanisms can change both the value of volatility information and the way a categorical regime label interacts with execution risk.": "Gerçek limit emir defterleri ana simülatör dışında kalan ek mekanizmalar içerir: kuyruk pozisyonu, yürütme önceliği, ters seçim, emir defteri dengesizliği, daha zengin gecikme etkileri, kümelenmiş ya da kendini-uyaran emir varışları, durağan olmayan likidite ve borsaya özgü işlem maliyeti ile mikro yapı kuralları. Bu mekanizmalar hem volatilite bilgisinin değerini hem de kategorik rejim etiketinin yürütme riskiyle etkileşme biçimini değiştirebilir.",
    "For that reason, the thesis should be read as controlled synthetic evidence, not as a real-market external validity claim. Richer LOB simulators, Hawkes-process LOB models, queue-reactive execution, and real-data evaluation are natural future work and could change the incremental value of regime labels.": "Bu nedenle tez, gerçek-piyasa dış geçerlilik iddiası olarak değil, kontrollü sentetik kanıt olarak okunmalıdır. Daha zengin LOB simülatörleri, Hawkes-süreci LOB modelleri, kuyruk-tepkisel yürütme ve gerçek-veri değerlendirmesi doğal gelecek çalışmalardır ve rejim etiketlerinin ek faydasını değiştirebilir.",
    "The out-of-sample protocol should also be interpreted carefully. It is a chronological 70/30 temporal hold-out within the same synthetic data-generating process, not a distributional-shift test across fundamentally different markets. Multiple seeds improve robustness to simulation randomness, but they do not establish external validity beyond the controlled synthetic environment.": "Örneklem dışı protokol de dikkatle yorumlanmalıdır. Bu, aynı sentetik veri-üretim süreci içinde kronolojik 70/30 zamansal hold-out'tur; temelden farklı piyasalar arasında dağılımsal-kayma testi değildir. Birden fazla seed simülasyon rastgeleliğine karşı sağlamlığı artırır; ancak kontrollü sentetik ortamın ötesinde dış geçerlilik kurmaz.",
    "10.4 Risks to Interpretation": "10.4 Yoruma İlişkin Riskler",
    "10.5 Future Work": "10.5 Gelecek Çalışmalar",
    "Evaluate stronger regime-dependent execution misspecification and richer non-stationary liquidity settings.": "Daha güçlü rejime bağlı yürütme yanlış-belirlenimini ve daha zengin durağan olmayan likidite ayarlarını değerlendir.",
    "Test richer limit-order-book simulators with queue position, order-book imbalance, latency priority, adverse selection, and self-exciting order flow.": "Kuyruk pozisyonu, emir defteri dengesizliği, gecikme önceliği, ters seçim ve kendini-uyaran emir akışı içeren daha zengin limit emir defteri simülatörlerini test et.",
    "Use Hawkes-process LOB models and real-data evaluation after thesis scope, not as current claims of this report.": "Hawkes-süreci LOB modellerini ve gerçek-veri değerlendirmesini tez kapsamı sonrasında kullan; bunlar bu raporun mevcut iddiaları değildir.",
    "Run representation-level diagnostics only as future mechanism work.": "Temsil-düzeyi tanıları yalnızca gelecek mekanizma çalışması olarak yürüt.",
    "Explore alternative observation encodings and policy architectures.": "Alternatif gözlem kodlamalarını ve politika mimarilerini keşfet.",
    "11 Reproducibility Checklist": "11 Yeniden Üretilebilirlik Kontrol Listesi",
    "These commands document reproducibility paths only; thesis_34 adaptation did not rerun experiments.": "Bu komutlar yalnızca yeniden üretilebilirlik yollarını belgeler; thesis_34 uyarlaması deneyleri yeniden çalıştırmamıştır.",
    "12 Conclusion": "12 Sonuç",
    "This template-adapted report preserves the thesis finding while reorganizing the manuscript into the requested technical-report structure. PPO learns strong risk-adjusted quoting behavior in the synthetic HFMM simulator, but the explicit categorical regime channel does not provide robust incremental value beyond sigma_hat. The result is best defended as evidence consistent with signal redundancy in the tested controlled environment.": "Bu şablona uyarlanmış rapor, manuskripti istenen teknik rapor yapısına yeniden düzenlerken tez bulgusunu korur. PPO sentetik HFMM simülatöründe güçlü risk-ayarlı kotasyon davranışı öğrenir; ancak açık kategorik rejim kanalı sigma_hat ötesinde güvenilir bir ek performans katkısı sağlamaz. Sonuç, test edilen kontrollü ortamda sinyal fazlalığı ile tutarlı kanıt olarak en iyi şekilde savunulur.",
    "References": "Kaynakça",
    "Appendix A Code Appendix": "Ek A Kod Eki",
    "Canonical commands": "Kanonik komutlar",
    "Appendix B Sanity Checks and Unit Tests": "Ek B Sağduyu Kontrolleri ve Birim Testleri",
    "Appendix C Extended Evidence Tables": "Ek C Genişletilmiş Kanıt Tabloları",
    "This appendix contains diagnostic material demoted from the main body to keep the report template focused.": "Bu ek, rapor şablonunu odaklı tutmak için ana metinden aşağı taşınan tanısal materyali içerir.",
    "Appendix Figure C1. Regime-wise action distribution diagnostic.": "Ek Şekil C1. Rejim bazında eylem dağılımı tanısı.",
    "Appendix Figure C2. Regime-wise Sharpe diagnostic.": "Ek Şekil C2. Rejim bazında Sharpe tanısı.",
    "The table below summarizes supporting diagnostics. These items support interpretation and provenance; they do not replace the canonical frozen evidence chain used for the main claims.": "Aşağıdaki tablo destekleyici tanıları özetler. Bu maddeler yorumu ve köken bilgisini destekler; ana iddialar için kullanılan kanonik dondurulmuş kanıt zincirinin yerini almaz.",
    "This separation keeps the appendix readable while preserving the boundary between primary experiment evidence and supporting diagnostics.": "Bu ayrım, birincil deney kanıtı ile destekleyici tanılar arasındaki sınırı korurken ekin okunabilir kalmasını sağlar.",
}


TRANSLATIONS.update(
    {
        "HFMM-RL MSc Thesis Draft": "HFMM-RL MSc Tez Taslağı",
        "Field": "Alan",
        "Value": "Değer",
        "Author": "Yazar",
        "Programme": "Program",
        "Financial Engineering MSc Programme": "Financial Engineering MSc Programme",
        "Institution": "Kurum",
        "Supervisor": "Danışman",
        "[Supervisor to be inserted]": "[Danışman eklenecek]",
        "Matriculation No.": "Öğrenci No.",
        "[Matriculation number to be inserted]": "[Öğrenci numarası eklenecek]",
        "Submission / Draft date": "Teslim / Taslak tarihi",
        "May 2026": "Mayıs 2026",
        "Git Repository": "Git Deposu",
        "[Repository URL to be inserted]": "[Depo URL'si eklenecek]",
        "Current draft": "Mevcut taslak",
        "Previous source draft": "Önceki kaynak taslak",
        "Frozen baseline": "Dondurulmuş temel sürüm",
        "Decision log": "Karar günlüğü",
        "Symbol": "Sembol",
        "Meaning": "Anlam",
        "Mid-price at time step t.": "t zaman adımındaki orta fiyat.",
        "Inventory after step t.": "t adımından sonraki envanter.",
        "Cash account after step t.": "t adımından sonraki nakit hesabı.",
        "Marked-to-market wealth or equity, W_t = X_t + q_t S_t.": "Piyasa fiyatına göre değerlenen servet ya da özsermaye, W_t = X_t + q_t S_t.",
        "Rolling realized-volatility estimate observed by the PPO policy.": "PPO politikasının gözlemlediği kayan gerçekleşmiş-volatilite tahmini.",
        "Reward at time t.": "t zamanındaki ödül.",
        "Quoted half-spread in ticks.": "Tick cinsinden kote edilen yarı-spread.",
        "Quote skew in ticks.": "Tick cinsinden kotasyon skew değeri.",
        "Poisson fill-intensity scale and decay parameters.": "Poisson dolum-yoğunluğu ölçek ve azalma parametreleri.",
        "Inventory penalty coefficient.": "Envanter cezası katsayısı.",
        "Avellaneda-Stoikov inventory-risk aversion parameter.": "Avellaneda-Stoikov envanter-riskinden kaçınma parametresi.",
        "Remaining horizon.": "Kalan ufuk.",
        "Poisson fill intensity at quote distance delta.": "delta kotasyon mesafesindeki Poisson dolum yoğunluğu.",
        "Quote distance variables.": "Kotasyon mesafesi değişkenleri.",
        "PPO probability ratio.": "PPO olasılık oranı.",
        "Advantage estimate.": "Avantaj tahmini.",
        "Standard-normal innovation in the mid-price process.": "Orta fiyat sürecindeki standart-normal yenilik.",
        "Latent volatility regime.": "Gizil volatilite rejimi.",
        "Markov transition probability.": "Markov geçiş olasılığı.",
        "Abbreviation": "Kısaltma",
        "Arithmetic Brownian motion.": "Aritmetik Brownian hareketi.",
        "Avellaneda-Stoikov baseline.": "Avellaneda-Stoikov referans modeli.",
        "High-frequency market making.": "Yüksek frekanslı piyasa yapıcılığı.",
        "Out-of-sample evaluation.": "Örneklem dışı değerlendirme.",
        "Proximal Policy Optimization.": "Proximal Policy Optimization.",
        "Realized volatility.": "Gerçekleşmiş volatilite.",
        "Two One-Sided Tests equivalence procedure.": "Two One-Sided Tests eşdeğerlik prosedürü.",
        "Work package.": "İş paketi.",
        "Term": "Terim",
        "Definition": "Tanım",
        "Regime-aware PPO": "Rejim-farkında PPO",
        "Regime-blind PPO": "Rejim-kör PPO",
        "PPO policy variant that receives a categorical regime one-hot channel.": "Kategorik rejim one-hot kanalı alan PPO politika varyantı.",
        "PPO policy variant that omits the regime one-hot but still observes sigma_hat.": "Rejim one-hot kanalını dışarıda bırakan fakat sigma_hat gözlemlemeye devam eden PPO politika varyantı.",
        "Ablation variant using the continuous volatility signal without categorical regime labels.": "Kategorik rejim etiketleri olmadan sürekli volatilite sinyalini kullanan ablasyon varyantı.",
        "Ablation variant using sigma_hat and estimated categorical regime labels.": "sigma_hat ve tahmini kategorik rejim etiketlerini kullanan ablasyon varyantı.",
        "Ablation variant using sigma_hat and the true regime label.": "sigma_hat ve gerçek rejim etiketini kullanan ablasyon varyantı.",
        "Strategy or variant": "Strateji veya varyant",
        "Observed information": "Gözlemlenen bilgi",
        "Purpose": "Amaç",
        "No learned state dependence": "Öğrenilmiş durum bağımlılığı yok",
        "Fixed-spread baseline.": "Sabit-spread referans stratejisi.",
        "Inventory, volatility, horizon": "Envanter, volatilite, ufuk",
        "Analytical inventory-risk baseline.": "Analitik envanter-riski referans modeli.",
        "sigma_hat and estimated regime one-hot": "sigma_hat ve tahmini rejim one-hot",
        "Original regime-aware PPO.": "Özgün rejim-farkında PPO.",
        "sigma_hat without regime one-hot": "rejim one-hot olmadan sigma_hat",
        "Original regime-blind PPO.": "Özgün rejim-kör PPO.",
        "Continuous sigma_hat only": "Yalnızca sürekli sigma_hat",
        "Tests whether sigma_hat alone carries the signal.": "Yalnızca sigma_hat'in sinyali taşıyıp taşımadığını test eder.",
        "sigma_hat and estimated regime one-hot": "sigma_hat ve tahmini rejim one-hot",
        "Tests estimated categorical incremental value.": "Tahmini kategorik etiketin ilave bilgi değerini test eder.",
        "sigma_hat and true regime one-hot": "sigma_hat ve gerçek rejim one-hot",
        "Tests perfect-label incremental value.": "Kusursuz etiketin ilave bilgi değerini test eder.",
        "Categorical labels without sigma_hat": "sigma_hat olmadan kategorik etiketler",
        "Anchor label-only variants.": "Yalnızca-etiket dayanak varyantları.",
        "Component": "Bileşen",
        "Reference file": "Referans dosya",
        "Run lifecycle": "Çalıştırma yaşam döngüsü",
        "Simulator": "Simülatör",
        "Regime generation and detectors": "Rejim üretimi ve dedektörler",
        "Gymnasium environment": "Gymnasium ortamı",
        "PPO training": "PPO eğitimi",
        "WP5 evaluation": "WP5 değerlendirmesi",
        "WP6 sweep": "WP6 süpürmesi",
        "Detector": "Dedektör",
        "Role": "Rol",
        "Caveat": "Not",
        "Main causal detector": "Ana nedensel dedektör",
        "Rolling RV threshold; used in main WP4/WP5/WP6 pipelines.": "Kayan RV eşiği; ana WP4/WP5/WP6 boru hatlarında kullanılır.",
        "Auxiliary robustness detector": "Yardımcı sağlamlık dedektörü",
        "Offline dwell smoothing; not the main causal detector.": "Çevrimdışı dwell yumuşatma; ana nedensel dedektör değildir.",
        "Robustness detector": "Sağlamlık dedektörü",
        "Higher classification accuracy but no reliable PPO advantage.": "Daha yüksek sınıflandırma doğruluğu, ancak güvenilir PPO avantajı yok.",
        "Hyperparameter": "Hiperparametre",
        "Canonical full-run value": "Kanonik tam-çalışma değeri",
        "Strategy": "Strateji",
        "Final equity": "Final equity",
        "Fill rate": "Fill rate",
        "Variant": "Varyant",
        "Signals": "Sinyaller",
        "sigma_hat only": "yalnızca sigma_hat",
        "sigma_hat + regime_true": "sigma_hat + regime_true",
        "regime_hat only": "yalnızca regime_hat",
        "sigma_hat + regime_hat": "sigma_hat + regime_hat",
        "regime_true only": "yalnızca regime_true",
        "Comparison": "Karşılaştırma",
        "Metric or test": "Metrik veya test",
        "Canonical result": "Kanonik sonuç",
        "Interpretation": "Yorum",
        "No significant Sharpe improvement.": "Anlamlı Sharpe iyileşmesi yok.",
        "Favors PPO-blind.": "PPO-blind lehine.",
        "No oracle-label Sharpe improvement.": "Oracle-etiket Sharpe iyileşmesi yok.",
        "Practical equivalence supported.": "Pratik eşdeğerlik destekleniyor.",
        "Higher accuracy does not create advantage.": "Daha yüksek doğruluk avantaj yaratmıyor.",
        "Detector choice does not explain result.": "Dedektör seçimi sonucu açıklamaz.",
        "Favors sigma_only.": "sigma_only lehine.",
        "No significant Sharpe difference.": "Anlamlı Sharpe farkı yok.",
        "Aware vs blind conclusion": "Aware/blind sonucu",
        "Key statistic": "Ana istatistik",
        "No reliable aware Sharpe advantage": "Güvenilir aware Sharpe avantajı yok",
        "No reliable aware Sharpe advantage despite higher accuracy": "Daha yüksek doğruluğa rağmen güvenilir aware Sharpe avantajı yok",
        "condition": "koşul",
        "Risk": "Risk",
        "Mitigation": "Azaltım",
        "Overclaiming mechanism": "Mekanizmayı fazla iddia etmek",
        "Use signal-redundancy and categorical-channel degradation as interpretations, not proofs of PPO internals.": "Sinyal fazlalığı ve kategorik kanalın hafif performans bozucu etkisini PPO iç yapısının kanıtı değil, yorum olarak kullan.",
        "Detector causality confusion": "Dedektör nedenselliği karışıklığı",
        "State that rv_baseline is the main causal detector and rv_dwell is auxiliary/offline.": "rv_baseline'ın ana nedensel dedektör, rv_dwell'in ise yardımcı/çevrimdışı olduğunu belirt.",
        "Synthetic external validity": "Sentetik dış geçerlilik",
        "Frame all claims as controlled synthetic-market claims.": "Tüm iddiaları kontrollü sentetik-piyasa iddiaları olarak çerçevele.",
        "Template over-compression": "Şablonun aşırı sıkıştırması",
        "Keep the literature and evidence spine, but map it into template headings.": "Literatür ve kanıt omurgasını koru, ancak şablon başlıklarına eşle.",
        "Mid-price mark-to-market assumption": "Orta fiyat mark-to-market varsayımı",
        "Report it as a simulator evaluation convention; real liquidation would require spread/depth/queue/adverse-selection considerations.": "Bunu bir simülatör değerlendirme konvansiyonu olarak raporla; gerçek likidasyon spread/derinlik/kuyruk/ters-seçim değerlendirmeleri gerektirir.",
        "Item": "Madde",
        "Status in thesis_34": "thesis_34 içindeki durum",
        "thesis-v29-frozen tag; commit 9681faa; thesis_29 remains untouched.": "thesis-v29-frozen etiketi; commit 9681faa; thesis_29 dokunulmadan kalır.",
        "thesis_34 is a defense-clarity polish draft built on thesis_33; frozen evidence remains unchanged.": "thesis_34, thesis_33 üzerine kurulu bir savunma-açıklığı cilalama taslağıdır; dondurulmuş kanıt değişmeden kalır.",
        "Run ID structure": "Run ID yapısı",
        "Config snapshots": "Yapılandırma anlık görüntüleri",
        "Every run snapshots config; config_snapshot_all.md inventories active config files.": "Her çalışma yapılandırmanın anlık görüntüsünü alır; config_snapshot_all.md aktif yapılandırma dosyalarının envanterini tutar.",
        "Seeds": "Seed'ler",
        "WP5 main uses seeds 1-20; WP6 full uses seeds 42-61.": "WP5 main seed 1-20 kullanır; WP6 full seed 42-61 kullanır.",
        "Train/test split": "Eğitim/test ayrımı",
        "70/30 chronological split on exogenous synthetic series.": "Dışsal sentetik seri üzerinde kronolojik 70/30 ayrım.",
        "Evidence manifest": "Kanıt manifestosu",
        "EVIDENCE_MANIFEST.md records protected evidence and remediation invariants.": "EVIDENCE_MANIFEST.md korunan kanıtı ve düzeltme değişmez koşullarını (invariants) kaydeder.",
        "Protected CSV hashes": "Korunan CSV hash'leri",
        "Listed below and verified by file hash checks.": "Aşağıda listelenmiş ve dosya hash kontrolleriyle doğrulanmıştır.",
        "Codebase snapshot": "Kod tabanı anlık görüntüsü",
        "No-rerun policy": "Yeniden çalıştırmama politikası",
        "No PPO/WP5/WP6 reruns and no frozen evidence-figure regeneration; only the methodology-only synthetic-environment illustration was generated.": "PPO/WP5/WP6 yeniden çalıştırması yok ve dondurulmuş kanıt-şekli yeniden üretimi yok; yalnızca metodolojiye yönelik sentetik-ortam gösterimi üretildi.",
        "Protected artifact": "Korunan artefakt",
        "Expected SHA256": "Beklenen SHA256",
        "Check": "Kontrol",
        "Reproducibility command": "Yeniden üretilebilirlik komutu",
        "Create environment.": "Ortam oluştur.",
        "Activate environment on Windows PowerShell.": "Windows PowerShell üzerinde ortamı etkinleştir.",
        "Install dependencies.": "Bağımlılıkları kur.",
        "Run main WP5 evaluation.": "Ana WP5 değerlendirmesini çalıştır.",
        "Run WP6 sweep.": "WP6 süpürmesini çalıştır.",
        "Resume WP6 sweep.": "WP6 süpürmesini devam ettir.",
        "Lint source code.": "Kaynak kodu lint et.",
        "Area": "Alan",
        "Files": "Dosyalar",
        "Run management": "Çalıştırma yönetimi",
        "Simulation": "Simülasyon",
        "Regimes": "Rejimler",
        "Environment": "Ortam",
        "Evaluation": "Değerlendirme",
        "Signal audit and sweep": "Sinyal denetimi ve süpürmesi",
        "Verifies stable CSV metric schema behavior.": "Kararlı CSV metrik şeması davranışını doğrular.",
        "Verifies resume-mode config validation.": "Devam modu yapılandırma doğrulamasını doğrular.",
        "WP3 sanity checks": "WP3 sağduyu kontrolleri",
        "Exercises naive, AS, and random policies through the Gymnasium environment.": "naive, AS ve rastgele politikaları Gymnasium ortamı üzerinden çalıştırır.",
        "Figure/path checks": "Şekil/yol kontrolleri",
        "Ensures thesis_34 references existing frozen figure files.": "thesis_34'ün mevcut dondurulmuş şekil dosyalarına referans verdiğini doğrular.",
        "Forbidden phrase scan": "Yasak ifade taraması",
        "Guards against overclaims in the generated draft.": "Üretilen taslakta aşırı iddialara karşı koruma sağlar.",
        "Diagnostic": "Tanı",
        "Status": "Durum",
        "Post-hoc signal diagnostics": "Post-hoc sinyal tanıları",
        "Supporting only": "Yalnızca destekleyici",
        "Consistent with signal redundancy but not primary evidence.": "Sinyal fazlalığı ile tutarlı, ancak birincil kanıt değil.",
        "WP5.5 signal audit": "WP5.5 sinyal denetimi",
        "Offline calibration": "Çevrimdışı kalibrasyon",
        "Supports WP6 degradation design; no PPO training.": "WP6 sinyal bozulması tasarımını destekler; PPO eğitimi yok.",
        "Detector robustness statistics": "Dedektör seçimine karşı sağlamlık istatistikleri",
        "Protected evidence": "Korunan kanıt",
        "Addresses poor-detector objection without changing main detector caveat.": "Ana dedektör notunu değiştirmeden zayıf-dedektör itirazını ele alır.",
        "WP6 paired summaries": "WP6 eşleştirilmiş özetleri",
        "Supports refined interpretation of the signal-informativeness sweep.": "Sinyal bilgi-değeri taramasının daha incelikli yorumunu destekler.",
    }
)


def iter_all_paragraphs(doc: Document):
    paragraphs = list(doc.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                paragraphs.extend(cell.paragraphs)
    for section in doc.sections:
        paragraphs.extend(section.header.paragraphs)
        paragraphs.extend(section.footer.paragraphs)
    return paragraphs


def has_drawing(paragraph) -> bool:
    return any(run._r.xpath(".//w:drawing") for run in paragraph.runs)


def replace_paragraph_text(paragraph, new_text: str) -> None:
    if has_drawing(paragraph):
        return
    if paragraph.runs:
        paragraph.runs[0].text = new_text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(new_text)


def disable_auto_hyphenation(doc: Document) -> None:
    settings = doc.settings.element
    for element in settings.findall(qn("w:autoHyphenation")):
        settings.remove(element)
    auto_hyphenation = OxmlElement("w:autoHyphenation")
    auto_hyphenation.set(qn("w:val"), "false")
    settings.append(auto_hyphenation)


def export_pdf_with_word() -> None:
    tmp_pdf = PDF.with_name(PDF.stem + "_export_tmp.pdf")
    if tmp_pdf.exists():
        tmp_pdf.unlink()
    ps = (
        "$ErrorActionPreference = 'Stop'; "
        f"$tmp = '{str(tmp_pdf)}'; "
        "if (Test-Path -LiteralPath $tmp) { Remove-Item -LiteralPath $tmp -Force }; "
        "$word = New-Object -ComObject Word.Application; "
        "$word.Visible = $false; "
        "try { "
        f"$doc = $word.Documents.Open('{str(DST)}', $false, $true); "
        "$doc.ExportAsFixedFormat($tmp, 17); "
        "$doc.Close($false); "
        "} finally { $word.Quit() }; "
        f"Move-Item -LiteralPath $tmp -Destination '{str(PDF)}' -Force"
    )
    subprocess.run(["powershell", "-Command", ps], cwd=str(ROOT), check=True, capture_output=True, text=True)


def export_pdf() -> bool:
    if sys.platform.startswith("win"):
        try:
            export_pdf_with_word()
            return True
        except Exception as word_err:
            print(f"Microsoft Word PDF export failed: {word_err}")
    return False


def main() -> int:
    if not SRC.exists():
        raise FileNotFoundError(SRC)

    shutil.copy2(SRC, DST)
    doc = Document(DST)
    before_shapes = len(doc.inline_shapes)

    replaced = 0
    for paragraph in iter_all_paragraphs(doc):
        text = paragraph.text
        if text in TRANSLATIONS:
            replace_paragraph_text(paragraph, TRANSLATIONS[text])
            replaced += 1

    after_shapes = len(doc.inline_shapes)
    if after_shapes != before_shapes:
        raise RuntimeError(f"Figure preservation failed: inline_shapes {before_shapes} -> {after_shapes}")

    disable_auto_hyphenation(doc)
    doc.save(DST)
    print(f"Wrote {DST.relative_to(ROOT)}")
    print(f"Translated text blocks: {replaced}")
    print(f"Preserved inline_shapes: {after_shapes}")

    if export_pdf():
        print(f"Wrote {PDF.relative_to(ROOT)}")
    else:
        print(f"PDF export failed; DOCX remains available at {DST.relative_to(ROOT)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
