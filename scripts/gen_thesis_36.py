"""Generate thesis_36.docx/pdf from thesis_35 with low-risk format polish.

This is a documentation-only generator. It does not run PPO, WP5, WP6,
detector robustness, ablations, misspecification checks, signal-informativeness
sweeps, protected CSV generation, or protected performance-figure generation.
It copies thesis_35 and applies only final submission structure/formatting
changes while preserving figures, formulas, numerical claims, and evidence.
"""

from __future__ import annotations

import re
import shutil
import subprocess
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "manuscript" / "thesis_35.docx"
DST = ROOT / "manuscript" / "thesis_36.docx"
PDF = ROOT / "manuscript" / "thesis_36.pdf"


MAIN_SECTIONS = [
    "Abstract",
    "Symbols, Abbreviations, and Glossary",
    "1 Introduction",
    "2 Mathematical Formulation",
    "3 Use-Case Scenario",
    "4 System Architecture",
    "5 Layered Experimental Model",
    "6 Algorithm Specification and Pseudocode",
    "7 Experimental Setup",
    "8 Results and Visualisation",
    "9 Evaluation Metrics",
    "10 Discussion",
    "11 Reproducibility Checklist",
    "12 Conclusion",
    "References",
    "Appendix A Code Appendix",
    "Appendix B Unit Tests and Validation Checks",
    "Appendix C Extended Evidence Tables",
    "Appendix D Author Contribution Statement",
]

TABLE_CAPTIONS = {
    4: "Table 3.1 Strategy variants",
    5: "Table 6.1 Reference implementation map",
    6: "Table 7.1 Detector variants",
    7: "Table 7.2 PPO hyperparameters",
    8: "Table 8.1 Main OOS strategy metrics",
    9: "Table 8.2 Five-variant ablation metrics",
    10: "Table 8.3 Compact statistical summary",
    11: "Table 8.4 Detector robustness summary",
    12: "Table 8.5 Signal information-value sweep",
    13: "Table 10.1 Risks to interpretation",
    14: "Table 11.1 Reproducibility checklist",
    15: "Table 11.2 Protected artifact hash checks",
    16: "Table 11.3 Reproducibility commands",
    17: "Appendix Table A.1 Code map",
    18: "Appendix Table B.1 Unit test and validation checklist",
    19: "Appendix Table C.1 Supporting diagnostics",
}

ALGORITHMS = [
    ("6.2 Synthetic Market and Regime Generation", "Algorithm 1: Synthetic market and regime generation"),
    ("6.3 Market-Making Environment Step", "Algorithm 2: Market-making environment step"),
    ("6.4 PPO Training and OOS Evaluation Protocol", "Algorithm 3: PPO training and OOS evaluation protocol"),
    ("6.5 Experiment 5: Signal-Informativeness Sweep", "Algorithm 4: Signal information-value sweep"),
]


def iter_all_paragraphs(doc: Document):
    paragraphs = list(doc.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                paragraphs.extend(cell.paragraphs)
    for section in doc.sections:
        paragraphs.extend(section.header.paragraphs)
        paragraphs.extend(section.footer.paragraphs)
    return paragraphs


def has_drawing(paragraph) -> bool:
    return any(run._r.xpath(".//w:drawing") for run in paragraph.runs)


def safe_set_paragraph_text(paragraph, text: str) -> None:
    """Replace paragraph text only when the paragraph has no drawing runs."""
    if has_drawing(paragraph):
        return
    if paragraph.text == text:
        return
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            if not run._r.xpath(".//w:drawing"):
                run.text = ""
    else:
        paragraph.add_run(text)


def find_paragraph(doc: Document, exact_text: str):
    for paragraph in doc.paragraphs:
        if paragraph.text.strip() == exact_text:
            return paragraph
    raise RuntimeError(f"Could not find paragraph: {exact_text}")


def add_paragraph_before_element(doc: Document, element, text: str = "", style: str | None = None):
    paragraph = doc.add_paragraph(text, style=style)
    element.addprevious(paragraph._p)
    return paragraph


def insert_paragraph_after(anchor, text: str = "", style: str | None = None):
    paragraph = anchor._parent.add_paragraph(text, style=style)
    anchor._p.addnext(paragraph._p)
    return paragraph


def delete_paragraph(paragraph) -> None:
    paragraph._element.getparent().remove(paragraph._element)
    paragraph._p = paragraph._element = None


def style_caption(paragraph) -> None:
    paragraph.style = "Caption" if "Caption" in [s.name for s in paragraph.part.document.styles] else paragraph.style
    if paragraph.runs:
        paragraph.runs[0].bold = True
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(3)


def style_code(paragraph) -> None:
    for run in paragraph.runs:
        run.font.name = "Consolas"
        run.font.size = Pt(9)
    paragraph.paragraph_format.left_indent = Pt(18)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)


def add_toc_field(paragraph) -> None:
    run = paragraph.add_run()
    fld_char = OxmlElement("w:fldChar")
    fld_char.set(qn("w:fldCharType"), "begin")
    run._r.append(fld_char)

    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = 'TOC \\o "1-3" \\h \\z \\u'
    run._r.append(instr)

    fld_char = OxmlElement("w:fldChar")
    fld_char.set(qn("w:fldCharType"), "separate")
    run._r.append(fld_char)

    paragraph.add_run("Right-click and update field if page numbers are not refreshed.")

    run = paragraph.add_run()
    fld_char = OxmlElement("w:fldChar")
    fld_char.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char)


def ensure_cover_row(table, label: str, value: str) -> None:
    for row in table.rows:
        if row.cells[0].text.strip() == label:
            row.cells[1].text = value
            return
    row = table.add_row()
    row.cells[0].text = label
    row.cells[1].text = value


def remove_cover_row(table, label: str) -> None:
    for row in list(table.rows):
        if row.cells[0].text.strip() == label:
            table._tbl.remove(row._tr)
            return


def update_title_page(doc: Document) -> list[str]:
    cover = doc.tables[0]
    placeholders = []
    known = {
        "Author": "Onur Emiroglu",
        "Programme": "Financial Engineering MSc Programme",
        "Institution": "Karlsruhe Institute of Technology",
        "Course / Module": "[Course / module to be inserted]",
        "Supervisor": "[Supervisor to be inserted]",
        "Advisor / Co-supervisor": "[Advisor / co-supervisor to be inserted, if applicable]",
        "Matriculation No.": "[Matriculation number to be inserted]",
        "Submission / Draft date": "Draft: 2026-06-06",
        "Git Repository": "https://github.com/OnurEmiroglu/THESIS",
    }
    for label, value in known.items():
        ensure_cover_row(cover, label, value)
        if value.startswith("["):
            placeholders.append(label)
    for label in ("Current draft", "Previous source draft", "Frozen baseline", "Decision log"):
        remove_cover_row(cover, label)
    for paragraph in doc.paragraphs:
        if paragraph.text.startswith("Draft integrity note:"):
            safe_set_paragraph_text(
                paragraph,
                "Draft integrity note: thesis_36 is a final Student Report Format Guide compliance pass. "
                "Frozen experimental evidence is reused and not regenerated.",
            )
    return placeholders


def clean_abstract_acronyms(doc: Document) -> None:
    abstract = find_paragraph(doc, "Abstract")
    abstract_index = next(i for i, p in enumerate(doc.paragraphs) if p._p is abstract._p)
    for paragraph in doc.paragraphs[abstract_index + 1 : abstract_index + 4]:
        text = paragraph.text.replace(
            "TOST +/-0.10 p = 0.00067; 90% CI",
            "two one-sided tests (TOST) +/-0.10 p = 0.00067; 90% confidence interval (CI)",
        )
        if text != paragraph.text:
            safe_set_paragraph_text(paragraph, text)


def enforce_clean_page_starts(doc: Document) -> None:
    find_paragraph(doc, "Contents").paragraph_format.page_break_before = True
    find_paragraph(doc, "Appendix B Unit Tests and Validation Checks").paragraph_format.page_break_before = True


EQUATIONS = {
    "(1)": "(1)   Sₜ₊₁ = Sₜ + σₜ√Δt εₜ,    εₜ ∼ N(0, 1)",
    "(2)": "(2)   P(zₜ₊₁ = j | zₜ = i) = Pᵢⱼ,    zₜ ∈ {L, M, H}",
    "(2a)": "(2a)   P = [[0.9967, 0.0023, 0.0010], [0.0042, 0.9917, 0.0041], [0.0010, 0.0030, 0.9960]]",
    "(3)": "(3)   σ̂ₜ = √[(1/w) Σᵗᵢ₌ₜ₋w₊₁ (ΔSᵢ)²]",
    "(4)": "(4)   λ(δ) = A exp(−kδ)",
    "(5)": "(5)   P(fill | δ) = 1 − exp[−λ(δ)Δt]",
    "(6)": "(6)   h = hᵢₙₓ + 1,    m = mᵢₙₓ − 2",
    "(7)": "(7)   δᵦᵢ𝒹 = max(1, h + m),    δₐₛₖ = max(1, h − m)",
    "(8)": "(8)   qₜ₊₁ = qₜ + Fᵇⁱᵈₜ − Fᵃˢᵏₜ",
    "(9)": "(9)   Xₜ₊₁ = Xₜ − FᵇⁱᵈₜPᵇⁱᵈₜ + FᵃˢᵏₜPᵃˢᵏₜ − feesₜ",
    "(10)": "(10)   Wₜ = Xₜ + qₜSₜ",
    "(11)": "(11)   Rₜ = Wₜ₊₁ − Wₜ − ηq²ₜ₊₁",
    "(12)": "(12)   rₜ = Sₜ − qₜγσ²τ",
    "(13)": "(13)   δAS = 0.5γσ²τ + (1/γ) ln(1 + γ/k)",
    "(14)": "(14)   ρₜ(θ) = πθ(aₜ | sₜ) / πθold(aₜ | sₜ)",
    "(15)": "(15)   LCLIP(θ) = Eₜ[min(ρₜ(θ)Âₜ, clip(ρₜ(θ), 1−ε, 1+ε)Âₜ)]",
}


def format_equations(doc: Document) -> None:
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        number = next((key for key in EQUATIONS if text.startswith(key + " ")), None)
        if number is None:
            continue
        safe_set_paragraph_text(paragraph, EQUATIONS[number])
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_before = Pt(6)
        paragraph.paragraph_format.space_after = Pt(4)
        paragraph.paragraph_format.keep_with_next = True
        for run in paragraph.runs:
            run.font.name = "Cambria Math"
            run.font.size = Pt(11)


def insert_table_of_contents(doc: Document) -> None:
    abstract = find_paragraph(doc, "Abstract")
    paragraphs = list(doc.paragraphs)
    abstract_idx = next(i for i, paragraph in enumerate(paragraphs) if paragraph._p is abstract._p)
    if abstract_idx > 0:
        previous = paragraphs[abstract_idx - 1]
        if not previous.text.strip() and previous._p.xpath(".//w:br"):
            delete_paragraph(previous)
    heading = add_paragraph_before_element(doc, abstract._p, "Contents")
    if heading.runs:
        heading.runs[0].bold = True
        heading.runs[0].font.size = Pt(16)
    toc = add_paragraph_before_element(doc, abstract._p)
    add_toc_field(toc)
    spacer = add_paragraph_before_element(doc, abstract._p)
    spacer.add_run().add_break(WD_BREAK.PAGE)
    heading.paragraph_format.keep_with_next = True


def add_table_captions(doc: Document) -> None:
    for index, caption in sorted(TABLE_CAPTIONS.items(), reverse=True):
        paragraph = add_paragraph_before_element(doc, doc.tables[index]._tbl, caption)
        style_caption(paragraph)


def add_algorithm_blocks(doc: Document) -> None:
    for heading_text, caption_text in reversed(ALGORITHMS):
        heading = find_paragraph(doc, heading_text)
        caption = insert_paragraph_after(heading, caption_text)
        style_caption(caption)
    in_algorithm = False
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text.startswith("Algorithm "):
            in_algorithm = True
            continue
        if re.match(r"^6\.\d", text):
            in_algorithm = False
        if in_algorithm and text and not re.match(r"^6\.\d", text):
            style_code(paragraph)


def add_hyperparameter_protocol(doc: Document) -> None:
    heading = find_paragraph(doc, "7.6 PPO Hyperparameters")
    paragraphs = [
        "7.6.1 Hyperparameter Selection Protocol",
        "The thesis does not perform a broad hyperparameter optimization campaign. This is intentional because the central comparison is signal-channel design, not PPO tuning.",
        "PPO hyperparameters are held fixed across PPO variants to avoid giving one information condition a tuning advantage.",
        "eta = 0.001 is the canonical inventory-penalty scale selected after the earlier inventory-penalty ablation.",
        "No separate validation set is introduced because the study uses a fixed train/OOS evaluation protocol; this should be treated as a limitation and a future extension if full tuning is pursued.",
        "This is a methodological choice, not a claim that the PPO settings are globally optimal.",
    ]
    anchor = heading
    for text in paragraphs:
        style = "Heading 2" if text.startswith("7.6.1") else None
        anchor = insert_paragraph_after(anchor, text, style=style)


def remove_duplicate_hyperparameter_paragraph(doc: Document) -> None:
    old_prefix = "The PPO hyperparameters are held fixed across PPO variants to keep the signal-channel comparison fair."
    for paragraph in list(doc.paragraphs):
        if paragraph.text.strip().startswith(old_prefix):
            delete_paragraph(paragraph)
            return


def strengthen_reproducibility(doc: Document) -> None:
    heading = find_paragraph(doc, "11 Reproducibility Checklist")
    paragraphs = [
        "The reproducibility path uses Python 3 with a local virtual environment created by python -m venv .venv and dependencies installed from requirements.txt. The exact interpreter version should be recorded at submission if institutionally required.",
        "Main commands are documented in Table 11.3. Exact reproducibility depends on using the project repository, config snapshots, fixed seeds, logged run metadata, and the protected evidence files recorded in EVIDENCE_MANIFEST.md.",
        "Protected evidence is checked by SHA256. No Dockerfile or environment.yml is present, so this report does not claim Docker or Conda reproducibility.",
        "No external CI service is claimed; validation is documented through local reproducibility checks and protected evidence hashes.",
        "Formal line-coverage percentage is not claimed; the tests are targeted regression and reproducibility checks.",
        "Repository and manuscript provenance: https://github.com/OnurEmiroglu/THESIS; current draft manuscript/thesis_36.pdf and .docx; source draft manuscript/thesis_35.pdf and .docx; frozen baseline manuscript/thesis_29.pdf at tag thesis-v29-frozen; decision log manuscript/decisions_log_13.pdf.",
        "Repository license status: [License information to be inserted]. The repository is thesis research code and evidence; no broader licensing claim is made here.",
    ]
    anchor = heading
    for text in paragraphs:
        anchor = insert_paragraph_after(anchor, text)


def expand_appendix_b(doc: Document) -> None:
    find_paragraph(doc, "Appendix B Unit Tests and Validation Checks")
    table = doc.tables[18]
    while len(table.columns) < 4:
        table.add_column(Pt(90))
    entries = [
        ("tests/test_csv_metric_logger.py", "Reject CSV metric schema drift and accept consistent rows.", "pytest tests/test_csv_metric_logger.py -q; all 3 tests pass.", "Automated pytest"),
        ("tests/test_resume_validation.py", "Validate matching, mismatched, forced, and legacy resume configurations.", "pytest tests/test_resume_validation.py -q; all 5 tests pass.", "Automated pytest"),
        ("Source lint", "Check source files for configured Ruff violations.", "ruff check src/ exits successfully.", "Automated local check"),
        ("WP3 sanity configuration", "Exercise naive, AS, and random policies through MMEnv.", "Documented sanity run completes and emits expected artifacts.", "Documented validation; not rerun for thesis_36"),
        ("Protected SHA256 audit", "Detect changes to the four canonical protected CSV evidence files.", "All four hashes match EVIDENCE_MANIFEST.md.", "Automated local hash check"),
        ("Document artifact audit", "Check captions, algorithms, placeholders, figures, and stale wording.", "Requested document-generation checks pass; remaining placeholders are listed.", "Automated/documented validation"),
    ]
    while len(table.rows) < len(entries) + 1:
        table.add_row()
    headers = ("Test or validation category", "Purpose", "Expected pass criterion", "Status/type")
    for col, value in enumerate(headers):
        table.cell(0, col).text = value
    for row_index, entry in enumerate(entries, start=1):
        for col, value in enumerate(entry):
            table.cell(row_index, col).text = value
    while len(table.rows) > len(entries) + 1:
        table._tbl.remove(table.rows[-1]._tr)
    appendix_c = find_paragraph(doc, "Appendix C Extended Evidence Tables")
    paragraphs = [
        "Available automated tests are the two files listed in Appendix Table B.1. Run the complete suite with pytest -q, or run either file individually with pytest <path> -q.",
        "The WP3 sanity configuration is an experiment-like validation path and is documented here for reproducibility; it was not rerun for this formatting-only thesis_36 generation.",
        "Formal line-coverage percentage is not claimed; the tests are targeted regression and reproducibility checks.",
        "No GitHub Actions workflow or other external continuous-integration (CI) service is claimed. Checks are run locally or recorded as documented validation.",
    ]
    for text in reversed(paragraphs):
        add_paragraph_before_element(doc, appendix_c._p, text)


def add_appendix_d(doc: Document) -> None:
    heading = doc.add_paragraph("Appendix D Author Contribution Statement", style="Heading 1")
    heading.paragraph_format.page_break_before = True
    doc.add_paragraph(
        "This thesis is submitted as a single-author MSc thesis by Onur Emiroglu. The reported implementation, "
        "experiments, documentation, and interpretation are organized in the project repository and were prepared "
        "under normal thesis supervision. Supervisory feedback is reflected in the final framing and limitations, "
        "but no multi-author contribution matrix is required."
    )


def update_thesis36_labels(doc: Document) -> None:
    replacements = {
        "Status in thesis_35": "Status in thesis_36",
        "thesis_35 is a final format-guide compliance polish draft built on thesis_34; frozen evidence remains unchanged.": (
            "thesis_36 is a final low-risk format-compliance polish built on thesis_35; frozen evidence remains unchanged."
        ),
        "The thesis_35 adaptation reuses frozen performance results.": (
            "The thesis_35/thesis_36 adaptation reuses frozen performance results."
        ),
        "the thesis_34 adaptation embeds existing figures": (
            "the thesis_35/thesis_36 adaptation embeds existing figures"
        ),
        "The thesis_34 adaptation embeds existing figures": (
            "The thesis_35/thesis_36 adaptation embeds existing figures"
        ),
        "the thesis_35 adaptation embeds existing figures": (
            "the thesis_35/thesis_36 adaptation embeds existing figures"
        ),
        "These commands document reproducibility paths only; thesis_35 adaptation did not rerun experiments.": (
            "These commands document reproducibility paths only; thesis_35/thesis_36 adaptation did not rerun experiments."
        ),
        "Ensures thesis_35 references existing frozen figure files.": (
            "Ensures thesis_36 references existing frozen figure files."
        ),
        "Appendix Figure C1.": "Appendix Figure C.1.",
        "Appendix Figure C2.": "Appendix Figure C.2.",
    }
    targets = list(doc.paragraphs)
    for table in doc.tables[1:]:
        for row in table.rows:
            for cell in row.cells:
                targets.extend(cell.paragraphs)
    for paragraph in targets:
        if has_drawing(paragraph):
            continue
        text = paragraph.text
        new_text = text
        for old, new in replacements.items():
            new_text = new_text.replace(old, new)
        if new_text != text:
            safe_set_paragraph_text(paragraph, new_text)


def round_ablation_table(doc: Document) -> None:
    table = doc.tables[9]
    for row in table.rows[1:]:
        for cell in row.cells[2:]:
            text = cell.text.strip()
            if re.fullmatch(r"-?\d+\.\d{4,}", text):
                cell.text = f"{float(text):.3f}"


def strengthen_appendix_a(doc: Document) -> None:
    caption = find_paragraph(doc, "Appendix Table A.1 Code map")
    paragraphs = [
        "Repository structure: run.py dispatches JSON jobs; src/ contains WP1-WP6 implementation modules; config/ contains experiment configurations; tests/ contains the available regression tests; results/ and docs/internal/ contain run outputs and evidence summaries; manuscript/ contains report artifacts.",
        "Environment setup: python -m venv .venv; activate the environment; pip install -r requirements.txt. No Dockerfile, Conda environment.yml, or formal environment lockfile is present.",
        "Configuration coverage includes baseline, synthetic-regime, sanity, PPO, OOS evaluation, ablation, detector, misspecification, WP5.5 audit/calibration, and WP6 sweep JSON files under config/.",
        "Reproducibility controls include fixed seeds, a chronological 70/30 train/test split, per-run config snapshots and metadata, stable CSV schemas, resume validation, and the four protected SHA256 entries in EVIDENCE_MANIFEST.md.",
        "Frozen-evidence policy: the commands below document provenance only. Do not rerun PPO/WP5/WP6 or regenerate protected CSVs or performance figures for documentation-only changes.",
        "Repository: https://github.com/OnurEmiroglu/THESIS. License: [License information to be inserted].",
    ]
    for text in reversed(paragraphs):
        add_paragraph_before_element(doc, caption._p, text)


def standardize_table_fonts(doc: Document) -> None:
    for table_index in (7, 9, 18):
        for row in doc.tables[table_index].rows:
            for cell in row.cells:
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_before = Pt(0)
                    paragraph.paragraph_format.space_after = Pt(0)
                    for run in paragraph.runs:
                        run.font.name = "Arial"
                        run.font.size = Pt(8.5)


def update_pdf_export_fields_with_word() -> None:
    tmp_pdf = PDF.with_name(PDF.stem + "_export_tmp.pdf")
    if tmp_pdf.exists():
        tmp_pdf.unlink()
    ps = (
        "$ErrorActionPreference = 'Stop'; "
        f"$tmp = '{str(tmp_pdf)}'; "
        "if (Test-Path -LiteralPath $tmp) { Remove-Item -LiteralPath $tmp -Force }; "
        "$word = New-Object -ComObject Word.Application; "
        "$word.Visible = $false; "
        "try { "
        f"$doc = $word.Documents.Open('{str(DST)}', $false, $false); "
        "$doc.Fields.Update() | Out-Null; "
        "foreach ($toc in $doc.TablesOfContents) { $toc.Update() | Out-Null }; "
        "$doc.Save(); "
        "$doc.ExportAsFixedFormat($tmp, 17); "
        "$doc.Close($false); "
        "} finally { $word.Quit() }; "
        f"Move-Item -LiteralPath $tmp -Destination '{str(PDF)}' -Force"
    )
    subprocess.run(["powershell", "-Command", ps], cwd=str(ROOT), check=True, capture_output=True, text=True)


def export_pdf() -> bool:
    if sys.platform.startswith("win"):
        try:
            update_pdf_export_fields_with_word()
            return True
        except Exception as err:
            print(f"Microsoft Word PDF export failed: {err}")
    return False


def main() -> int:
    if not SRC.exists():
        raise FileNotFoundError(SRC)
    shutil.copy2(SRC, DST)
    doc = Document(DST)
    before_shapes = len(doc.inline_shapes)

    placeholders = update_title_page(doc)
    enforce_clean_page_starts(doc)
    clean_abstract_acronyms(doc)
    format_equations(doc)
    strengthen_reproducibility(doc)
    strengthen_appendix_a(doc)
    expand_appendix_b(doc)
    update_thesis36_labels(doc)
    standardize_table_fonts(doc)

    after_shapes = len(doc.inline_shapes)
    if after_shapes != before_shapes:
        raise RuntimeError(f"Figure preservation failed: inline_shapes {before_shapes} -> {after_shapes}")
    doc.save(DST)
    print(f"Wrote {DST.relative_to(ROOT)}")
    print(f"Preserved inline_shapes: {after_shapes}")
    print("Placeholders requiring user input: " + ", ".join(placeholders))

    if export_pdf():
        print(f"Wrote {PDF.relative_to(ROOT)}")
    else:
        print(f"PDF export failed; DOCX remains available at {DST.relative_to(ROOT)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
