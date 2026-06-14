"""Build the named final submission artifacts from thesis_38.

This is a document-only transformation. It does not run experiments or
regenerate evidence, tables, figures, or protected CSV files.
"""

from __future__ import annotations

import argparse
import hashlib
import shutil
import subprocess
from pathlib import Path
from zipfile import ZipFile

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "manuscript" / "thesis_38.docx"
OUTPUT_DOCX = ROOT / "manuscript" / "Onur_Emiroglu_MSc_Thesis.docx"
OUTPUT_PDF = ROOT / "manuscript" / "Onur_Emiroglu_MSc_Thesis.pdf"

OLD_HEADING = "8.4 Reward-Shaping and Misspecification Checks"
NEW_HEADING = "8.4 Experiment 4: Reward-Shaping and Misspecification Checks"

PROTECTED_HASHES = {
    ROOT / "results" / "metrics_detector_compare.csv": (
        "28E7AD40BB47214F8576132846E9E1D4CD643F623CF1187743091FC367A206ED"
    ),
    ROOT / "docs" / "internal" / "wp6_sweep_full" / "summary_condition_variant.csv": (
        "6DD627E81637A49A60163F58AC1D3EF23B8D694E39AC55BA64FBF808E978C6EA"
    ),
    ROOT
    / "docs"
    / "internal"
    / "wp6_sweep_full"
    / "summary_paired_combined_vs_sigma.csv": (
        "4BABCAAACE1DD5228C674E2CED9D977236F8D3ACB503C098FAEB06FF6C10B796"
    ),
    ROOT
    / "docs"
    / "internal"
    / "wp6_sweep_full"
    / "summary_paired_combined_vs_regime.csv": (
        "2087FEFBE5DC39AF23372EA2D8999AC0F1071D0FEE90BC1B3668F2158130E8F9"
    ),
}


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest().upper()


def verify_protected_hashes() -> None:
    mismatches = []
    for path, expected in PROTECTED_HASHES.items():
        actual = sha256(path)
        if actual != expected:
            mismatches.append(f"{path.relative_to(ROOT)}: {actual} != {expected}")
    if mismatches:
        raise RuntimeError("Protected evidence mismatch:\n" + "\n".join(mismatches))


def nonempty(value: str) -> str:
    cleaned = value.strip()
    if not cleaned or cleaned.upper() == "[INSERT]":
        raise argparse.ArgumentTypeError("must be an exact non-placeholder value")
    return cleaned


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Create the named final submission DOCX and PDF from thesis_38."
    )
    parser.add_argument("--supervisor", required=True, type=nonempty)
    parser.add_argument("--matriculation-no", required=True, type=nonempty)
    parser.add_argument("--course-module", default="MSc Thesis", type=nonempty)
    advisor = parser.add_mutually_exclusive_group(required=True)
    advisor.add_argument("--advisor", type=nonempty)
    advisor.add_argument("--advisor-not-applicable", action="store_true")
    return parser.parse_args()


def media_hashes(path: Path) -> dict[str, str]:
    with ZipFile(path) as archive:
        return {
            name: hashlib.sha256(archive.read(name)).hexdigest().upper()
            for name in archive.namelist()
            if name.startswith("word/media/")
        }


def set_title_fields(doc: Document, values: dict[str, str]) -> None:
    if not doc.tables:
        raise RuntimeError("The source document has no title-page table.")
    table = doc.tables[0]
    found: set[str] = set()
    for row in table.rows:
        if len(row.cells) < 2:
            continue
        label = row.cells[0].text.strip()
        if label in values:
            row.cells[1].text = values[label]
            found.add(label)
    missing = set(values) - found
    if missing:
        raise RuntimeError(f"Title-page fields not found: {sorted(missing)}")


def clarify_heading(doc: Document) -> None:
    matches = [p for p in doc.paragraphs if p.text.strip() == OLD_HEADING]
    if len(matches) != 1:
        raise RuntimeError(
            f"Expected one exact Section 8.4 heading, found {len(matches)}."
        )
    paragraph = matches[0]
    if paragraph.runs:
        paragraph.runs[0].text = NEW_HEADING
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.text = NEW_HEADING


def iter_text(doc: Document):
    for paragraph in doc.paragraphs:
        yield paragraph.text
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                yield cell.text
    for section in doc.sections:
        for part in (section.header, section.footer):
            for paragraph in part.paragraphs:
                yield paragraph.text


def validate_docx(doc: Document, expected_shapes: int) -> None:
    text = "\n".join(iter_text(doc))
    forbidden = {
        "[To be provided]": text.count("[To be provided]"),
        "[To be provided, or Not applicable]": text.count(
            "[To be provided, or Not applicable]"
        ),
        "Draft": text.count("Draft"),
        "+/-": text.count("+/-"),
    }
    remaining = {term: count for term, count in forbidden.items() if count}
    if remaining:
        raise RuntimeError(f"Forbidden text remains: {remaining}")
    if text.count(NEW_HEADING) != 1:
        raise RuntimeError("Clarified Section 8.4 heading is not present exactly once.")
    if len(doc.inline_shapes) != expected_shapes:
        raise RuntimeError(
            f"Inline shape count changed: {expected_shapes} -> {len(doc.inline_shapes)}"
        )


def export_pdf() -> None:
    temporary_pdf = OUTPUT_PDF.with_suffix(".tmp.pdf")
    escaped_docx = str(OUTPUT_DOCX).replace("'", "''")
    escaped_pdf = str(OUTPUT_PDF).replace("'", "''")
    escaped_tmp = str(temporary_pdf).replace("'", "''")
    script = (
        f"$docx = '{escaped_docx}'; "
        f"$pdf = '{escaped_pdf}'; "
        f"$tmp = '{escaped_tmp}'; "
        "if (Test-Path -LiteralPath $tmp) { Remove-Item -LiteralPath $tmp -Force }; "
        "$word = New-Object -ComObject Word.Application; "
        "$word.Visible = $false; "
        "try { "
        "$doc = $word.Documents.Open($docx, $false, $false); "
        "$doc.Fields.Update() | Out-Null; "
        "foreach ($toc in $doc.TablesOfContents) { $toc.Update() | Out-Null }; "
        "$doc.Save(); "
        "$doc.ExportAsFixedFormat($tmp, 17); "
        "$doc.Close($false); "
        "} finally { $word.Quit() }; "
        "Move-Item -LiteralPath $tmp -Destination $pdf -Force"
    )
    subprocess.run(
        ["powershell", "-Command", script],
        cwd=ROOT,
        check=True,
        capture_output=True,
        text=True,
    )


def main() -> int:
    args = parse_args()
    advisor = "Not applicable" if args.advisor_not_applicable else args.advisor

    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)

    verify_protected_hashes()
    source_hash = sha256(SOURCE)
    source_media = media_hashes(SOURCE)

    shutil.copy2(SOURCE, OUTPUT_DOCX)
    doc = Document(OUTPUT_DOCX)
    expected_shapes = len(doc.inline_shapes)
    set_title_fields(
        doc,
        {
            "Supervisor": args.supervisor,
            "Matriculation No.": args.matriculation_no,
            "Course / Module": args.course_module,
            "Advisor / Co-supervisor": advisor,
        },
    )
    clarify_heading(doc)
    validate_docx(doc, expected_shapes)
    doc.save(OUTPUT_DOCX)

    exported = Document(OUTPUT_DOCX)
    validate_docx(exported, expected_shapes)
    if media_hashes(OUTPUT_DOCX) != source_media:
        raise RuntimeError("Embedded media changed relative to thesis_38.")

    export_pdf()

    if sha256(SOURCE) != source_hash:
        raise RuntimeError("thesis_38.docx changed during submission generation.")
    verify_protected_hashes()

    print(f"Wrote {OUTPUT_DOCX.relative_to(ROOT)}")
    print(f"Wrote {OUTPUT_PDF.relative_to(ROOT)}")
    print(f"Inline shapes preserved: {expected_shapes}")
    print("Protected evidence hashes: 4/4 MATCH")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
