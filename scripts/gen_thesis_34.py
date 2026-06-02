"""Generate thesis_34.docx/pdf from thesis_33 with defense-clarity polish.

This is a documentation-only generator. It does not run PPO, WP5, WP6, detector
robustness, ablations, misspecification, signal-informativeness sweeps, or any
protected evidence pipeline. It preserves figures and tables from thesis_33 and
adds only reader-facing clarification text.
"""

from __future__ import annotations

import shutil
import subprocess
import sys
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "manuscript" / "thesis_33.docx"
DST = ROOT / "manuscript" / "thesis_34.docx"
PDF = ROOT / "manuscript" / "thesis_34.pdf"

DRAFT_NOTE = (
    "Draft integrity note: thesis_34 is a documentation and defense-clarity polish pass. "
    "Frozen experimental evidence is reused and not regenerated."
)
MARKET_CONTEXT = (
    "Market context. The simulator is asset-class agnostic and uses normalized price units. It should be read as a "
    "stylized liquid electronic limit-order-book market rather than as a calibrated model of a specific venue such "
    "as equities, futures, crypto, or BIST. The fee and latency parameters are controlled friction assumptions used "
    "to make strategy comparisons internally consistent. They are not intended to identify a particular exchange "
    "microstructure."
)
MTM_CAVEAT = (
    "The equity measure is marked to the mid-price for consistency across simulated strategies. This is an "
    "evaluation convention inside the controlled simulator and should not be interpreted as immediately liquidatable "
    "wealth in a real limit-order book. In practice, liquidating inventory would generally require crossing the "
    "spread, facing queue priority and available depth, and potentially incurring additional adverse-selection or "
    "market-impact costs. These effects are part of the real-market boundary discussed later, not part of the "
    "current controlled experiment."
)
FIGURE_88_INTERPRETATION = (
    "One possible interpretation is that the degraded continuous signal still preserves enough ordinal volatility "
    "information for quote-width adaptation, whereas the categorical label is coarser and does not add useful "
    "variation in this calibration. In noisy or lagged settings, adding the label may also increase the observation "
    "dimension without improving the control-relevant signal. This should be read as a performance pattern in the "
    "tested calibration band, not as proof of the PPO policy's internal representation mechanism."
)
COMPUTE_NOTE = (
    "The implementation prioritizes reproducibility and auditability over maximum simulation throughput. Future "
    "versions could improve runtime through vectorized environments, parallel rollouts, or lower-level simulator "
    "components for the execution loop."
)
SECTION5_BRIDGE = (
    "Section 4 describes the software/evidence architecture, while this section restates the same workflow as a "
    "conceptual layered experimental model."
)
AS_CAVEAT = (
    "The AS implementation is used as a canonical analytical inventory-aware comparator rather than as an "
    "exhaustively optimized trading system. The PPO-versus-AS comparison establishes a reference point for learned "
    "quoting and inventory control, but the central thesis claim is the within-PPO signal-design comparison between "
    "sigma_hat, estimated labels, and oracle labels."
)
ORACLE_PURE_NOTE = (
    "The label-quality objection is also bounded by the oracle_pure result. oracle_pure receives the true categorical "
    "regime label but not sigma_hat, yet it remains below sigma_only in mean Sharpe-like performance in the "
    "five-variant ablation. This suggests that the limitation is not only detector noise: in this calibration, the "
    "continuous volatility signal is more useful for quote control than the discretized regime category."
)
MULTIPLE_COMPARISONS_NOTE = (
    "Because the manuscript reports several paired comparisons, isolated marginal p-values should be read "
    "descriptively rather than as a stand-alone familywise discovery claim. The main conclusion is based on the "
    "repeated pattern across the main OOS evaluation, detector robustness, oracle ablations, reward-shaping, "
    "misspecification, and signal-informativeness diagnostics."
)
TOST_BOUND_NOTE = (
    "The equivalence bounds should be read as practical smallest-effect-size thresholds for the Sharpe-like metric, "
    "not as universal constants. A wider +/-0.10 band is used for the main five-variant ablation, where the question "
    "is whether oracle regime information creates a practically meaningful improvement over sigma_hat alone. Tighter "
    "+/-0.05 bounds are used in narrower robustness or misspecification checks. The thesis therefore does not treat "
    "TOST as a stand-alone proof; it reports TOST together with paired tests, confidence intervals, and directional "
    "seed-level evidence."
)
INTERPRETATION_REFINEMENT = (
    "The evidence should therefore be read slightly more precisely than pure redundancy. If a categorical label only "
    "repeated the useful information in sigma_hat, combined and sigma_only would be expected to be approximately "
    "indistinguishable. In several experiments, however, combined is directionally below sigma_only. This pattern is "
    "consistent with signal redundancy plus mild categorical-channel degradation: the label may be coarser than the "
    "continuous volatility estimate, may add observation dimensionality, or may interact with policy learning without "
    "adding control-relevant information. This is an interpretation of the observed performance pattern, not a proof "
    "of the PPO policy's internal mechanism."
)
OOS_LIMITATION = (
    "The out-of-sample protocol should also be interpreted carefully. It is a chronological 70/30 temporal hold-out "
    "within the same synthetic data-generating process, not a distributional-shift test across fundamentally "
    "different markets. Multiple seeds improve robustness to simulation randomness, but they do not establish "
    "external validity beyond the controlled synthetic environment."
)
HYPERPARAMETER_NOTE = (
    "The PPO hyperparameters are held fixed across PPO variants to keep the signal-channel comparison fair. They "
    "should be read as canonical experiment settings rather than as a claim of globally optimal PPO tuning. The "
    "inventory penalty eta = 0.001 is the canonical reward scale used in the main experiments after the "
    "inventory-penalty ablation."
)
SHARPE_FORMULA = "(16)  Sharpe-like = mean(Delta W_t) / std(Delta W_t) * sqrt(1 / dt)"
SHARPE_DEFINITION = (
    "Let Delta W_t = W_t - W_{t-1}. The reported Sharpe-like metric is computed as mean(Delta W_t) / "
    "std(Delta W_t), using the sample standard deviation with ddof = 1, multiplied by sqrt(1 / dt) when the standard "
    "deviation is positive; otherwise it is reported as 0. It is therefore a simulator-scale risk-adjusted PnL "
    "metric rather than an annualized market Sharpe ratio."
)
ABSTRACT_QUANT_SENTENCE = (
    "In the five-variant ablation, sigma_only achieved the highest mean Sharpe-like value (0.753), while oracle_full "
    "did not provide a statistically significant or practically meaningful improvement over sigma_hat alone (paired "
    "t-test p = 0.115; TOST +/-0.10 p = 0.00067; 90% CI [-0.001, +0.063])."
)
LAGGED_DEGRADATION_SENTENCE = (
    "The strongest degradation pattern appears in the lagged condition, where combined is materially below sigma_only "
    "(mean difference about -0.101; Cohen's dz about -0.91). This suggests that a stale categorical channel can be "
    "harmful when it adds delayed/coarse state information rather than control-relevant volatility information."
)
MISSPEC_TOST_QUALIFICATION = (
    "In the misspecification check, the mean difference between sigma_only and oracle_full is very small, but the "
    "+/-0.05 TOST result is close to the threshold; it should therefore be read as supportive but less decisive than "
    "the main five-variant ablation equivalence result."
)


def iter_all_paragraphs(doc: Document):
    paragraphs = list(doc.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                paragraphs.extend(cell.paragraphs)
    for section in doc.sections:
        paragraphs.extend(section.header.paragraphs)
        paragraphs.extend(section.footer.paragraphs)
    return paragraphs


def has_drawing(paragraph) -> bool:
    return any(run._r.xpath(".//w:drawing") for run in paragraph.runs)


def replace_text_in_paragraph(paragraph, old: str, new: str) -> None:
    if has_drawing(paragraph):
        return
    full_text = "".join(run.text for run in paragraph.runs)
    if old not in full_text:
        return
    replaced = full_text.replace(old, new)
    if paragraph.runs:
        paragraph.runs[0].text = replaced
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(replaced)


def find_paragraph(doc: Document, exact_text: str):
    for paragraph in doc.paragraphs:
        if paragraph.text.strip() == exact_text:
            return paragraph
    raise RuntimeError(f"Could not find paragraph: {exact_text}")


def insert_paragraph_before(doc: Document, ref_para, text: str = "", style: str | None = None):
    paragraph = doc.add_paragraph(text, style=style)
    ref_para._p.addprevious(paragraph._p)
    return paragraph


def insert_paragraph_after(anchor, text: str, style: str | None = None):
    paragraph = anchor._parent.add_paragraph(text, style=style)
    anchor._p.addnext(paragraph._p)
    return paragraph


def disable_auto_hyphenation(doc: Document) -> None:
    settings = doc.settings.element
    for element in settings.findall(qn("w:autoHyphenation")):
        settings.remove(element)
    auto_hyphenation = OxmlElement("w:autoHyphenation")
    auto_hyphenation.set(qn("w:val"), "false")
    settings.append(auto_hyphenation)

    for paragraph in iter_all_paragraphs(doc):
        p_pr = paragraph._p.get_or_add_pPr()
        if p_pr.find(qn("w:suppressAutoHyphens")) is None:
            suppress = OxmlElement("w:suppressAutoHyphens")
            suppress.set(qn("w:val"), "true")
            p_pr.append(suppress)


def clean_encoding_artifacts(doc: Document) -> None:
    replacements = {
        "\ufffe": "-",
        "\ufffd": "",
        "\u00ad": "-",
        "out-of\ufffesample": "out-of-sample",
        "information\ufffedesign": "information-design",
        "order\ufffebook": "order-book",
        "real\ufffeorder-book": "real-order-book",
        "Sharpe\ufffelike": "Sharpe-like",
        "one\ufffehot": "one-hot",
        "practical\ufffeequivalence": "practical-equivalence",
        "signal\ufffeinformativeness": "signal-informativeness",
        "out-of\u00adsample": "out-of-sample",
        "information\u00addesign": "information-design",
        "order\u00adbook": "order-book",
        "real\u00adorder-book": "real-order-book",
        "Sharpe\u00adlike": "Sharpe-like",
        "one\u00adhot": "one-hot",
        "practical\u00adequivalence": "practical-equivalence",
        "signal\u00adinformativeness": "signal-informativeness",
        "out-ofï¿¾sample": "out-of-sample",
        "informationï¿¾design": "information-design",
        "orderï¿¾book": "order-book",
        "realï¿¾order-book": "real-order-book",
        "Sharpeï¿¾like": "Sharpe-like",
        "oneï¿¾hot": "one-hot",
        "practicalï¿¾equivalence": "practical-equivalence",
        "signalï¿¾informativeness": "signal-informativeness",
    }
    for paragraph in iter_all_paragraphs(doc):
        for run in paragraph.runs:
            if run._r.xpath(".//w:drawing"):
                continue
            text = run.text
            for old, new in replacements.items():
                text = text.replace(old, new)
            if text != run.text:
                run.text = text


def update_cover_and_footer(doc: Document) -> None:
    if not doc.tables:
        raise RuntimeError("Expected cover metadata table")
    cover = doc.tables[0]
    for row in cover.rows:
        label = row.cells[0].text.strip()
        if label == "Current draft":
            row.cells[1].text = "manuscript/thesis_34.pdf / .docx"
        elif label == "Previous source draft":
            row.cells[1].text = "manuscript/thesis_33.pdf / .docx"

    for paragraph in doc.paragraphs:
        if paragraph.text.strip() == "Academic-polish draft; frozen evidence is not regenerated.":
            paragraph.text = DRAFT_NOTE

    for section in doc.sections:
        for paragraph in section.footer.paragraphs:
            if "Academic-polish draft; frozen evidence is not regenerated." in paragraph.text:
                paragraph.text = ""


def update_version_text(doc: Document) -> None:
    replacements = {
        "Status in thesis_33": "Status in thesis_34",
        "thesis_33 is a new template-adapted draft, not a replacement for thesis_29 or thesis_30.": (
            "thesis_34 is a defense-clarity polish draft built on thesis_33; frozen evidence remains unchanged."
        ),
        "These commands document reproducibility paths only; thesis_33 adaptation did not rerun experiments.": (
            "These commands document reproducibility paths only; thesis_34 adaptation did not rerun experiments."
        ),
        "Evaluation writes CSV metrics, figures, and summaries. The thesis_33 adaptation embeds existing figures and quotes frozen numerical evidence; it does not regenerate protected evidence artifacts.": (
            "Evaluation writes CSV metrics, figures, and summaries. The thesis_34 adaptation embeds existing figures and quotes frozen numerical evidence; it does not regenerate protected evidence artifacts."
        ),
        "The thesis_33 adaptation reuses frozen performance results. It does not rerun PPO training, WP5/WP6 experiments, detector robustness, ablations, misspecification checks, protected CSV generation, or frozen performance-figure generation. The only newly generated visuals are methodology-only illustrations produced by `scripts/figures/gen_synthetic_environment_figure.py` and `scripts/figures/gen_pipeline_architecture_figure.py`.": (
            "The thesis_34 adaptation reuses frozen performance results. It does not rerun PPO training, WP5/WP6 experiments, detector robustness, ablations, misspecification checks, protected CSV generation, or frozen performance-figure generation. The only newly generated visuals are methodology-only illustrations produced by `scripts/figures/gen_synthetic_environment_figure.py` and `scripts/figures/gen_pipeline_architecture_figure.py`."
        ),
        "Ensures thesis_33 references existing frozen figure files.": (
            "Ensures thesis_34 references existing frozen figure files."
        ),
        "Defense-safe caveat": "Caveat",
        "Defense-safe reading": "Interpretation",
        "(16)  Sharpe-like = mean(g_t) / std(g_t)": SHARPE_FORMULA,
        "Here g_t denotes the logged per-step PnL or return-like series used by the project metric convention. This is the primary risk-adjusted performance metric.": SHARPE_DEFINITION,
    }
    for paragraph in iter_all_paragraphs(doc):
        for old, new in replacements.items():
            replace_text_in_paragraph(paragraph, old, new)


def add_market_context(doc: Document) -> None:
    anchor = find_paragraph(
        doc,
        "Canonical experiments are driven by JSON configs under `config/`. Shared parameters include mid0 = 100.0, tick_size = 0.01, dt = 0.2, baseline sigma_mid_ticks = 0.8, A = 5.0, k = 1.5, fee_bps = 0.2, and latency_steps = 1.",
    )
    insert_paragraph_after(anchor, MARKET_CONTEXT)


def add_midprice_mtm_caveat(doc: Document) -> None:
    anchor = find_paragraph(
        doc,
        "Here W_t is mark-to-market wealth or equity, X_t is cash, q_t is inventory, and S_t is the mid-price.",
    )
    insert_paragraph_after(anchor, MTM_CAVEAT)


def add_figure_88_interpretation(doc: Document) -> None:
    anchor = find_paragraph(
        doc,
        "Figure 8.8. Experiment 5 monotonic-gap plot: the expected narrowing of the sigma_only versus combined gap does not appear in the tested calibration band.",
    )
    inserted = insert_paragraph_after(anchor, FIGURE_88_INTERPRETATION)
    inserted.paragraph_format.keep_with_next = True
    lagged = insert_paragraph_after(inserted, LAGGED_DEGRADATION_SENTENCE)
    lagged.paragraph_format.keep_with_next = True


def add_compute_note(doc: Document) -> None:
    anchor = find_paragraph(
        doc,
        "Simulation and deterministic evaluation are linear in the number of environment steps. PPO training cost is approximately linear in total timesteps, policy-network forward/backward passes, and the number of seed-variant cells. The WP6 full sweep is expensive because it multiplies conditions, variants, and seeds; this is why its outputs are treated as frozen evidence.",
    )
    insert_paragraph_after(anchor, COMPUTE_NOTE)


def add_midprice_risk_row(doc: Document) -> None:
    for table in doc.tables:
        if len(table.rows) >= 1 and len(table.columns) >= 2:
            headers = [cell.text.strip() for cell in table.rows[0].cells]
            if headers[:2] == ["Risk", "Mitigation"]:
                for row in table.rows:
                    if row.cells[0].text.strip() == "Mid-price mark-to-market assumption":
                        row.cells[1].text = (
                            "Report it as a simulator evaluation convention; real liquidation would require "
                            "spread/depth/queue/adverse-selection considerations."
                        )
                        return
                row = table.add_row()
                row.cells[0].text = "Mid-price mark-to-market assumption"
                row.cells[1].text = (
                    "Report it as a simulator evaluation convention; real liquidation would require "
                    "spread/depth/queue/adverse-selection considerations."
                )
                return
    raise RuntimeError("Could not find risks table")


def add_eta_hyperparameter_row(doc: Document) -> None:
    for table in doc.tables:
        if len(table.rows) >= 1 and len(table.columns) >= 2:
            headers = [cell.text.strip() for cell in table.rows[0].cells]
            if headers[:2] == ["Hyperparameter", "Canonical full-run value"]:
                for row in table.rows:
                    if row.cells[0].text.strip() == "eta":
                        row.cells[1].text = "0.001"
                        return
                row = table.add_row()
                row.cells[0].text = "eta"
                row.cells[1].text = "0.001"
                return
    raise RuntimeError("Could not find PPO hyperparameter table")


def add_hardening_paragraphs(doc: Document) -> None:
    insert_paragraph_after(find_paragraph(doc, "5 Layered Experimental Model"), SECTION5_BRIDGE)
    insert_paragraph_after(
        find_paragraph(
            doc,
            "AS has higher raw equity but carries substantially larger inventory tail risk; therefore the main comparison is risk-adjusted performance and inventory control, not raw equity alone.",
        ),
        AS_CAVEAT,
    )
    insert_paragraph_after(
        find_paragraph(
            doc,
            "The strongest ablation result is that sigma_only has the highest mean Sharpe-like value, while oracle_full does not significantly beat it. TOST supports practical equivalence under the +/-0.10 Sharpe-like bound.",
        ),
        ORACLE_PURE_NOTE,
    )
    insert_paragraph_after(
        find_paragraph(
            doc,
            "Here d_i is the paired seed-level difference between two strategies, sd(d_i) is its sample standard deviation, and n is the number of paired seeds.",
        ),
        MULTIPLE_COMPARISONS_NOTE,
    )
    insert_paragraph_after(
        find_paragraph(
            doc,
            "TOST evaluates whether the plausible range of a paired difference lies inside a pre-specified practical-equivalence band. At alpha = 0.05, the corresponding equivalence reading uses the 90% confidence interval. The Five-Variant Signal Ablation uses TOST to support practical equivalence for sigma_only versus oracle_full; the Signal-Informativeness Sweep also uses related paired tests to examine non-equivalence and mean indistinguishability patterns.",
        ),
        TOST_BOUND_NOTE,
    )
    insert_paragraph_after(
        find_paragraph(
            doc,
            "The interpretation is that the continuous realized-volatility proxy already supplies the economically useful volatility information needed by the policy in this environment. The categorical label may be coarser or redundant when presented alongside sigma_hat.",
        ),
        INTERPRETATION_REFINEMENT,
    )
    insert_paragraph_after(
        find_paragraph(
            doc,
            "For that reason, the thesis should be read as controlled synthetic evidence, not as a real-market external validity claim. Richer LOB simulators, Hawkes-process LOB models, queue-reactive execution, and real-data evaluation are natural future work and could change the incremental value of regime labels.",
        ),
        OOS_LIMITATION,
    )
    insert_paragraph_after(find_paragraph(doc, "7.6 PPO Hyperparameters"), HYPERPARAMETER_NOTE)


def add_micro_polish_sentences(doc: Document) -> None:
    evidence_sentence = (
        "The evidence comes from the canonical frozen experiments: out-of-sample evaluation, detector robustness, "
        "oracle-label ablations, regime-conditional reward shaping, mild model misspecification, and a "
        "signal-informativeness sweep."
    )
    abstract_anchor = find_paragraph(doc, "Abstract")
    paragraphs = list(doc.paragraphs)
    abstract_idx = next(i for i, paragraph in enumerate(paragraphs) if paragraph._p is abstract_anchor._p)
    for paragraph in paragraphs[abstract_idx + 1 :]:
        if evidence_sentence in paragraph.text:
            replace_text_in_paragraph(paragraph, evidence_sentence, f"{evidence_sentence} {ABSTRACT_QUANT_SENTENCE}")
            break
    else:
        raise RuntimeError("Could not find abstract evidence sentence")

    insert_paragraph_after(
        find_paragraph(
            doc,
            "The regime-conditional eta run tests whether explicit labels become useful when the reward penalizes high-volatility inventory more strongly. sigma_only still beats combined on Sharpe-like performance (p = 0.0016). Under mild regime-dependent execution misspecification, sigma_only and oracle_full remain statistically indistinguishable and practically equivalent under the reported TOST bound.",
        ),
        MISSPEC_TOST_QUALIFICATION,
    )


def export_pdf_with_word() -> None:
    tmp_pdf = PDF.with_name(PDF.stem + "_export_tmp.pdf")
    if tmp_pdf.exists():
        tmp_pdf.unlink()
    ps = (
        "$ErrorActionPreference = 'Stop'; "
        f"$tmp = '{str(tmp_pdf)}'; "
        "if (Test-Path -LiteralPath $tmp) { Remove-Item -LiteralPath $tmp -Force }; "
        "$word = New-Object -ComObject Word.Application; "
        "$word.Visible = $false; "
        "try { "
        f"$doc = $word.Documents.Open('{str(DST)}', $false, $true); "
        "$doc.ExportAsFixedFormat($tmp, 17); "
        "$doc.Close($false); "
        "} finally { $word.Quit() }; "
        f"Move-Item -LiteralPath $tmp -Destination '{str(PDF)}' -Force"
    )
    subprocess.run(["powershell", "-Command", ps], cwd=str(ROOT), check=True, capture_output=True, text=True)


def export_pdf() -> bool:
    if sys.platform.startswith("win"):
        try:
            export_pdf_with_word()
            return True
        except Exception as word_err:
            print(f"Microsoft Word PDF export failed: {word_err}")
    return False


def main() -> int:
    if not SRC.exists():
        raise FileNotFoundError(SRC)

    shutil.copy2(SRC, DST)
    doc = Document(DST)
    before_shapes = len(doc.inline_shapes)

    update_cover_and_footer(doc)
    update_version_text(doc)
    add_market_context(doc)
    add_midprice_mtm_caveat(doc)
    add_figure_88_interpretation(doc)
    add_compute_note(doc)
    add_midprice_risk_row(doc)
    add_eta_hyperparameter_row(doc)
    add_hardening_paragraphs(doc)
    add_micro_polish_sentences(doc)
    clean_encoding_artifacts(doc)
    disable_auto_hyphenation(doc)

    after_shapes = len(doc.inline_shapes)
    if after_shapes != before_shapes:
        raise RuntimeError(f"Figure preservation failed: inline_shapes {before_shapes} -> {after_shapes}")

    doc.save(DST)
    print(f"Wrote {DST.relative_to(ROOT)}")
    print(f"Preserved inline_shapes: {after_shapes}")

    if export_pdf():
        print(f"Wrote {PDF.relative_to(ROOT)}")
    else:
        print(f"PDF export failed; DOCX remains available at {DST.relative_to(ROOT)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
