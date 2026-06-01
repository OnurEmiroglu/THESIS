"""Generate thesis_31.docx as an English technical-report adaptation.

This generator is documentation-only. It does not run experiments, retrain PPO,
regenerate plots, or modify protected evidence. It embeds existing frozen figure
files and reproduces numerical claims from the documented thesis evidence base.
"""

from __future__ import annotations

import hashlib
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "manuscript" / "thesis_31.docx"

FIGURES = {
    "fig1": ROOT / "results/plots/thesis/fig1_sharpe_inv.png",
    "fig2": ROOT / "results/plots/thesis/fig2_paired_seed.png",
    "fig3": ROOT / "results/plots/thesis_23/fig6_ablation_summary.png",
    "fig4": ROOT / "results/plots/thesis_23/fig7_oracle_paired_seed.png",
    "fig5": ROOT / "results/plots/thesis/fig4_detector_robustness.png",
    "fig6": ROOT / "results/plots/thesis_23/fig8_eta_regime_summary.png",
    "fig7": ROOT / "results/plots/thesis_23/fig9_misspec_summary.png",
    "fig8": ROOT / "docs/internal/wp6_sweep_full/plots/monotonic_gap.png",
    "fig9": ROOT / "docs/internal/wp6_sweep_full/plots/paired_seed_combined_vs_sigma.png",
    "fig10": ROOT / "docs/internal/wp6_sweep_full/plots/paired_seed_combined_vs_regime.png",
    "figA": ROOT / "results/plots/thesis/fig5_action_analysis.png",
    "figB": ROOT / "results/plots/thesis/fig3_regime_sharpe.png",
}

PROTECTED_HASHES = {
    "results/metrics_detector_compare.csv": "28E7AD40BB47214F8576132846E9E1D4CD643F623CF1187743091FC367A206ED",
    "docs/internal/wp6_sweep_full/summary_condition_variant.csv": "6DD627E81637A49A60163F58AC1D3EF23B8D694E39AC55BA64FBF808E978C6EA",
    "docs/internal/wp6_sweep_full/summary_paired_combined_vs_sigma.csv": "4BABCAAACE1DD5228C674E2CED9D977236F8D3ACB503C098FAEB06FF6C10B796",
    "docs/internal/wp6_sweep_full/summary_paired_combined_vs_regime.csv": "2087FEFBE5DC39AF23372EA2D8999AC0F1071D0FEE90BC1B3668F2158130E8F9",
}

CENTRAL_CLAIM = (
    "In the tested controlled synthetic HFMM environment, explicit categorical "
    "volatility-regime labels do not provide robust incremental value once "
    "sigma_hat is already observed by the PPO policy. The results are "
    "consistent with signal redundancy."
)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa: list[int]) -> None:
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    grid = tbl.tblGrid
    if grid is None:
        grid = OxmlElement("w:tblGrid")
        tbl.insert(0, grid)
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            cell.width = Pt(widths_dxa[i] / 20)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[i]))
            tc_w.set(qn("w:type"), "dxa")
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[int] | None = None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = h
        set_cell_shading(hdr[i], "F2F4F7")
        for p in hdr[i].paragraphs:
            for run in p.runs:
                run.bold = True
                run.font.size = Pt(9)
    for row_values in rows:
        row = table.add_row().cells
        for i, value in enumerate(row_values):
            row[i].text = str(value)
            for p in row[i].paragraphs:
                for run in p.runs:
                    run.font.size = Pt(9)
    set_table_geometry(table, widths or balanced_widths(len(headers)))
    doc.add_paragraph()
    return table


def balanced_widths(n: int) -> list[int]:
    base = 9360 // n
    widths = [base] * n
    widths[-1] += 9360 - sum(widths)
    return widths


def add_caption(doc: Document, text: str) -> None:
    p = doc.add_paragraph(text)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.style = doc.styles["Caption"]


def add_figure(doc: Document, key: str, caption: str) -> None:
    path = FIGURES[key]
    if not path.exists():
        p = doc.add_paragraph(f"[Missing figure: {path.relative_to(ROOT)}]")
        p.style = doc.styles["Caption"]
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(path), width=Inches(6.1))
    add_caption(doc, caption)


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def add_numbered(doc: Document, items: list[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Number")


def add_code_block(doc: Document, title: str, lines: list[str]) -> None:
    doc.add_paragraph(title, style="Heading 3")
    for line in lines:
        p = doc.add_paragraph()
        run = p.add_run(line)
        run.font.name = "Consolas"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
        run.font.size = Pt(9)
    doc.add_paragraph()


def add_equation(doc: Document, equation: str, definitions: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(equation)
    run.bold = True
    run.font.name = "Cambria Math"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Cambria Math")
    run.font.size = Pt(11)
    doc.add_paragraph(definitions)


def style_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    title = styles["Title"]
    title.font.name = "Calibri"
    title.font.size = Pt(24)
    title.font.bold = True
    title.font.color.rgb = RGBColor(11, 37, 69)
    title.paragraph_format.space_after = Pt(12)

    for name, size, color, before, after in [
        ("Heading 1", 16, RGBColor(46, 116, 181), 16, 8),
        ("Heading 2", 13, RGBColor(46, 116, 181), 12, 6),
        ("Heading 3", 12, RGBColor(31, 77, 120), 8, 4),
    ]:
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)

    caption = styles["Caption"]
    caption.font.name = "Calibri"
    caption.font.size = Pt(9)
    caption.font.italic = True
    caption.font.color.rgb = RGBColor(85, 85, 85)
    caption.paragraph_format.space_after = Pt(8)

    header = section.header.paragraphs[0]
    header.text = "HFMM-RL Thesis 31 Template Draft"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs:
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(85, 85, 85)

    footer = section.footer.paragraphs[0]
    footer.text = "Template-adapted draft; frozen evidence is not regenerated."
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in footer.runs:
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(85, 85, 85)


def title_page(doc: Document) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("High-Frequency Market Making via Reinforcement Learning under Different Volatility Regimes")
    run.bold = True
    run.font.size = Pt(22)
    run.font.color.rgb = RGBColor(11, 37, 69)
    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Technical Report Template Adaptation - thesis_31")
    run.font.size = Pt(15)
    run.bold = True
    doc.add_paragraph()
    rows = [
        ["Project", "Synthetic HFMM reinforcement-learning thesis"],
        ["Draft", "manuscript/thesis_31.docx and manuscript/thesis_31.pdf"],
        ["Frozen baseline", "manuscript/thesis_29.pdf, tag thesis-v29-frozen"],
        ["Source draft", "manuscript/thesis_30.pdf / .docx"],
        ["Decision log", "manuscript/decisions_log_13.pdf"],
        ["Scope", "Template adaptation only; no experiment reruns or figure regeneration"],
    ]
    add_table(doc, ["Field", "Value"], rows, [2200, 7160])
    doc.add_paragraph(CENTRAL_CLAIM)
    doc.add_page_break()


def abstract(doc: Document) -> None:
    doc.add_heading("Abstract", level=1)
    doc.add_paragraph(
        "This report studies whether explicit volatility-regime information improves a PPO-based high-frequency "
        "market-making agent in a controlled synthetic limit-order-book environment. The simulator uses an arithmetic "
        "Brownian-motion mid-price, a sticky Markov volatility-regime process, rolling realized-volatility detection, "
        "and Poisson-arrival fills. The PPO agent chooses bid and ask quotes through a discrete half-spread and skew "
        "parameterization, and it is compared with a fixed-spread strategy and an Avellaneda-Stoikov analytical baseline."
    )
    doc.add_paragraph(
        CENTRAL_CLAIM
        + " The evidence comes from the canonical WP5/WP6 experiments: out-of-sample evaluation, detector robustness, "
        "oracle-label ablations, regime-conditional reward shaping, mild model misspecification, and a signal-informativeness "
        "sweep. The strongest interpretation is therefore conditional and synthetic-market bounded: the explicit regime "
        "label mostly repackages volatility information already present in the continuous signal. The report makes no "
        "live-trading deployment claim and does not claim to prove the internal PPO mechanism."
    )


def front_matter(doc: Document) -> None:
    doc.add_heading("Symbols, Abbreviations, and Glossary", level=1)
    add_table(
        doc,
        ["Symbol", "Meaning"],
        [
            ["S_t", "Mid-price at time step t."],
            ["q_t", "Inventory after step t."],
            ["X_t", "Cash account after step t."],
            ["W_t", "Marked-to-market wealth or equity, W_t = X_t + q_t S_t."],
            ["sigma_hat", "Rolling realized-volatility estimate observed by the PPO policy."],
            ["R_t", "Reward at time t."],
            ["h", "Quoted half-spread in ticks."],
            ["m", "Quote skew in ticks."],
            ["A, k", "Poisson fill-intensity scale and decay parameters."],
            ["eta", "Inventory penalty coefficient."],
        ],
        [1600, 7760],
    )
    add_table(
        doc,
        ["Abbreviation", "Meaning"],
        [
            ["ABM", "Arithmetic Brownian motion."],
            ["AS", "Avellaneda-Stoikov baseline."],
            ["HFMM", "High-frequency market making."],
            ["OOS", "Out-of-sample evaluation."],
            ["PPO", "Proximal Policy Optimization."],
            ["RV", "Realized volatility."],
            ["TOST", "Two One-Sided Tests equivalence procedure."],
            ["WP", "Work package."],
        ],
        [1800, 7560],
    )
    add_table(
        doc,
        ["Term", "Definition"],
        [
            ["Regime-aware PPO", "PPO policy variant that receives a categorical regime one-hot channel."],
            ["Regime-blind PPO", "PPO policy variant that omits the regime one-hot but still observes sigma_hat."],
            ["sigma_only", "Ablation variant using the continuous volatility signal without categorical regime labels."],
            ["combined", "Ablation variant using sigma_hat and estimated categorical regime labels."],
            ["oracle_full", "Ablation variant using sigma_hat and the true regime label."],
        ],
        [2200, 7160],
    )
    doc.add_page_break()


def introduction(doc: Document) -> None:
    doc.add_heading("1 Introduction", level=1)
    doc.add_heading("1.1 Problem Statement and Motivation", level=2)
    doc.add_paragraph(
        "Market makers continuously choose bid and ask quotes while balancing spread capture against inventory risk. "
        "When volatility changes, the same quote width can become either too passive or too exposed. A natural response "
        "is to provide a learning agent with regime labels such as low, medium, and high volatility. The central question "
        "is whether that categorical label adds decision-relevant information when the policy already observes a continuous "
        "volatility estimate."
    )
    doc.add_heading("1.2 Research Objective", level=2)
    doc.add_paragraph(
        "The objective is to test the incremental value of explicit categorical volatility-regime labels in PPO market "
        "making under a controlled synthetic environment. The comparison is intentionally information-design focused: "
        "the regime label is evaluated against a continuous sigma_hat signal, not against an observation space with no "
        "volatility information."
    )
    doc.add_heading("1.3 Background and Related Work", level=2)
    doc.add_paragraph(
        "Classical market-making models, especially the Avellaneda-Stoikov framework and later inventory-risk extensions, "
        "formalize the tradeoff between spread income and inventory exposure. They provide analytical baselines and clarify "
        "why volatility and inventory jointly affect reservation prices and spreads."
    )
    doc.add_paragraph(
        "Reinforcement-learning approaches to market making replace closed-form quoting rules with learned policies that "
        "can condition on state variables, simulated fills, and reward shaping. Prior DRL market-making work motivates the "
        "use of PPO-style policies, but it also raises a representation question: adding state channels can help only when "
        "they provide incremental information that the policy can exploit."
    )
    doc.add_paragraph(
        "Volatility regimes are a common way to summarize non-stationarity. In this thesis, regimes are generated by a "
        "three-state Markov process and estimated by rolling realized volatility. The literature motivates regime awareness, "
        "but the tested design asks a narrower question: whether the categorical regime label improves performance beyond "
        "the continuous realized-volatility proxy."
    )
    doc.add_paragraph(
        "Because the main empirical result is a bounded null/equivalence-style finding, the report uses paired tests and "
        "TOST equivalence tests rather than relying only on non-significant p-values. This supports a defense-safe reading "
        "of practical equivalence where the evidence warrants it."
    )
    doc.add_heading("1.4 Contributions", level=2)
    add_numbered(
        doc,
        [
            "A controlled HFMM simulator with Poisson-arrival fills, fees, latency, and Markov volatility regimes.",
            "A Gymnasium environment in which PPO controls quote half-spread and skew.",
            "An OOS comparison of fixed-spread, Avellaneda-Stoikov, regime-aware PPO, and regime-blind PPO strategies.",
            "A five-variant ablation separating continuous volatility information from estimated and oracle regime labels.",
            "Detector, reward-shaping, misspecification, and signal-informativeness checks that bound the interpretation.",
        ],
    )
    doc.add_heading("1.5 Scope and Limitations", level=2)
    doc.add_paragraph(
        "The thesis is a controlled synthetic-market study. It does not claim live trading viability, real-order-book "
        "external validity, or a proven PPO representation mechanism. The result is conditional on the simulator, signal "
        "design, PPO hyperparameters, and degradation calibration tested here."
    )
    doc.add_heading("1.6 Report Structure", level=2)
    doc.add_paragraph(
        "Sections 2-6 define the mathematical model, use case, architecture, layered model, and algorithms. Sections 7-9 "
        "describe the experimental setup, results, and metrics. Sections 10-12 discuss the interpretation, reproducibility, "
        "and conclusion. Appendices provide code maps, sanity checks, and extended evidence."
    )


def mathematical_formulation(doc: Document) -> None:
    doc.add_heading("2 Mathematical Formulation", level=1)
    doc.add_heading("2.1 Synthetic Mid-Price Process", level=2)
    add_equation(
        doc,
        "S_{t+1} = S_t + sigma_t sqrt(dt) epsilon_t,   epsilon_t ~ N(0, 1)",
        "Here S_t is the mid-price, sigma_t is the step volatility in ticks or price units according to the simulator configuration, dt is the time step, and epsilon_t is an independent standard-normal shock.",
    )
    doc.add_heading("2.2 Markov Volatility-Regime Process", level=2)
    add_equation(
        doc,
        "P(z_{t+1}=j | z_t=i) = P_ij,   z_t in {L, M, H}",
        "Here z_t is the latent volatility regime at step t, L/M/H denote low, medium, and high volatility, and P_ij is the sticky transition probability from regime i to regime j.",
    )
    doc.add_paragraph(
        "The regime controls the volatility multiplier applied to the base sigma parameter. The canonical full experiments "
        "use sigma multipliers [0.6, 1.0, 1.8] for L, M, and H."
    )
    doc.add_heading("2.3 Realized Volatility Signal and Regime Detection", level=2)
    add_equation(
        doc,
        "sigma_hat_t = sqrt((1 / w) sum_{i=t-w+1}^{t} (Delta S_i)^2)",
        "Here sigma_hat_t is the rolling realized-volatility signal, w is the rolling window length, and Delta S_i is the mid-price increment over step i.",
    )
    doc.add_paragraph(
        "Estimated regime labels are assigned by thresholding sigma_hat after a warmup period. The main WP4/WP5/WP6 "
        "pipelines use the causal rv_baseline detector. The rv_dwell detector is retained only as an auxiliary/offline "
        "robustness comparison, while the HMM detector is an additional robustness variant."
    )
    doc.add_heading("2.4 Limit-Order Fill Model", level=2)
    add_equation(
        doc,
        "lambda(delta) = A exp(-k delta)",
        "Here lambda(delta) is the fill intensity at quote distance delta, A is the baseline arrival scale, and k controls the exponential decay as quotes move farther from the mid-price.",
    )
    add_equation(
        doc,
        "P(fill | delta) = 1 - exp(-lambda(delta) dt)",
        "Here P(fill | delta) is the per-step fill probability and dt is the simulator step length.",
    )
    doc.add_heading("2.5 Quote Parameterization", level=2)
    add_equation(
        doc,
        "h = h_idx + 1,   m = m_idx - 2",
        "Here h_idx and m_idx are the two discrete PPO action components, h is the half-spread in ticks, and m is the quote skew in ticks.",
    )
    add_equation(
        doc,
        "delta_bid = max(1, h + m),   delta_ask = max(1, h - m)",
        "Here delta_bid and delta_ask are the bid and ask quote distances in ticks. The max operator enforces a minimum quote distance of one tick.",
    )
    doc.add_heading("2.6 Inventory, Cash, Equity, and Reward", level=2)
    add_equation(
        doc,
        "q_{t+1} = q_t + F_t^{bid} - F_t^{ask}",
        "Here q_t is inventory, F_t^{bid} is the bid-side fill indicator or count, and F_t^{ask} is the ask-side fill indicator or count.",
    )
    add_equation(
        doc,
        "X_{t+1} = X_t - F_t^{bid} P_t^{bid} + F_t^{ask} P_t^{ask} - fees_t",
        "Here X_t is cash, P_t^{bid} and P_t^{ask} are the executed bid and ask prices, and fees_t denotes transaction costs.",
    )
    add_equation(
        doc,
        "W_t = X_t + q_t S_t",
        "Here W_t is mark-to-market wealth or equity, X_t is cash, q_t is inventory, and S_t is the mid-price.",
    )
    add_equation(
        doc,
        "R_t = W_{t+1} - W_t - eta q_{t+1}^2",
        "Here R_t is the reward and eta is the inventory-penalty coefficient. Fees are included in the cash update and are not counted a second time in the reward.",
    )
    doc.add_heading("2.7 Avellaneda-Stoikov Baseline", level=2)
    add_equation(
        doc,
        "r_t = S_t - q_t gamma sigma^2 tau",
        "Here r_t is the reservation price, gamma is inventory-risk aversion, sigma is volatility, and tau is remaining horizon.",
    )
    add_equation(
        doc,
        "delta_AS = 0.5 gamma sigma^2 tau + (1 / gamma) log(1 + gamma / k)",
        "Here delta_AS is the AS half-spread and k is the same fill-intensity decay parameter used in the execution model. The implemented deltas are clipped to configured minimum and maximum bounds.",
    )
    doc.add_heading("2.8 PPO Objective", level=2)
    add_equation(
        doc,
        "rho_t(theta) = pi_theta(a_t | s_t) / pi_theta_old(a_t | s_t)",
        "Here rho_t(theta) is the PPO probability ratio, pi_theta is the current policy, pi_theta_old is the behavior policy used to collect the rollout, a_t is the action, and s_t is the state.",
    )
    add_equation(
        doc,
        "L_CLIP(theta) = E_t[min(rho_t(theta) A_hat_t, clip(rho_t(theta), 1-eps, 1+eps) A_hat_t)]",
        "Here theta denotes policy parameters, rho_t(theta) is the PPO probability ratio, A_hat_t is the advantage estimate, and eps is the clipping parameter. The report uses this only as a concise training-objective reference.",
    )


def use_case_and_architecture(doc: Document) -> None:
    doc.add_heading("3 Use-Case Scenario", level=1)
    doc.add_heading("3.1 Synthetic HFMM Scenario", level=2)
    doc.add_paragraph(
        "The use case is a synthetic market maker posting one bid and one ask quote at each step. The environment supplies "
        "mid-price dynamics, volatility estimates, and regime labels; the agent receives stochastic fills and is evaluated "
        "on wealth, risk-adjusted performance, inventory tails, and fill behavior."
    )
    doc.add_heading("3.2 Trading Agents and Strategy Variants", level=2)
    add_table(
        doc,
        ["Strategy or variant", "Observed information", "Purpose"],
        [
            ["Naive", "No learned state dependence", "Fixed-spread baseline."],
            ["Avellaneda-Stoikov", "Inventory, volatility, horizon", "Analytical inventory-risk baseline."],
            ["ppo_aware", "sigma_hat and estimated regime one-hot", "Original regime-aware PPO."],
            ["ppo_blind", "sigma_hat without regime one-hot", "Original regime-blind PPO."],
            ["sigma_only", "Continuous sigma_hat only", "Tests whether sigma_hat alone carries the signal."],
            ["combined", "sigma_hat and estimated regime one-hot", "Tests estimated categorical incremental value."],
            ["oracle_full", "sigma_hat and true regime one-hot", "Tests perfect-label incremental value."],
            ["regime_only / oracle_pure", "Categorical labels without sigma_hat", "Anchor label-only variants."],
        ],
        [2100, 3300, 3960],
    )
    doc.add_heading("3.3 Information-Design Question", level=2)
    doc.add_paragraph(CENTRAL_CLAIM)

    doc.add_heading("4 System Architecture", level=1)
    doc.add_heading("4.1 Run Lifecycle and Reproducibility Layer", level=2)
    doc.add_paragraph(
        "`run.py` dispatches jobs from JSON configuration files, creates a timestamped run directory, snapshots the config, "
        "records git metadata, writes logs and metrics, and finalizes run status. Long WP6 jobs support resume validation "
        "against the saved config snapshot."
    )
    doc.add_heading("4.2 Synthetic Market and Detector Layer", level=2)
    doc.add_paragraph(
        "`src/wp1/sim.py` implements the fill and inventory simulator. `src/wp2/synth_regime.py` generates synthetic "
        "regime paths, mid-prices, rolling realized volatility, and detector outputs."
    )
    doc.add_heading("4.3 Gymnasium Environment Layer", level=2)
    doc.add_paragraph(
        "`src/wp3/env.py` wraps the simulator as a Gymnasium environment. The observation vector is "
        "[q_norm, sigma_hat, tau, regime_L, regime_M, regime_H]. Warmup, invalid, or disabled regime labels produce a "
        "zero one-hot vector rather than an artificial medium label."
    )
    doc.add_heading("4.4 Strategy and Training Layer", level=2)
    doc.add_paragraph(
        "WP4 trains PPO policies; WP5 evaluates strategies and ablations; WP6 runs the signal-informativeness sweep. "
        "The fixed-spread and AS baselines provide non-learning comparators."
    )
    doc.add_heading("4.5 Evaluation and Evidence Layer", level=2)
    doc.add_paragraph(
        "Evaluation writes CSV metrics, figures, and summaries. The thesis_31 adaptation embeds existing figures and "
        "quotes frozen numerical evidence; it does not regenerate evidence artifacts."
    )


def layered_model_and_algorithms(doc: Document) -> None:
    doc.add_heading("5 Layered Experimental Model", level=1)
    layers = [
        ("5.1 Layer 1: Market State Generation", "Generate Markov regimes and mid-price paths under controlled volatility multipliers."),
        ("5.2 Layer 2: Signal Construction", "Construct sigma_hat from rolling realized volatility and derive estimated regime labels."),
        ("5.3 Layer 3: Observation Encoding", "Encode inventory, volatility, time-to-horizon, and optional regime one-hot channels."),
        ("5.4 Layer 4: Action and Execution", "Decode PPO actions into bid/ask quote distances and sample Poisson-arrival fills."),
        ("5.5 Layer 5: Learning and Evaluation", "Train PPO on the chronological train segment and evaluate deterministic policies OOS."),
    ]
    for heading, text in layers:
        doc.add_heading(heading, level=2)
        doc.add_paragraph(text)

    doc.add_heading("6 Algorithm Specification and Pseudocode", level=1)
    doc.add_heading("6.1 Plain-Language Overview", level=2)
    doc.add_paragraph(
        "The experiment creates a synthetic market path, estimates volatility signals, trains policies on the first 70% of "
        "the path, and evaluates all strategies on the held-out 30%. Ablations modify which volatility channels the PPO "
        "policy can observe."
    )
    add_code_block(
        doc,
        "6.2 Synthetic Market and Regime Generation",
        [
            "Input: seed, transition matrix P, base volatility, volatility multipliers, n_steps",
            "Initialize z_0 and S_0",
            "For t = 0 ... n_steps - 1:",
            "    sample z_{t+1} from P[z_t]",
            "    set sigma_t = base_sigma * multiplier[z_t]",
            "    sample mid-price increment and update S_{t+1}",
            "Compute rolling sigma_hat after warmup window",
            "Calibrate thresholds on warmup data and assign regime_hat",
            "Output: exogenous table with mid, sigma_hat, regime_true, regime_hat",
        ],
    )
    add_code_block(
        doc,
        "6.3 Market-Making Environment Step",
        [
            "Input: action (h_idx, m_idx), current simulator state, exogenous row",
            "Decode h = h_idx + 1 and m = m_idx - 2",
            "Set delta_bid = max(1, h + m), delta_ask = max(1, h - m)",
            "Compute fill intensities and per-step fill probabilities",
            "Sample bid and ask fills",
            "Update inventory and cash, including fees",
            "Advance mid-price and compute W_{t+1}",
            "Return observation, reward = delta_equity - eta * inventory^2, done flag, info",
        ],
    )
    add_code_block(
        doc,
        "6.4 PPO Training and OOS Evaluation Protocol",
        [
            "For each seed and variant:",
            "    generate one exogenous synthetic path",
            "    split path chronologically into train and test segments",
            "    configure observation channels for the variant",
            "    train PPO on the train segment for the configured timesteps",
            "    evaluate the deterministic policy on the test segment",
            "    log Sharpe-like ratio, final equity, inventory p99, fill rate, and per-regime metrics",
            "Aggregate seed-paired statistics across variants",
        ],
    )
    add_code_block(
        doc,
        "6.5 WP6 Signal-Informativeness Sweep",
        [
            "For each degradation condition in {full, noisy, lagged, coarsened, none}:",
            "    transform sigma_hat according to the condition",
            "    for each valid variant and seed:",
            "        train PPO and evaluate OOS",
            "Aggregate condition-variant means and paired comparisons",
            "Test whether combined gains value as sigma_hat is degraded",
        ],
    )
    doc.add_heading("6.6 Reference Implementation Map", level=2)
    add_table(
        doc,
        ["Component", "Reference file"],
        [
            ["Run lifecycle", "src/run_context.py, run.py"],
            ["Simulator", "src/wp1/sim.py"],
            ["Regime generation and detectors", "src/wp2/synth_regime.py"],
            ["Gymnasium environment", "src/wp3/env.py"],
            ["PPO training", "src/wp4/job_w4_ppo.py"],
            ["WP5 evaluation", "src/wp5/job_w5_eval.py"],
            ["WP6 sweep", "src/wp6/job_w6_sweep_full.py"],
        ],
        [3300, 6060],
    )
    doc.add_heading("6.7 Complexity Analysis", level=2)
    doc.add_paragraph(
        "Simulation and deterministic evaluation are linear in the number of environment steps. PPO training cost is "
        "approximately linear in total timesteps, policy-network forward/backward passes, and the number of seed-variant "
        "cells. The WP6 full sweep is expensive because it multiplies conditions, variants, and seeds; this is why its "
        "outputs are treated as frozen evidence."
    )
    doc.add_heading("6.8 Executable Example Commands", level=2)
    add_code_block(
        doc,
        "Commands",
        [
            "python run.py --config config/w5_main.json",
            "python run.py --config config/w5_detector_full.json",
            "python run.py --config config/w5_eta_regime.json",
            "python run.py --config config/w5_misspec_mild.json",
            "python run.py --config config/w6_sweep_full.json",
            "python run.py --config config/w6_sweep_full.json --resume <run_id>",
        ],
    )


def experimental_setup(doc: Document) -> None:
    doc.add_heading("7 Experimental Setup", level=1)
    doc.add_heading("7.1 Canonical Configuration Protocol", level=2)
    doc.add_paragraph(
        "Canonical experiments are driven by JSON configs under `config/`. Shared parameters include mid0 = 100.0, "
        "tick_size = 0.01, dt = 0.2, baseline sigma_mid_ticks = 0.8, A = 5.0, k = 1.5, fee_bps = 0.2, and latency_steps = 1."
    )
    doc.add_heading("7.2 Data Generation and Preprocessing", level=2)
    doc.add_paragraph(
        "Each run generates synthetic mid-price and regime paths, computes rolling realized volatility, and assigns detector "
        "labels. The main reported pipelines use causal rv_baseline labels."
    )
    doc.add_heading("7.3 Train/Test Split", level=2)
    doc.add_paragraph(
        "WP5 and WP6 use a chronological 70/30 split on the exogenous series. PPO trains on the first segment and is "
        "evaluated deterministically on the held-out OOS segment."
    )
    doc.add_heading("7.4 Strategy Variants", level=2)
    doc.add_paragraph(
        "The main WP5 comparison evaluates naive, AS, ppo_aware, and ppo_blind. The ablation and WP6 designs use "
        "sigma_only, regime_only, combined, oracle_pure, and oracle_full."
    )
    doc.add_heading("7.5 Detector Variants", level=2)
    add_table(
        doc,
        ["Detector", "Role", "Defense-safe caveat"],
        [
            ["rv_baseline", "Main causal detector", "Rolling RV threshold; used in main WP4/WP5/WP6 pipelines."],
            ["rv_dwell", "Auxiliary robustness detector", "Offline dwell smoothing; not the main causal detector."],
            ["hmm", "Robustness detector", "Higher classification accuracy but no reliable PPO advantage."],
        ],
        [1800, 2500, 5060],
    )
    doc.add_heading("7.6 PPO Hyperparameters", level=2)
    add_table(
        doc,
        ["Hyperparameter", "Canonical full-run value"],
        [
            ["total_timesteps", "1,000,000"],
            ["learning_rate", "3e-4"],
            ["n_steps", "2048"],
            ["batch_size", "256"],
            ["n_epochs", "10"],
            ["gamma", "0.999"],
            ["clip_range", "0.2"],
            ["ent_coef", "0.01 for full WP5/WP6 runs"],
        ],
        [3000, 6360],
    )
    doc.add_heading("7.7 Protected Evidence and No-Rerun Policy", level=2)
    doc.add_paragraph(
        "The thesis_31 adaptation reuses frozen results. It does not rerun PPO, WP5, WP6, figure scripts, or protected "
        "CSV generation. Protected evidence hashes are documented in `EVIDENCE_MANIFEST.md` and repeated in the "
        "reproducibility checklist."
    )


def results(doc: Document) -> None:
    doc.add_heading("8 Results and Visualisation", level=1)
    doc.add_heading("8.1 Main WP5 OOS Results", level=2)
    doc.add_paragraph(
        "The main 20-seed OOS evaluation shows that PPO variants produce much higher risk-adjusted performance than the "
        "fixed-spread and AS baselines. AS can produce higher raw equity, but with substantially larger inventory exposure."
    )
    add_table(
        doc,
        ["Strategy", "Mean Sharpe-like", "Mean final equity", "Mean inv_p99", "Mean fill rate"],
        [
            ["ppo_aware", "0.715", "4.10", "2.00", "0.236"],
            ["ppo_blind", "0.740", "4.42", "2.05", "0.232"],
            ["AS", "0.105", "5.05", "29.95", "0.444"],
            ["naive", "0.127", "4.49", "21.20", "0.119"],
        ],
        [1900, 1800, 1900, 1800, 1960],
    )
    doc.add_paragraph(
        "AS has higher raw equity but carries substantially larger inventory tail risk; therefore the main comparison is "
        "risk-adjusted performance and inventory control, not raw equity alone."
    )
    add_figure(doc, "fig1", "Figure 1. Main OOS performance and inventory risk, reused from frozen thesis figures.")
    add_figure(doc, "fig2", "Figure 2. Seed-paired PPO-aware versus PPO-blind comparison.")

    doc.add_heading("8.2 Five-Variant Ablation", level=2)
    add_table(
        doc,
        ["Variant", "Observation meaning", "Mean Sharpe-like", "Mean final equity", "Mean inv_p99"],
        [
            ["ppo_sigma_only", "continuous sigma_hat only", "0.752987", "4.547923", "1.950000"],
            ["ppo_oracle_full", "sigma_hat + true regime one-hot", "0.722443", "4.067945", "2.050000"],
            ["ppo_regime_only", "estimated regime one-hot only", "0.697552", "4.005379", "2.150000"],
            ["ppo_combined", "sigma_hat + estimated regime one-hot", "0.696182", "3.905214", "1.800000"],
            ["ppo_oracle_pure", "true regime one-hot only", "0.683829", "3.925499", "1.950000"],
        ],
        [1800, 3000, 1500, 1500, 1560],
    )
    doc.add_paragraph(
        "The strongest ablation result is that sigma_only has the highest mean Sharpe-like value, while oracle_full does "
        "not significantly beat it. TOST supports practical equivalence under the +/-0.10 Sharpe-like bound."
    )
    add_table(
        doc,
        ["Comparison", "Metric or test", "Canonical result", "Defense-safe reading"],
        [
            ["PPO-aware vs PPO-blind", "Sharpe-like paired t-test", "p = 0.261", "No significant Sharpe improvement from explicit regime labels."],
            ["PPO-aware vs PPO-blind", "Final equity paired t-test", "p = 0.023", "Final equity favors PPO-blind."],
            ["ppo_sigma_only vs ppo_oracle_full", "Sharpe-like paired t-test", "p = 0.115", "No significant oracle-label Sharpe improvement over sigma_hat alone."],
            ["ppo_sigma_only vs ppo_oracle_full", "TOST +/-0.10", "p = 0.00067; 90% CI [-0.001, +0.063]", "Positive evidence of practical equivalence under the stated bound."],
            ["HMM detector", "Detector accuracy", "81.8%", "Higher detector accuracy still does not create a reliable PPO advantage."],
            ["Detector robustness", "ANOVA", "p = 0.997", "Detector choice does not explain the null result."],
            ["Regime-conditional eta", "sigma_only vs combined Sharpe", "p = 0.0016", "Result favors sigma_only."],
            ["Mild misspecification", "sigma_only vs oracle_full Sharpe t-test", "p = 0.881", "No significant Sharpe difference."],
            ["Mild misspecification", "TOST +/-0.05", "p = 0.042", "Practical equivalence is supported under the stated bound."],
        ],
        [2400, 2100, 2300, 2560],
    )
    doc.add_paragraph(
        "The t-test results show no significant Sharpe improvement from explicit regime labels, while the TOST result "
        "provides positive evidence of practical equivalence under the stated bound."
    )
    add_figure(doc, "fig3", "Figure 3. Five-variant ablation summary.")
    add_figure(doc, "fig4", "Figure 4. Oracle-label paired-seed comparison.")

    doc.add_heading("8.3 Detector Robustness", level=2)
    add_table(
        doc,
        ["Detector", "Aware vs blind conclusion", "Key statistic"],
        [
            ["rv_baseline", "No reliable aware Sharpe advantage", "p = 0.1142"],
            ["rv_dwell", "No reliable aware Sharpe advantage", "p = 0.1095"],
            ["hmm", "No reliable aware Sharpe advantage despite higher accuracy", "p = 0.0822"],
            ["ANOVA", "Detector choice does not explain the result", "F = 0.0034, p = 0.9966"],
        ],
        [1900, 4800, 2660],
    )
    add_figure(doc, "fig5", "Figure 5. Detector robustness across rv_baseline, rv_dwell, and HMM.")

    doc.add_heading("8.4 Reward-Shaping and Misspecification Checks", level=2)
    doc.add_paragraph(
        "The regime-conditional eta run tests whether explicit labels become useful when the reward penalizes high-volatility "
        "inventory more strongly. sigma_only still beats combined on Sharpe-like performance (p = 0.0016). Under mild "
        "regime-dependent execution misspecification, sigma_only and oracle_full remain statistically indistinguishable "
        "and practically equivalent under the reported TOST bound."
    )
    add_figure(doc, "fig6", "Figure 6. Regime-conditional eta summary.")
    add_figure(doc, "fig7", "Figure 7. Mild model-misspecification summary.")

    doc.add_heading("8.5 WP6 Signal-Informativeness Sweep", level=2)
    add_table(
        doc,
        ["condition", "sigma_only", "combined", "regime_only", "oracle_full", "oracle_pure"],
        [
            ["full", "0.7626 +/- 0.0596", "0.6898 +/- 0.0607", "0.6991 +/- 0.0668", "0.6811 +/- 0.0697", "0.6632 +/- 0.0769"],
            ["noisy", "0.7563 +/- 0.0616", "0.7125 +/- 0.0628", "0.6991 +/- 0.0668", "0.6736 +/- 0.1013", "0.6632 +/- 0.0769"],
            ["lagged", "0.7822 +/- 0.0454", "0.6812 +/- 0.0536", "0.6991 +/- 0.0668", "0.7264 +/- 0.0686", "0.6632 +/- 0.0769"],
            ["coarsened", "0.7827 +/- 0.0533", "0.6908 +/- 0.0827", "0.6991 +/- 0.0668", "0.7447 +/- 0.0549", "0.6632 +/- 0.0769"],
            ["none", "undefined", "0.6991 +/- 0.0668", "0.6991 +/- 0.0668", "0.6632 +/- 0.0769", "0.6632 +/- 0.0769"],
        ],
        [1450, 1600, 1600, 1600, 1600, 1510],
    )
    doc.add_paragraph(
        "WP6 did not support the original informativeness-threshold hypothesis within the tested calibration band. "
        "sigma_only remained high under informative conditions, while combined was directionally below sigma_only."
    )
    add_figure(doc, "fig8", "Figure 8. WP6 monotonic-gap plot.")
    add_figure(doc, "fig9", "Figure 9. WP6 paired-seed combined versus sigma_only.")
    add_figure(doc, "fig10", "Figure 10. WP6 paired-seed combined versus regime_only.")


def metrics_discussion_repro(doc: Document) -> None:
    doc.add_heading("9 Evaluation Metrics", level=1)
    doc.add_heading("9.1 Final Equity", level=2)
    doc.add_paragraph("Final equity is the terminal marked-to-market wealth. It is economically meaningful but not sufficient alone because high raw equity can be earned by taking large inventory risk.")
    doc.add_heading("9.2 Sharpe-Like Ratio", level=2)
    add_equation(
        doc,
        "Sharpe-like = mean(r_t) / std(r_t)",
        "Here r_t denotes the per-step or per-episode return series used by the project metric convention. This is the primary risk-adjusted performance metric.",
    )
    doc.add_heading("9.3 Inventory Tail Risk", level=2)
    add_equation(
        doc,
        "inv_p99 = percentile_99(|q_t|)",
        "Here q_t is inventory. inv_p99 is a practical tail-risk diagnostic for whether a strategy earns returns by carrying large inventory.",
    )
    doc.add_heading("9.4 Fill Rate", level=2)
    add_equation(
        doc,
        "fill_rate = filled quote opportunities / total quote opportunities",
        "The numerator and denominator follow the logged project convention. Fill rate is diagnostic rather than the primary thesis metric.",
    )
    doc.add_heading("9.5 Paired t-Tests", level=2)
    add_equation(
        doc,
        "t = mean(d_i) / (sd(d_i) / sqrt(n))",
        "Here d_i is the paired seed-level difference between two strategies, sd(d_i) is its sample standard deviation, and n is the number of paired seeds.",
    )
    doc.add_heading("9.6 TOST Equivalence Tests", level=2)
    doc.add_paragraph(
        "TOST evaluates whether the plausible range of a paired difference lies inside a pre-specified practical-equivalence "
        "band. At alpha = 0.05, the corresponding equivalence reading uses the 90% confidence interval. WP5 uses TOST "
        "to support practical equivalence for sigma_only versus oracle_full; WP6 also uses related paired tests to examine "
        "non-equivalence and mean indistinguishability patterns."
    )

    doc.add_heading("10 Discussion", level=1)
    doc.add_heading("10.1 Main Interpretation: Signal Redundancy", level=2)
    doc.add_paragraph(CENTRAL_CLAIM)
    doc.add_paragraph(
        "The interpretation is that the continuous realized-volatility proxy already supplies the economically useful "
        "volatility information needed by the policy in this environment. The categorical label may be coarser or redundant "
        "when presented alongside sigma_hat."
    )
    doc.add_heading("10.2 Why the Result Is Not a Weak Null", level=2)
    doc.add_paragraph(
        "The conclusion does not rest only on p > 0.05. It is supported by detector robustness, oracle labels, TOST "
        "equivalence, reward-channel checks, mild misspecification, and WP6 degradation tests."
    )
    doc.add_heading("10.3 Synthetic-Market Boundary", level=2)
    doc.add_paragraph(
        "The result is bounded to the tested synthetic HFMM simulator. It does not imply that volatility-regime structure "
        "lacks importance in finance or that categorical signals cannot help in other market-making designs."
    )
    doc.add_heading("10.4 Risks to Interpretation", level=2)
    add_table(
        doc,
        ["Risk", "Mitigation"],
        [
            ["Overclaiming mechanism", "Use signal-redundancy and categorical-channel degradation as interpretations, not proofs of PPO internals."],
            ["Detector causality confusion", "State that rv_baseline is the main causal detector and rv_dwell is auxiliary/offline."],
            ["Synthetic external validity", "Frame all claims as controlled synthetic-market claims."],
            ["Template over-compression", "Keep the literature and evidence spine, but map it into template headings."],
        ],
        [3000, 6360],
    )
    doc.add_heading("10.5 Future Work", level=2)
    add_bullets(
        doc,
        [
            "Evaluate stronger regime-dependent execution misspecification.",
            "Test real or richer limit-order-book data after thesis scope, not as a claim of this report.",
            "Run representation-level diagnostics only as future mechanism work.",
            "Explore alternative observation encodings and policy architectures.",
        ],
    )

    doc.add_heading("11 Reproducibility Checklist", level=1)
    add_table(
        doc,
        ["Item", "Status in thesis_31"],
        [
            ["Frozen baseline", "thesis-v29-frozen tag; commit 9681faa; thesis_29 remains untouched."],
            ["Template draft", "thesis_31 is a new template-adapted draft, not a replacement for thesis_29 or thesis_30."],
            ["Run ID structure", "YYYYMMDD-HHMMSS_seed<S>_<TAG>_<COMMIT>."],
            ["Config snapshots", "Every run snapshots config; config_snapshot_all.md inventories active config files."],
            ["Seeds", "WP5 main uses seeds 1-20; WP6 full uses seeds 42-61."],
            ["Train/test split", "70/30 chronological split on exogenous synthetic series."],
            ["Evidence manifest", "EVIDENCE_MANIFEST.md records protected evidence and remediation invariants."],
            ["Protected CSV hashes", "Listed below and verified by file hash checks."],
            ["Codebase snapshot", "docs/internal/codebase_snapshot.py."],
            ["No-rerun policy", "No PPO/WP5/WP6 reruns or figure regeneration during template adaptation."],
        ],
        [2800, 6560],
    )
    hash_rows = []
    for rel, expected in PROTECTED_HASHES.items():
        path = ROOT / rel
        actual = hashlib.sha256(path.read_bytes()).hexdigest().upper() if path.exists() else "MISSING"
        status = "MATCH" if actual == expected else "MISMATCH"
        hash_rows.append([rel, expected[:12] + "...", status])
    add_table(doc, ["Protected artifact", "Expected SHA256", "Check"], hash_rows, [5300, 2300, 1760])
    add_table(
        doc,
        ["Reproducibility command", "Purpose"],
        [
            ["python -m venv .venv", "Create environment."],
            ["& .\\.venv\\Scripts\\Activate.ps1", "Activate environment on Windows PowerShell."],
            ["pip install -r requirements.txt", "Install dependencies."],
            ["python run.py --config config/w5_main.json", "Run main WP5 evaluation."],
            ["python run.py --config config/w6_sweep_full.json", "Run WP6 sweep."],
            ["python run.py --config config/w6_sweep_full.json --resume <run_id>", "Resume WP6 sweep."],
            ["ruff check src/", "Lint source code."],
        ],
        [4200, 5160],
    )
    doc.add_paragraph(
        "These commands document reproducibility paths only; thesis_31 adaptation did not rerun experiments."
    )

    doc.add_heading("12 Conclusion", level=1)
    doc.add_paragraph(
        "This template-adapted report preserves the thesis finding while reorganizing the manuscript into the requested "
        "technical-report structure. PPO learns strong risk-adjusted quoting behavior in the synthetic HFMM simulator, but "
        "the explicit categorical regime channel does not provide robust incremental value beyond sigma_hat. The result is "
        "best defended as evidence consistent with signal redundancy in the tested controlled environment."
    )


def appendices(doc: Document) -> None:
    doc.add_heading("Appendix A Code Appendix", level=1)
    add_table(
        doc,
        ["Area", "Files"],
        [
            ["Run management", "run.py; src/run_context.py"],
            ["Simulation", "src/wp1/sim.py; src/wp1/w1_as_baseline.py"],
            ["Regimes", "src/wp2/synth_regime.py; src/wp2/job_w2_synth.py"],
            ["Environment", "src/wp3/env.py; src/wp3/w3_sanity.py"],
            ["PPO training", "src/wp4/job_w4_ppo.py"],
            ["Evaluation", "src/wp5/job_w5_eval.py; src/wp5/job_w5_detector_compare.py"],
            ["Signal audit and sweep", "src/wp5_5/*; src/wp6/job_w6_sweep_full.py"],
        ],
        [2300, 7060],
    )
    add_code_block(
        doc,
        "Canonical commands",
        [
            "python run.py --config config/w5_main.json",
            "python run.py --config config/w5_detector_full.json",
            "python run.py --config config/w5_eta_regime.json",
            "python run.py --config config/w5_misspec_mild.json",
            "python run.py --config config/w6_sweep_full.json",
        ],
    )
    doc.add_heading("Appendix B Sanity Checks and Unit Tests", level=1)
    add_table(
        doc,
        ["Check", "Purpose"],
        [
            ["tests/test_csv_metric_logger.py", "Verifies stable CSV metric schema behavior."],
            ["tests/test_resume_validation.py", "Verifies resume-mode config validation."],
            ["WP3 sanity checks", "Exercises naive, AS, and random policies through the Gymnasium environment."],
            ["Figure/path checks", "Ensures thesis_31 references existing frozen figure files."],
            ["Forbidden phrase scan", "Guards against overclaims in the generated draft."],
        ],
        [3300, 6060],
    )
    doc.add_heading("Appendix C Extended Evidence Tables", level=1)
    doc.add_paragraph(
        "This appendix contains diagnostic material demoted from the main body to keep the report template focused."
    )
    add_figure(doc, "figA", "Appendix Figure C1. Regime-wise action distribution diagnostic.")
    add_figure(doc, "figB", "Appendix Figure C2. Regime-wise Sharpe diagnostic.")
    add_table(
        doc,
        ["Diagnostic", "Status", "Interpretation"],
        [
            ["Post-hoc signal diagnostics", "Supporting only", "Consistent with signal redundancy but not primary evidence."],
            ["WP5.5 signal audit", "Offline calibration", "Supports WP6 degradation design; no PPO training."],
            ["Detector robustness statistics", "Protected evidence", "Addresses poor-detector objection without changing main detector caveat."],
            ["WP6 paired summaries", "Protected evidence", "Supports refined interpretation of the signal-informativeness sweep."],
        ],
        [2800, 2200, 4360],
    )


def validate_inputs() -> None:
    missing = [str(path.relative_to(ROOT)) for path in FIGURES.values() if not path.exists()]
    if missing:
        raise FileNotFoundError("Missing figure inputs: " + ", ".join(missing))
    missing_hash = [rel for rel in PROTECTED_HASHES if not (ROOT / rel).exists()]
    if missing_hash:
        raise FileNotFoundError("Missing protected artifacts: " + ", ".join(missing_hash))


def main() -> None:
    validate_inputs()
    doc = Document()
    style_document(doc)
    title_page(doc)
    abstract(doc)
    front_matter(doc)
    introduction(doc)
    mathematical_formulation(doc)
    use_case_and_architecture(doc)
    layered_model_and_algorithms(doc)
    experimental_setup(doc)
    results(doc)
    metrics_discussion_repro(doc)
    appendices(doc)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(f"Wrote {OUT.relative_to(ROOT)}")


if __name__ == "__main__":
    sys.exit(main())
